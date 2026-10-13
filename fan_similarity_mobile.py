import streamlit as st
import pandas as pd
import numpy as np
import plotly.graph_objects as go
from plotly.subplots import make_subplots
import io
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
import datetime
import json
from pathlib import Path
import hashlib
import sqlite3

# -----------------------------------------------------------------------------
# 물리/모델 상수 (magic number 중앙화)
# -----------------------------------------------------------------------------
R_SPECIFIC_AIR = 287.05  # 건공기 기체상수 [J/(kg·K)]
AIR_VISCOSITY = 1.81e-5  # 공기 점성계수(동적 점도) 근사 [Pa·s]
ACKERET_EXP = 0.2  # Ackeret 효율 보정 지수

CENTRIFUGAL_CATEGORY = "원심팬 계열 (터보, 시로코 등)"
AXIAL_CATEGORY = "축류팬 계열 (Axial)"
CROSSFLOW_CATEGORY = "횡류팬 계열 (Cross-flow)"

INPUT_STATE_PATH = Path(__file__).with_name(".fan_similarity_mobile_inputs.sqlite3")
LEGACY_INPUT_STATE_PATH = Path(__file__).with_name(".fan_similarity_v9_inputs.json")


def read_project_inputs(path, project="default", legacy_path=None):
    """작업별 저장값을 읽는다. 기본 작업의 v9 JSON은 읽기 전용으로 이관한다."""
    state = {}
    if path.exists():
        with sqlite3.connect(path, timeout=10) as connection:
            connection.execute(
                "CREATE TABLE IF NOT EXISTS inputs "
                "(project TEXT, name TEXT, value TEXT, PRIMARY KEY(project, name))"
            )
            state = {
                name: json.loads(value)
                for name, value in connection.execute(
                    "SELECT name, value FROM inputs WHERE project=?", (project,)
                )
            }
    if not state and project == "default" and legacy_path and legacy_path.exists():
        loaded = json.loads(legacy_path.read_text(encoding="utf-8"))
        state = loaded if isinstance(loaded, dict) else {}
    return state


def write_project_inputs(path, project, updates):
    """변경한 필드만 한 트랜잭션으로 저장해 다른 필드의 동시 편집을 보존한다."""
    encoded = [
        (project, key, json.dumps(value, ensure_ascii=False, allow_nan=False))
        for key, value in updates.items()
    ]
    with sqlite3.connect(path, timeout=10) as connection:
        connection.execute(
            "CREATE TABLE IF NOT EXISTS inputs "
            "(project TEXT, name TEXT, value TEXT, PRIMARY KEY(project, name))"
        )
        connection.executemany(
            "INSERT INTO inputs VALUES (?, ?, ?) ON CONFLICT(project, name) "
            "DO UPDATE SET value=excluded.value",
            encoded,
        )


def _read_input_state():
    """현재 세션의 작업 스냅샷을 유지하고 새 세션에서는 저장값을 복원한다."""
    project = st.session_state.get("input_project", "default").strip() or "default"
    snapshots = st.session_state.setdefault("_input_snapshots", {})
    if project in snapshots:
        return dict(snapshots[project])
    try:
        state = read_project_inputs(INPUT_STATE_PATH, project, LEGACY_INPUT_STATE_PATH)
        # v9 입력은 최초 v10 작업 저장 시 함께 이관한다.
        snapshots[project] = state
        return dict(state)
    except (OSError, ValueError, TypeError, sqlite3.Error) as error:
        st.warning(f"저장된 입력을 읽지 못했습니다: {error}")
        return {}


def _persist_input_state(**updates):
    """세션에서 변경한 필드만 저장한다. 같은 필드는 마지막 저장이 우선한다."""
    project = st.session_state.get("input_project", "default").strip() or "default"
    state = _read_input_state()
    changed = {key: value for key, value in updates.items() if key not in state or state[key] != value}
    if state.get("__schema_version__") != 10:
        changed = {**state, **changed, "__schema_version__": 10}
    if not changed:
        return
    try:
        write_project_inputs(INPUT_STATE_PATH, project, changed)
        state.update(changed)
        st.session_state["_input_snapshots"][project] = state
    except (OSError, ValueError, TypeError, sqlite3.Error) as error:
        st.warning(f"입력을 저장하지 못했습니다. 현재 화면의 계산은 계속합니다: {error}")


def _unwrap_editor_scalar(value):
    """Streamlit 1.52.x가 빈 열 입력을 [값]으로 반환한 데이터를 복구한다."""
    if isinstance(value, (list, tuple, np.ndarray, pd.Series)):
        values = list(value)
        if len(values) == 1:
            return _unwrap_editor_scalar(values[0])
        # 이 앱의 편집표는 모두 스칼라 셀이므로 다중값 컨테이너는 결측 처리한다.
        return None
    if isinstance(value, dict):
        return None
    return value


def _normalize_editor_dataframe(df, numeric_columns=()):
    """Excel 붙여넣기·구버전 저장값을 스칼라 숫자형 편집표로 정규화한다."""
    normalized = df.copy()
    for column in normalized.columns:
        normalized[column] = normalized[column].map(_unwrap_editor_scalar)
    for column in numeric_columns:
        if column in normalized.columns:
            normalized[column] = pd.to_numeric(normalized[column], errors="coerce").astype(float)
    return normalized.replace([np.inf, -np.inf], np.nan)


def _state_dataframe(state, name, fallback, numeric_columns=()):
    records = state.get(name)
    if not isinstance(records, list):
        return _normalize_editor_dataframe(fallback, numeric_columns)
    try:
        restored = pd.DataFrame(records)
        for column in fallback.columns:
            if column not in restored.columns:
                restored[column] = np.nan
        restored = restored[list(fallback.columns)]
        restored = _normalize_editor_dataframe(restored, numeric_columns)
        # 구버전 신규 직접입력 표에 들어 있던 63~71 dB 값은 실제 측정값이
        # 아니라 UI 예시값이었다. 전체 예시 PQ와 함께 그대로 저장된 경우에만
        # 소음을 결측으로 마이그레이션해 실제 데이터로 오인하지 않는다.
        if name == "set2_new_direct_curve":
            old_q = np.array([0, 12, 24, 36, 48, 60, 72, 84, 96], dtype=float)
            old_p = np.array([260, 255, 238, 210, 175, 130, 82, 35, 0], dtype=float)
            old_lw = np.array([63, 62, 61, 60, 61, 63, 65, 68, 71], dtype=float)
            if len(restored) == len(old_q):
                saved_q = pd.to_numeric(restored["풍량 Q (CMM)"], errors="coerce").to_numpy(float)
                saved_p = pd.to_numeric(restored["정압 P (Pa)"], errors="coerce").to_numpy(float)
                saved_lw = pd.to_numeric(restored["소음 Lw (dB(A))"], errors="coerce").to_numpy(float)
                if (
                    np.allclose(saved_q, old_q, equal_nan=False)
                    and np.allclose(saved_p, old_p, equal_nan=False)
                    and np.allclose(saved_lw, old_lw, equal_nan=False)
                ):
                    restored["소음 Lw (dB(A))"] = np.nan
        return restored
    except (ValueError, TypeError):
        return _normalize_editor_dataframe(fallback, numeric_columns)


def _dataframe_records(df, numeric_columns=()):
    """편집표를 저장하기 전에 리스트 셀 유입을 차단하고 JSON 스칼라로 만든다."""
    normalized = _normalize_editor_dataframe(df, numeric_columns)
    clean = normalized.astype(object).where(pd.notna(normalized), None)
    return clean.to_dict(orient="records")


def _numeric_editor_column_config(columns):
    """숫자 표의 스키마를 명시해 Excel 범위 붙여넣기 타입을 안정화한다."""
    config = {
        column: st.column_config.NumberColumn(format="%.0f" if column == "RPM" else "%.2f")
        for column in columns
    }
    efficiency_column = "정압효율 ηs (%)"
    if efficiency_column in config:
        config[efficiency_column] = st.column_config.NumberColumn(
            min_value=0.0, max_value=100.0, format="%.2f"
        )
    return config


def is_axial_category(fan_category):
    """축류팬 여부를 한 곳에서 판정하여 UI와 계산 모델을 일치시킨다."""
    return fan_category == AXIAL_CATEGORY


def create_ppt_report(
    fan_category,
    dim_label,
    base_D,
    base_X,
    base_N,
    base_Q,
    base_P,
    base_W,
    base_noise,
    target_D,
    target_X,
    target_N_calc,
    target_Q_res,
    target_P_res,
    target_W_res,
    db_t_sc,
    flow_unit,
    pressure_unit,
    power_unit,
    air_density_1,
    air_density_2,
    use_eff_correction,
    n_del,
    q_del,
    p_del,
    w_del,
    db_del,
    commentary_list,
):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    dark_gray = RGBColor(40, 40, 40)
    orange = RGBColor(13, 148, 136)  # v10 공통 강조색
    light_bg = RGBColor(248, 247, 245)  # F8F7F5
    blue = RGBColor(91, 138, 153)  # 5B8A99
    white = RGBColor(255, 255, 255)
    text_muted = RGBColor(120, 120, 120)

    # ----------------------------------------------------
    # SLIDE 1: Title (Cover)
    # ----------------------------------------------------
    blank_layout = prs.slide_layouts[6]
    slide1 = prs.slides.add_slide(blank_layout)
    bg_shape = slide1.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = light_bg
    bg_shape.line.fill.background()

    title_box = slide1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(3.0))
    tf = title_box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "FAN SIMILARITY SCALING REPORT"
    p.font.name = "Aptos"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = orange

    p2 = tf.add_paragraph()
    p2.text = "팬 상사법칙 기반 성능 및 소음 예측 보고서"
    p2.font.name = "Malgun Gothic"
    p2.font.size = Pt(36)
    p2.font.bold = True
    p2.font.color.rgb = dark_gray
    p2.space_before = Pt(10)
    p2.space_after = Pt(20)

    p3 = tf.add_paragraph()
    p3.text = f"분류: {fan_category} | 단위계: {flow_unit}, {pressure_unit}, {power_unit}"
    p3.font.name = "Malgun Gothic"
    p3.font.size = Pt(14)
    p3.font.color.rgb = text_muted

    p4 = tf.add_paragraph()
    today_str = datetime.date.today().strftime("%Y-%m-%d")
    p4.text = f"발행일: {today_str} | Generated by Engineering Toolkit"
    p4.font.name = "Aptos"
    p4.font.size = Pt(12)
    p4.font.color.rgb = text_muted
    p4.space_before = Pt(30)

    # ----------------------------------------------------
    # SLIDE 2: Comparison & Results (Table)
    # ----------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    title_box2 = slide2.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(11.9), Inches(1.0))
    tf2 = title_box2.text_frame
    p_t = tf2.paragraphs[0]
    p_t.text = "1:1 상사 스케일링 결과 비교"
    p_t.font.name = "Malgun Gothic"
    p_t.font.size = Pt(24)
    p_t.font.bold = True
    p_t.font.color.rgb = dark_gray

    p_st = tf2.add_paragraph()
    p_st.text = f"기준 스펙 대비 변경 스펙에서의 예측 거동 요약 정보"
    p_st.font.name = "Malgun Gothic"
    p_st.font.size = Pt(12)
    p_st.font.color.rgb = text_muted
    p_st.space_before = Pt(5)

    table_shape = slide2.shapes.add_table(8, 4, Inches(0.7), Inches(1.8), Inches(11.9), Inches(4.5))
    table = table_shape.table
    table.columns[0].width = Inches(3.5)
    table.columns[1].width = Inches(2.8)
    table.columns[2].width = Inches(2.8)
    table.columns[3].width = Inches(2.8)

    headers = ["구분 (Parameter)", "기준 사양 (Baseline)", "목표 사양 (Target)", "변화율 (Delta)"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = orange
        p = cell.text_frame.paragraphs[0]
        p.font.name = "Malgun Gothic"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = white
        p.alignment = PP_ALIGN.CENTER

    dimension_row = (
        ["독립 높이 H", "미적용", "미적용", "직경 D로 기하상사"]
        if is_axial_category(fan_category)
        else [
            f"팬 {dim_label} (Width/Length)",
            f"{base_X:.1f} mm",
            f"{target_X:.1f} mm",
            f"{((target_X/base_X)-1)*100:+.1f}%",
        ]
    )
    data = [
        [
            "팬 직경 D (Diameter)",
            f"{base_D:.1f} mm",
            f"{target_D:.1f} mm",
            f"{((target_D/base_D)-1)*100:+.1f}%",
        ],
        dimension_row,
        ["회전수 N (RPM)", f"{base_N:.0f} RPM", f"{target_N_calc:.0f} RPM", f"{n_del:+.1f}%"],
        [
            "풍량 Q (Flow Rate)",
            f"{base_Q:.2f} {flow_unit}",
            f"{target_Q_res:.2f} {flow_unit}",
            f"{q_del:+.1f}%",
        ],
        [
            "정압 P (Static Pressure)",
            f"{base_P:.2f} {pressure_unit}",
            f"{target_P_res:.2f} {pressure_unit}",
            f"{p_del:+.1f}%",
        ],
        [
            "축동력 W (Shaft Power)",
            f"{base_W:.1f} {power_unit}",
            f"{target_W_res:.1f} {power_unit}",
            f"{w_del:+.1f}%",
        ],
        [
            "소음 레벨 Lw (Sound Level)",
            f"{base_noise:.1f} dB(A)",
            f"{db_t_sc:.1f} dB(A)",
            f"{db_del:+.1f} dB(A)",
        ],
    ]

    for row_idx, row_data in enumerate(data):
        for col_idx, val in enumerate(row_data):
            cell = table.cell(row_idx + 1, col_idx)
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            p.font.name = "Malgun Gothic"
            p.font.size = Pt(11)
            p.alignment = PP_ALIGN.CENTER if col_idx > 0 else PP_ALIGN.LEFT
            cell.fill.solid()
            if row_idx % 2 == 0:
                cell.fill.fore_color.rgb = RGBColor(250, 249, 246)
            else:
                cell.fill.fore_color.rgb = white

            if col_idx == 3:
                p.font.bold = True
                # 동력(row 5)과 소음(row 6)은 '증가'가 나쁜 방향 → 증가 시 빨강, 감소 시 초록
                higher_is_bad = row_idx in (5, 6)
                if "+" in val:
                    p.font.color.rgb = RGBColor(198, 40, 40) if higher_is_bad else RGBColor(46, 125, 50)
                elif "-" in val:
                    p.font.color.rgb = RGBColor(46, 125, 50) if higher_is_bad else RGBColor(198, 40, 40)
            else:
                p.font.color.rgb = dark_gray

    # ----------------------------------------------------
    # SLIDE 3: Predictive Formulas & Equations
    # ----------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    title_box3 = slide3.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(11.9), Inches(1.0))
    tf3 = title_box3.text_frame
    p_t3 = tf3.paragraphs[0]
    p_t3.text = "예측 공식 및 물리 모델 근거"
    p_t3.font.name = "Malgun Gothic"
    p_t3.font.size = Pt(24)
    p_t3.font.bold = True
    p_t3.font.color.rgb = dark_gray

    p_st3 = tf3.add_paragraph()
    p_st3.text = f"수행된 1:1 예측에 사용된 확장 상사법칙(Affinity Laws) 수식 체계"
    p_st3.font.name = "Malgun Gothic"
    p_st3.font.size = Pt(12)
    p_st3.font.color.rgb = text_muted
    p_st3.space_before = Pt(5)

    box_width = Inches(5.7)
    box_height = Inches(4.5)

    box1 = slide3.shapes.add_textbox(Inches(0.7), Inches(1.8), box_width, box_height)
    tf_b1 = box1.text_frame
    tf_b1.word_wrap = True

    p = tf_b1.paragraphs[0]
    p.text = "■ 풍량 및 정압 예측 공식"
    p.font.name = "Malgun Gothic"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = orange

    p = tf_b1.add_paragraph()
    p.text = "1. 풍량 상사법칙 (Flow Rate Scaling):"
    p.font.name = "Malgun Gothic"
    p.font.size = Pt(12)
    p.font.bold = True
    p.space_before = Pt(10)

    p = tf_b1.add_paragraph()
    p.text = (
        "   Q2 = Q1 * (N2 / N1) * (D2 / D1)³"
        if is_axial_category(fan_category)
        else "   Q2 = Q1 * (N2 / N1) * (D2 / D1)² * (H2 / H1)"
    )
    p.font.name = "Courier New"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = blue

    p = tf_b1.add_paragraph()
    p.text = (
        "   - 축류팬은 독립 높이 변수를 사용하지 않으며, 기하학적으로 유사한 팬의 유량은 회전수의 1승과 직경의 3승에 비례합니다."
        if is_axial_category(fan_category)
        else "   - 직경의 2승과 폭(높이)의 1승에 비례하여 기하학적 유로 면적이 확대되고, 회전 속도에 비례해 토출 유량이 선형 스케일링됩니다."
    )
    p.font.name = "Malgun Gothic"
    p.font.size = Pt(10)
    p.font.color.rgb = dark_gray
    p.space_before = Pt(2)

    p = tf_b1.add_paragraph()
    p.text = "2. 정압 상사법칙 (Static Pressure Scaling):"
    p.font.name = "Malgun Gothic"
    p.font.size = Pt(12)
    p.font.bold = True
    p.space_before = Pt(15)

    p = tf_b1.add_paragraph()
    p.text = "   P2 = P1 * (N2 / N1)² * (D2 / D1)² * (ρ2 / ρ1)"
    p.font.name = "Courier New"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = blue

    p = tf_b1.add_paragraph()
    p.text = "   - 동압과 정압은 블레이드 끝단의 회전 속도(팁 속도)의 제곱에 비례하며, 작동 유체의 기체 밀도(ρ) 비율에 따라 정비례로 가중 보정됩니다."
    p.font.name = "Malgun Gothic"
    p.font.size = Pt(10)
    p.font.color.rgb = dark_gray
    p.space_before = Pt(2)

    box2 = slide3.shapes.add_textbox(Inches(6.9), Inches(1.8), box_width, box_height)
    tf_b2 = box2.text_frame
    tf_b2.word_wrap = True

    p = tf_b2.paragraphs[0]
    p.text = "■ 축동력 및 소음 예측 공식"
    p.font.name = "Malgun Gothic"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = orange

    p = tf_b2.add_paragraph()
    p.text = "3. 축동력 상사법칙 (Shaft Power Scaling):"
    p.font.name = "Malgun Gothic"
    p.font.size = Pt(12)
    p.font.bold = True
    p.space_before = Pt(10)

    p = tf_b2.add_paragraph()
    p.text = (
        "   W2 = W1 * (N2 / N1)³ * (D2 / D1)⁵ * (ρ2 / ρ1) * (η1 / η2)"
        if is_axial_category(fan_category)
        else "   W2 = W1 * (N2 / N1)³ * (D2 / D1)⁴ * (H2 / H1) * (ρ2 / ρ1) * (η1 / η2)"
    )
    p.font.name = "Courier New"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = blue

    p = tf_b2.add_paragraph()
    p.text = "   - 공기역학적 운동 에너지는 속도의 3승에 비례하며, 효율 보정(Ackeret 효율 변화율 η1/η2)을 반영하여 스케일 증가에 따른 점성 손실 개선분을 정밀하게 산출합니다."
    p.font.name = "Malgun Gothic"
    p.font.size = Pt(10)
    p.font.color.rgb = dark_gray
    p.space_before = Pt(2)

    p = tf_b2.add_paragraph()
    p.text = "4. 소음 레벨 상사법칙 (Sound Level Scaling):"
    p.font.name = "Malgun Gothic"
    p.font.size = Pt(12)
    p.font.bold = True
    p.space_before = Pt(15)

    p = tf_b2.add_paragraph()
    p.text = (
        "   Lw2 = Lw1 + 50 log10(N2/N1) + 70 log10(D2/D1) + 20 log10(ρ2/ρ1)"
        if is_axial_category(fan_category)
        else "   Lw2 = Lw1 + 50 log10(N2/N1) + 60 log10(D2/D1) + 10 log10(H2/H1) + 20 log10(ρ2/ρ1)"
    )
    p.font.name = "Courier New"
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = blue

    p = tf_b2.add_paragraph()
    p.text = "   - 소음 강도는 속도의 5승, 크기의 6승에 가까운 스케일링 경향을 보이며, 난류 소음원 강도 모델(Turbulent Noise Model)을 기반으로 물리 보정 계수가 중첩 적용됩니다."
    p.font.name = "Malgun Gothic"
    p.font.size = Pt(10)
    p.font.color.rgb = dark_gray
    p.space_before = Pt(2)

    # ----------------------------------------------------
    # SLIDE 4: Physical Commentary & Recommendations
    # ----------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    title_box4 = slide4.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(11.9), Inches(1.0))
    tf4 = title_box4.text_frame
    p_t4 = tf4.paragraphs[0]
    p_t4.text = "물리 현상 해석 및 기술 검토"
    p_t4.font.name = "Malgun Gothic"
    p_t4.font.size = Pt(24)
    p_t4.font.bold = True
    p_t4.font.color.rgb = dark_gray

    p_st4 = tf4.add_paragraph()
    p_st4.text = f"수행된 예측 조건에 대한 공학적 해석 코멘트"
    p_st4.font.name = "Malgun Gothic"
    p_st4.font.size = Pt(12)
    p_st4.font.color.rgb = text_muted
    p_st4.space_before = Pt(5)

    comm_box = slide4.shapes.add_textbox(Inches(0.7), Inches(1.8), Inches(11.9), Inches(4.5))
    tf_c = comm_box.text_frame
    tf_c.word_wrap = True

    p = tf_c.paragraphs[0]
    p.text = "■ 설계 타당성 공학 해설 (Engineering Commentary)"
    p.font.name = "Malgun Gothic"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = orange

    for idx, comment in enumerate(commentary_list):
        p = tf_c.add_paragraph()
        p.text = comment
        p.font.name = "Malgun Gothic"
        p.font.size = Pt(12)
        p.font.color.rgb = dark_gray
        p.space_before = Pt(12)

    p_eff_note = tf_c.add_paragraph()
    if use_eff_correction:
        p_eff_note.text = f"• 효율 보정 옵션 적용: Ackeret 공식을 통해 스케일 효과에 의한 레이놀즈수 차이를 보정하였습니다. 직경이 {base_D:.0f}mm에서 {target_D:.0f}mm로 변화함에 따라 고정 스케일 가정을 넘어선 동력 보정이 완수되었습니다."
    else:
        p_eff_note.text = "• 효율 보정 옵션 미적용: 형상 치수와 무관하게 시스템 효율이 1:1로 고정 상사된다고 단순 가정하고 산출한 스펙입니다."
    p_eff_note.font.name = "Malgun Gothic"
    p_eff_note.font.size = Pt(11)
    p_eff_note.font.italic = True
    p_eff_note.font.color.rgb = blue
    p_eff_note.space_before = Pt(20)

    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


# -----------------------------------------------------------------------------
# 1. PAGE CONFIGURATION & THEME CUSTOMIZATION
# -----------------------------------------------------------------------------
def render_theme():
    st.set_page_config(
        page_title="Fan Studio Mobile",
        page_icon="💨",
        layout="wide",
        initial_sidebar_state="collapsed",
    )

    st.markdown(
        """
    <style>
    @import url('https://cdn.jsdelivr.net/gh/orioncactus/pretendard@v1.3.9/dist/web/variable/pretendardvariable-dynamic-subset.min.css');
    :root {
      --ink:#0F172A; --ink2:#334155; --muted:#64748B; --line:#E5E9F0; --bg:#F4F6F9; --card:#FFFFFF;
      --accent:#0D9488; --accent-strong:#0F766E; --accent-soft:#E6F6F4; --navy:#0F1B2D;
      --font:'Pretendard Variable',Pretendard,-apple-system,BlinkMacSystemFont,'Apple SD Gothic Neo','Noto Sans KR','Malgun Gothic',sans-serif;
      --radius:16px;
    }
    html, body { -webkit-text-size-adjust:100%; }
    .stApp { background:var(--bg); color:var(--ink); font-family:var(--font); }
    .stApp p, .stApp label, .stApp input, .stApp textarea, .stApp button, .stApp li, .stApp td, .stApp th,
    .stApp [data-testid="stMarkdownContainer"], .stApp [data-testid="stCaptionContainer"], .stApp strong,
    .stApp [data-testid="stMetricValue"] *, [data-baseweb="select"], [data-baseweb="popover"] * { font-family:var(--font) !important; }
    .stApp p, .stApp li { line-height:1.6; }
    /* 상단 크롬 숨김 · 화면 여백 */
    header[data-testid="stHeader"] { display:none; }
    [data-testid="stSidebar"], [data-testid="stSidebarCollapsedControl"] { display:none; }
    .block-container { max-width:880px; padding:12px 14px 72px !important; }
    h1,h2,h3,h4,h5,h6 { font-family:var(--font) !important; color:var(--ink) !important; letter-spacing:-.02em; }
    h1 { font-weight:800 !important; font-size:1.45rem !important; }
    h2 { font-size:1.2rem !important; } h3 { font-size:1.08rem !important; font-weight:700 !important; }
    h4 { font-size:.98rem !important; font-weight:700 !important; } h5, h6 { font-size:.92rem !important; }
    hr { margin:.9rem 0 !important; border-color:var(--line) !important; }
    /* 헤더 */
    .m-hero { display:flex; align-items:center; gap:12px; padding:16px 18px; margin:2px 0 12px; border-radius:var(--radius);
              background:linear-gradient(135deg,#0F1B2D 0%,#16314F 100%); color:#fff; }
    .m-hero .logo { width:42px; height:42px; flex:0 0 42px; border-radius:12px; background:rgba(94,234,212,.14);
                    display:flex; align-items:center; justify-content:center; font-size:1.35rem; }
    .m-hero .eyebrow { font-size:.66rem; color:#5EEAD4; letter-spacing:.14em; font-weight:700; }
    .m-hero h1 { color:#F8FAFC !important; font-size:1.18rem !important; margin:2px 0 0 !important; padding:0 !important; line-height:1.3; }
    .m-hero p { color:#B8C4D6; margin:2px 0 0; font-size:.78rem; line-height:1.4; }
    .m-chips { display:flex; gap:6px; flex-wrap:wrap; margin:0 0 10px; }
    .m-chips span { background:var(--card); border:1px solid var(--line); border-radius:999px; padding:5px 11px;
                    font-size:.74rem; color:var(--ink2); white-space:nowrap; font-variant-numeric:tabular-nums; }
    .m-chips span b { color:var(--accent-strong); font-weight:700; }
    /* 입력 위젯: 터치 친화 크기, iOS 확대 방지(16px) */
    [data-testid="stWidgetLabel"] p { font-size:.82rem !important; font-weight:600; color:var(--ink2); }
    [data-testid="stNumberInput"] input, [data-testid="stTextInput"] input, [data-testid="stTextArea"] textarea,
    [data-baseweb="select"] > div { font-size:16px !important; font-variant-numeric:tabular-nums; }
    [data-testid="stNumberInput"] input, [data-testid="stTextInput"] input { min-height:44px; }
    [data-baseweb="select"] > div { min-height:44px; border-radius:10px !important; }
    [data-baseweb="input"], [data-baseweb="base-input"] { border-radius:10px !important; }
    [data-testid="stNumberInput"] button { min-width:40px; }
    [data-baseweb="checkbox"]:has(input:checked)>span:first-child { background-color:var(--accent) !important; border-color:var(--accent) !important; }
    [data-testid="stRadio"] label, [data-testid="stCheckbox"] label { min-height:32px; align-items:center; }
    /* 카드형 컨테이너 */
    [data-testid="stVerticalBlockBorderWrapper"] { background:var(--card); border-color:var(--line) !important; border-radius:var(--radius) !important; }
    [class*="st-key-panel_"] { background:var(--card); border:1px solid var(--line) !important; border-radius:var(--radius) !important; padding:14px !important; }
    [data-testid="stExpander"] { background:var(--card); border-radius:14px; }
    [data-testid="stExpander"] details { border-color:var(--line) !important; border-radius:14px !important; }
    [data-testid="stExpander"] summary { min-height:48px; padding:10px 14px !important; }
    [data-testid="stExpander"] summary p { font-size:.92rem !important; font-weight:650; color:var(--ink); }
    .st-key-m_settings [data-testid="stExpander"] details { border:1.5px solid #CDEDEA !important; background:#FBFEFE; }
    /* 섹션 탭: 가로 스크롤 알약형 + 상단 고정 */
    .stTabs [data-baseweb="tab-list"] { gap:6px; overflow-x:auto; scrollbar-width:none; -webkit-overflow-scrolling:touch;
        position:sticky; top:0; z-index:50; background:var(--bg); padding:8px 0 10px; border-bottom:none; }
    .stTabs [data-baseweb="tab-list"]::-webkit-scrollbar { display:none; }
    .stTabs [data-baseweb="tab"] { flex:0 0 auto; height:38px; padding:0 14px; border-radius:999px; background:var(--card);
        border:1px solid var(--line); color:var(--muted); font-size:.86rem; font-weight:650; white-space:nowrap; }
    .stTabs [data-baseweb="tab"] p { font-size:.86rem !important; }
    .stTabs [aria-selected="true"] { background:var(--navy) !important; border-color:var(--navy) !important; color:#fff !important; }
    .stTabs [aria-selected="true"] p { color:#fff !important; }
    .stTabs [data-baseweb="tab-highlight"], .stTabs [data-baseweb="tab-border"] { display:none; }
    .stTabs [data-baseweb="tab-panel"] { padding-top:6px; }
    /* 설정 패널 안의 탭은 밑줄형 */
    .st-key-m_settings .stTabs [data-baseweb="tab-list"] { position:static; background:transparent; padding:0 0 6px; }
    .st-key-m_settings .stTabs [data-baseweb="tab"] { height:34px; font-size:.82rem; }
    /* 지표 카드 */
    .metric-box { background:var(--card) !important; border:1px solid var(--line) !important; border-radius:14px !important;
                  padding:14px 16px !important; text-align:left !important; min-height:0; box-sizing:border-box; margin-bottom:8px; }
    .metric-box .title { color:var(--muted); font-size:.76rem; font-weight:600; }
    .metric-box .value { font-family:var(--font) !important; font-size:1.55rem; font-weight:800; color:var(--ink); margin:4px 0 2px;
                         font-variant-numeric:tabular-nums; letter-spacing:-.03em; overflow-wrap:anywhere; line-height:1.25; }
    .metric-box .delta { font-size:.76rem; font-weight:500; color:var(--muted); }
    .delta-positive { color:var(--accent-strong) !important; } .delta-negative { color:#B45309 !important; } .delta-neutral { color:var(--muted) !important; }
    [data-testid="stMetric"] { background:var(--card); border:1px solid var(--line); border-radius:14px; padding:12px 14px; }
    [data-testid="stMetricLabel"] p { font-size:.76rem !important; color:var(--muted); }
    [data-testid="stMetricValue"] { font-variant-numeric:tabular-nums; font-size:1.4rem; font-weight:800; }
    /* 차트 · 표 */
    [data-testid="stPlotlyChart"] { background:var(--card); border:1px solid var(--line); border-radius:var(--radius); padding:4px; overflow:hidden; }
    .m-chart-title { font-size:.95rem; font-weight:750; color:var(--ink); margin:10px 2px -4px; line-height:1.4; letter-spacing:-.01em; }
    [data-testid="stDataFrame"], [data-testid="stDataEditor"] { border-radius:12px; overflow:hidden; }
    /* 버튼: 전체 너비 · 44px 이상 */
    div.stButton>button, div.stDownloadButton>button, [data-testid="stFormSubmitButton"]>button {
        width:100%; min-height:46px; border-radius:12px; font-weight:700; font-size:.92rem; }
    div.stButton>button[kind="primary"], div.stDownloadButton>button[kind="primary"] { background:var(--accent-strong); border-color:var(--accent-strong); }
    [data-testid="stFileUploaderDropzone"] { border-radius:12px; padding:14px; }
    [data-testid="stDownloadButton"], [data-testid="stButton"] { width:100%; }
    /* 알림·캡션 */
    [data-testid="stAlert"] { border-radius:12px; }
    [data-testid="stAlert"] p { font-size:.86rem; }
    [data-testid="stCaptionContainer"], [data-testid="stCaptionContainer"] p { color:var(--muted) !important; font-size:.76rem; line-height:1.55; }
    /* 세트 요약 카드 */
    .set-chart-summary { display:grid; grid-template-columns:repeat(auto-fit,minmax(150px,1fr)); gap:8px; margin:4px 0 10px; }
    .set-chart-summary>div { padding:12px 14px; background:var(--card); border:1px solid var(--line); border-radius:14px; }
    .set-chart-summary span { display:block; color:var(--muted); font-size:.72rem; margin-bottom:4px; }
    .set-chart-summary strong { color:var(--ink2); font-size:1.2rem; font-weight:700; font-variant-numeric:tabular-nums; }
    .set-chart-summary b { color:var(--accent); font-weight:800; }
    .set-chart-summary em { color:#CBD5E1; font-style:normal; padding:0 4px; }
    .set-chart-summary small { color:var(--muted); font-size:.72rem; font-weight:500; }
    /* 좁은 화면: 열 간격 축소 */
    @media (max-width:640px) {
      [data-testid="stHorizontalBlock"] { gap:.5rem !important; }
      .block-container { padding:8px 12px 72px !important; }
      .metric-box .value { font-size:1.4rem; }
      [data-testid="stPlotlyChart"] { padding:0; border-radius:14px; }
    }
    </style>
    """,
        unsafe_allow_html=True,
    )


# -----------------------------------------------------------------------------
# 2. UNIT CONVERSIONS & GENERAL DATA DICTIONARY
# -----------------------------------------------------------------------------
# Standard internal units: Flow = CMM (m3/min), Pressure = Pa, Power = W, Noise = dB(A)
UNIT_FLOW = {"CMM (m³/min)": 1.0, "CMH (m³/hr)": 60.0, "CFM (ft³/min)": 35.3147}

UNIT_PRESSURE = {"Pa": 1.0, "mmAq": 1.0 / 9.80665, "kPa": 0.001}

UNIT_POWER = {"W": 1.0, "kW": 0.001, "HP": 1.0 / 745.699}


def create_set_prediction_powerpoint_report(
    fan_category,
    dim_label,
    target_q,
    current_rpm,
    new_rpm,
    current_p,
    new_p,
    current_power,
    new_power,
    current_noise,
    new_noise,
    current_d,
    current_x,
    new_d,
    new_x,
    current_k,
    system_exponent,
    fit_rmse_pct,
    current_q_curve,
    current_p_curve,
    current_system_curve,
    new_q_curve,
    new_p_curve,
    new_system_curve,
    current_noise_curve,
    new_noise_curve,
    current_power_source,
    new_power_source,
    current_noise_source,
    new_noise_source,
    component_note="변경 부품 없음",
):
    """세트 예측 결과를 편집 가능한 데이터 내장형 차트가 포함된 PPTX로 만든다."""

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    navy = RGBColor(28, 39, 54)
    orange = RGBColor(13, 148, 136)
    teal = RGBColor(70, 133, 146)
    warm_bg = RGBColor(247, 246, 243)
    white = RGBColor(255, 255, 255)
    soft_gray = RGBColor(226, 225, 221)
    text_gray = RGBColor(98, 101, 105)
    green = RGBColor(54, 128, 103)

    background = slide.background
    background.fill.solid()
    background.fill.fore_color.rgb = warm_bg

    def add_box(x, y, w, h, color=white, radius=True, line_color=None):
        shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.color.rgb = line_color or color
        return shape

    def add_text(
        x,
        y,
        w,
        h,
        text,
        size=10,
        color=navy,
        bold=False,
        align=PP_ALIGN.LEFT,
        font_name="Malgun Gothic",
        valign=MSO_ANCHOR.TOP,
    ):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.vertical_anchor = valign
        tf.margin_left = Inches(0.03)
        tf.margin_right = Inches(0.03)
        tf.margin_top = Inches(0.02)
        tf.margin_bottom = Inches(0.02)
        paragraph = tf.paragraphs[0]
        paragraph.text = text
        paragraph.alignment = align
        paragraph.font.name = font_name
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color
        return box

    def fmt(value, digits=1, suffix=""):
        return "—" if value is None or not np.isfinite(value) else f"{value:.{digits}f}{suffix}"

    header = add_box(0, 0, 13.333, 0.82, navy, radius=False)
    add_text(0.48, 0.16, 8.6, 0.35, "FAN SYSTEM PERFORMANCE REPORT", 19, white, True)
    add_text(
        0.50, 0.51, 8.8, 0.19, "세트 실측 보정 · 목표 풍량 자동 RPM · 교체 예측", 8.5, RGBColor(198, 207, 216)
    )
    add_text(
        10.0, 0.23, 2.83, 0.28, datetime.date.today().strftime("%Y.%m.%d"), 10, white, True, PP_ALIGN.RIGHT
    )

    def add_operating_card(x, title, color, rpm, pressure, power, noise, diameter, x_dim):
        add_box(x, 1.04, 2.88, 1.68, white, True, soft_gray)
        add_box(x, 1.04, 0.09, 1.68, color, radius=False)
        add_text(x + 0.22, 1.16, 2.35, 0.25, title, 11, color, True)
        add_text(x + 0.22, 1.48, 2.38, 0.42, f"{fmt(rpm, 0)} RPM", 22, navy, True)
        add_text(
            x + 0.22,
            1.94,
            2.43,
            0.26,
            f"Q {target_q:.2f} CMM   ·   P {fmt(pressure, 1)} Pa",
            8.5,
            text_gray,
            True,
        )
        add_text(
            x + 0.22,
            2.23,
            2.43,
            0.24,
            f"소전 {fmt(power, 1)} W   ·   소음 {fmt(noise, 1)} dB(A)",
            8.5,
            text_gray,
        )
        geometry = (
            f"D {diameter:.0f} mm"
            if is_axial_category(fan_category)
            else f"D {diameter:.0f} / {dim_label} {x_dim:.0f} mm"
        )
        add_text(x + 0.22, 2.49, 2.43, 0.18, geometry, 7.5, text_gray)

    add_operating_card(
        0.48,
        "CURRENT · 현재",
        teal,
        current_rpm,
        current_p,
        current_power,
        current_noise,
        current_d,
        current_x,
    )
    add_operating_card(3.48, "NEW · 신규", orange, new_rpm, new_p, new_power, new_noise, new_d, new_x)

    rpm_delta = (new_rpm / current_rpm - 1.0) * 100.0 if current_rpm else np.nan
    power_delta = (
        (new_power / current_power - 1.0) * 100.0 if current_power and new_power is not None else np.nan
    )
    noise_delta = new_noise - current_noise if current_noise is not None and new_noise is not None else np.nan
    add_box(0.48, 2.88, 5.88, 0.66, navy, True, navy)
    add_text(0.72, 3.01, 1.65, 0.18, "CHANGE SNAPSHOT", 8, RGBColor(185, 195, 204), True)
    add_text(
        2.05,
        2.99,
        4.02,
        0.24,
        f"RPM {rpm_delta:+.1f}%     소전 {power_delta:+.1f}%     소음 {noise_delta:+.1f} dB",
        10.5,
        white,
        True,
        PP_ALIGN.RIGHT,
    )

    add_box(0.48, 3.72, 5.88, 2.62, white, True, soft_gray)
    add_text(0.72, 3.91, 5.30, 0.25, "CALCULATION LOGIC · 적용 수식", 10, navy, True)
    if is_axial_category(fan_category):
        formula_lines = [
            "Q₂ = Q₁·(N₂/N₁)·(D₂/D₁)³",
            "ΔP₂ = ΔP₁·(N₂/N₁)²·(D₂/D₁)²·(ρ₂/ρ₁)",
            "W₂ = W₁·(N₂/N₁)³·(D₂/D₁)⁵·(ρ₂/ρ₁)",
            "L₂ = L₁ + 50log(N₂/N₁) + 70log(D₂/D₁) + 20log(ρ₂/ρ₁)",
        ]
    else:
        formula_lines = [
            "Q₂ = Q₁·(N₂/N₁)·(D₂/D₁)²·(X₂/X₁)",
            "ΔP₂ = ΔP₁·(N₂/N₁)²·(D₂/D₁)²·(ρ₂/ρ₁)",
            "W₂ = W₁·(N₂/N₁)³·(D₂/D₁)⁴·(X₂/X₁)·(ρ₂/ρ₁)",
            "L₂ = L₁ + 50log(N₂/N₁) + 60log(D₂/D₁) + 10log(X₂/X₁) + 20log(ρ₂/ρ₁)",
        ]
    formula_lines += [
        f"세트 저항: ΔP = {current_k:.5g}·Q^{system_exponent:.3f}  (적합 RMSE {fit_rmse_pct:.2f}%)",
        "부품 기준점식: ΔP₂ = ΔP_ref·[(Q/Q_ref)·(A₁/A₂)]ᵐ",
        "면풍속식: ΔP = a₁v+a₂v², v=(Q/60)/A  /  K식: ΔP=Kρv²/2",
        "자동 RPM: P_fan(Q목표, N) = P_set(Q목표)가 되는 N을 수치해석",
        "소비전력: P_in = W_shaft/(η_motor·η_ctrl)·C_power + P_aux",
        "Lw 미입력 시: L_set(N)=a+b·logN 회귀 + 팬 종류별 형상·밀도 상사",
        "raw 소음 회귀 불가 시에만 KN + 10logQ_fan + 20logP_fan 보조식 사용",
    ]
    formula_text = "\n".join(formula_lines)
    formula_box = add_text(0.73, 4.20, 5.27, 1.53, formula_text, 7.15, navy, False, font_name="Malgun Gothic")
    for paragraph in formula_box.text_frame.paragraphs:
        paragraph.space_after = Pt(2)
    add_text(0.73, 5.74, 5.24, 0.44, f"전력: {new_power_source}\n소음: {new_noise_source}", 7.4, text_gray)

    add_box(6.62, 1.04, 6.22, 5.30, white, True, soft_gray)
    add_text(6.88, 1.20, 5.72, 0.24, "EDITABLE PERFORMANCE CHARTS · 데이터 내장형 차트", 10, navy, True)

    def add_native_xy_chart(x, y, w, h, title, x_title, y_title, series_specs):
        chart_data = XyChartData()
        valid_specs = []
        for spec in series_specs:
            x_values = np.asarray(spec["x"], dtype=float)
            y_values = np.asarray(spec["y"], dtype=float)
            valid = np.isfinite(x_values) & np.isfinite(y_values)
            if not np.any(valid):
                continue
            data_series = chart_data.add_series(spec["name"])
            for x_value, y_value in zip(x_values[valid], y_values[valid]):
                data_series.add_data_point(float(x_value), float(y_value))
            valid_specs.append(spec)

        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), chart_data
        ).chart
        chart.has_title = True
        chart.chart_title.text_frame.text = title
        chart.chart_title.text_frame.paragraphs[0].font.name = "Malgun Gothic"
        chart.chart_title.text_frame.paragraphs[0].font.size = Pt(9)
        chart.chart_title.text_frame.paragraphs[0].font.bold = True
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.name = "Malgun Gothic"
        chart.legend.font.size = Pt(6.5)
        chart.category_axis.has_title = True
        chart.category_axis.axis_title.text_frame.text = x_title
        chart.value_axis.has_title = True
        chart.value_axis.axis_title.text_frame.text = y_title
        chart.category_axis.tick_labels.font.name = "Malgun Gothic"
        chart.category_axis.tick_labels.font.size = Pt(6.5)
        chart.value_axis.tick_labels.font.name = "Malgun Gothic"
        chart.value_axis.tick_labels.font.size = Pt(6.5)
        chart.value_axis.has_major_gridlines = True
        chart.value_axis.major_gridlines.format.line.color.rgb = soft_gray
        for series, spec in zip(chart.series, valid_specs):
            color = spec["color"]
            if spec.get("marker", False):
                series.format.line.fill.background()
                series.marker.style = XL_MARKER_STYLE.DIAMOND
                series.marker.size = 8
                series.marker.format.fill.solid()
                series.marker.format.fill.fore_color.rgb = color
                series.marker.format.line.color.rgb = color
            else:
                series.format.line.color.rgb = color
                series.format.line.width = Pt(spec.get("width", 1.8))
                if spec.get("dash", False):
                    series.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        return chart

    pq_specs = [
        {
            "name": f"현재 팬 {current_rpm:.0f} RPM",
            "x": current_q_curve,
            "y": current_p_curve,
            "color": teal,
            "width": 2.2,
        },
        {
            "name": "현재 세트 저항",
            "x": current_q_curve,
            "y": current_system_curve,
            "color": teal,
            "width": 1.3,
            "dash": True,
        },
        {
            "name": f"신규 팬 {new_rpm:.0f} RPM",
            "x": new_q_curve,
            "y": new_p_curve,
            "color": orange,
            "width": 2.2,
        },
        {
            "name": "변경 세트 저항",
            "x": new_q_curve,
            "y": new_system_curve,
            "color": orange,
            "width": 1.3,
            "dash": True,
        },
        {"name": "현재 운전점", "x": [target_q], "y": [current_p], "color": navy, "marker": True},
        {"name": "신규 운전점", "x": [target_q], "y": [new_p], "color": orange, "marker": True},
    ]
    add_native_xy_chart(
        6.83, 1.48, 5.74, 2.95, "PQ·세트 저항 및 운전점", "세트 풍량 Q (CMM)", "정압 ΔP (Pa)", pq_specs
    )

    noise_specs = [
        {
            "name": "현재 세트 소음",
            "x": current_q_curve,
            "y": current_noise_curve,
            "color": teal,
            "width": 2.0,
        },
        {"name": "신규 세트 소음", "x": new_q_curve, "y": new_noise_curve, "color": orange, "width": 2.0},
    ]
    if current_noise is not None:
        noise_specs.append(
            {"name": "현재 목표 소음", "x": [target_q], "y": [current_noise], "color": navy, "marker": True}
        )
    if new_noise is not None:
        noise_specs.append(
            {"name": "신규 목표 소음", "x": [target_q], "y": [new_noise], "color": orange, "marker": True}
        )
    add_native_xy_chart(
        6.83, 4.48, 5.74, 1.60, "세트 소음 예측", "세트 풍량 Q (CMM)", "소음 dB(A)", noise_specs
    )

    add_box(0.48, 6.55, 12.36, 0.56, RGBColor(236, 238, 236), True, RGBColor(236, 238, 236))
    add_text(0.72, 6.69, 1.12, 0.18, "ENGINEERING NOTE", 7.6, green, True)
    note = (
        f"{component_note}  ·  직경/형상 ±20% 또는 raw RPM 범위 밖에서는 단품시험·CFD 확인 권장  ·  "
        "KN 소음식은 광대역 경향을 예측하며 공진·실속·BPF 톤소음은 별도 검증"
    )
    add_text(1.94, 6.66, 10.52, 0.25, note, 7.2, text_gray)
    add_text(
        0.50, 7.24, 12.30, 0.12, "Generated by Fan Similarity v10", 6.5, text_gray, False, PP_ALIGN.RIGHT
    )

    output = io.BytesIO()
    prs.save(output)
    output.seek(0)
    return output.getvalue()


def unit_number_input(
    label,
    *,
    quantity,
    unit,
    value=0.0,
    min_value=None,
    max_value=None,
    step=None,
    key=None,
    value_is_display=False,
    **kwargs,
):
    """스칼라 입력의 원본값을 CMM/Pa/W로 유지하고 표시 단위만 변경한다."""
    factors = {"flow": UNIT_FLOW, "pressure": UNIT_PRESSURE, "power": UNIT_POWER}[quantity]
    factor = factors[unit]
    if key is None:
        stem = label.rsplit("(", 1)[0].strip()
        key = "unit_" + hashlib.sha256((quantity + stem).encode()).hexdigest()[:16]
    records = st.session_state.setdefault("_unit_input_records", {})
    previous = records.get(key)
    if previous is None:
        canonical = float(value) / factor if value_is_display else float(value)
    else:
        canonical = previous["canonical"]
        if key in st.session_state:
            canonical = float(st.session_state[key]) / previous["factor"]
    display = canonical * factor
    if previous is not None and previous["factor"] != factor:
        st.session_state[key] = display
    parameters = dict(value=display, key=key, **kwargs)
    if previous is not None and previous["factor"] != factor:
        # Streamlit은 state와 명시적 기본값을 함께 지정하면 경고하므로 변경 시 기본값은 생략한다.
        parameters["value"] = "min"
    for name, val in [("min_value", min_value), ("max_value", max_value), ("step", step)]:
        if val is not None:
            parameters[name] = float(val) * factor
    parameters.setdefault("format", "%.4f" if factor < 0.01 else "%.2f")
    result = st.number_input(label, **parameters)
    records[key] = dict(canonical=float(result) / factor, factor=factor)
    return result


def pq_display_frame(canonical, flow_unit, pressure_unit, power_unit):
    return pd.DataFrame(
        {
            f"풍량 ({flow_unit.split(' ')[0]})": canonical["q"] * UNIT_FLOW[flow_unit],
            f"정압 ({pressure_unit})": canonical["p"] * UNIT_PRESSURE[pressure_unit],
            f"동력 ({power_unit})": canonical["w"] * UNIT_POWER[power_unit],
            "소음 (dB(A))": canonical["noise"],
        }
    )


def pq_canonical_frame(display, flow_unit, pressure_unit, power_unit):
    clean = _normalize_editor_dataframe(display, display.columns)
    return pd.DataFrame(
        {
            "q": clean.iloc[:, 0] / UNIT_FLOW[flow_unit],
            "p": clean.iloc[:, 1] / UNIT_PRESSURE[pressure_unit],
            "w": clean.iloc[:, 2] / UNIT_POWER[power_unit],
            "noise": clean.iloc[:, 3],
        }
    )


def render_pq_editor(flow_unit, pressure_unit, power_unit):
    """입력 파일/열/원본 단위 변경 시에만 원본을 교체한다. 단위 전환은 편집값을 보존한다."""
    canonical = pd.DataFrame(
        dict(
            q=[0.0, 10.0, 20.0, 30.0, 40.0, 50.0, 60.0, 70.0, 80.0],
            p=[25.0, 24.5, 23.5, 22.0, 19.5, 16.0, 11.5, 6.0, 0.0],
            w=[100.0, 120.0, 135.0, 145.0, 150.0, 145.0, 130.0, 105.0, 75.0],
            noise=[62.0, 60.0, 59.0, 58.0, 59.0, 61.0, 63.0, 66.0, 70.0],
        )
    )
    source = "example"
    uploaded = st.file_uploader(
        "측정 데이터 가져오기",
        type=["xlsx", "xls", "csv"],
        key="pq_upload",
        help="열과 원본 단위를 확인한 뒤 입력합니다. 표를 직접 편집할 수도 있습니다.",
    )
    if uploaded is not None:
        try:
            content = uploaded.getvalue()
            frame = read_uploaded_table(content, uploaded.name)
            if len(frame.columns) < 3:
                raise ValueError(
                    "단품 탭에는 풍량·정압·동력 열이 필요합니다. Q/P만 있다면 세트 예측 탭을 사용하십시오."
                )
            columns = list(frame.columns)

            def guess(words, fallback):
                return next(
                    (i for i, c in enumerate(columns) if any(w in str(c).lower() for w in words)), fallback
                )

            with st.expander("열 매핑 · 원본 단위", expanded=True):
                c1, c2 = st.columns(2)
                with c1:
                    qc = st.selectbox(
                        "풍량 열",
                        columns,
                        index=guess(["풍량", "flow", "cmm", "cmh", "cfm"], 0),
                        key="pq_qcol",
                    )
                    pc = st.selectbox(
                        "정압 열", columns, index=guess(["정압", "pressure", "static"], 1), key="pq_pcol"
                    )
                    wc = st.selectbox(
                        "동력 열", columns, index=guess(["동력", "power", "watt"], 2), key="pq_wcol"
                    )
                    noise_options = [None] + columns
                    nc = st.selectbox(
                        "소음 열 (선택)",
                        noise_options,
                        index=next(
                            (
                                i + 1
                                for i, c in enumerate(columns)
                                if any(w in str(c).lower() for w in ["소음", "noise", "lw", "db"])
                            ),
                            0,
                        ),
                        format_func=lambda v: "없음" if v is None else str(v),
                        key="pq_ncol",
                    )

                def unit_index(column, quantity):
                    name = str(column).lower()
                    if quantity == "flow":
                        return 1 if "cmh" in name or "m³/hr" in name else 2 if "cfm" in name else 0
                    if quantity == "pressure":
                        return 1 if "mmaq" in name else 2 if "kpa" in name else 0
                    return 1 if "kw" in name else 2 if "hp" in name else 0

                with c2:
                    fu = st.selectbox(
                        "원본 풍량 단위", list(UNIT_FLOW), index=unit_index(qc, "flow"), key="pq_import_fu"
                    )
                    pu = st.selectbox(
                        "원본 정압 단위",
                        list(UNIT_PRESSURE),
                        index=unit_index(pc, "pressure"),
                        key="pq_import_pu",
                    )
                    wu = st.selectbox(
                        "원본 동력 단위", list(UNIT_POWER), index=unit_index(wc, "power"), key="pq_import_wu"
                    )
            selected = [qc, pc, wc] + ([nc] if nc is not None else [])
            if len(set(selected)) != len(selected):
                raise ValueError("풍량·정압·동력·소음은 서로 다른 열로 지정하십시오.")
            canonical = pd.DataFrame(
                {
                    "q": pd.to_numeric(frame[qc], errors="coerce") / UNIT_FLOW[fu],
                    "p": pd.to_numeric(frame[pc], errors="coerce") / UNIT_PRESSURE[pu],
                    "w": pd.to_numeric(frame[wc], errors="coerce") / UNIT_POWER[wu],
                    "noise": pd.to_numeric(frame[nc], errors="coerce") if nc is not None else np.nan,
                }
            )
            source = hashlib.sha256(content + repr((qc, pc, wc, nc, fu, pu, wu)).encode()).hexdigest()
        except Exception as error:
            st.error(f"파일 입력 확인: {error}")
            return None
    else:
        st.caption("초기 표는 예시 데이터입니다. 실제 측정값으로 편집하거나 파일을 가져오십시오.")
    units = (flow_unit, pressure_unit, power_unit)
    state = st.session_state.get("_pq_state")
    if state is None or state["source"] != source:
        state = dict(source=source, units=None, canonical=canonical, revision=0)
    if state["units"] != units:
        state["seed"] = pq_display_frame(state["canonical"], *units)
        state["units"] = units
        state["revision"] += 1
    # seed는 같은 단위에서 고정해야 data_editor의 편집 패치가 중복 적용되지 않는다.
    editor_key = f"pq_editor_{source[:12]}_{state['revision']}"
    edited = st.data_editor(
        state["seed"],
        num_rows="dynamic",
        width="stretch",
        hide_index=True,
        height=330,
        key=editor_key,
        column_config=_numeric_editor_column_config(state["seed"].columns),
    )
    state["canonical"] = pq_canonical_frame(edited, *units)
    st.session_state["_pq_state"] = state
    template = pq_display_frame(canonical.iloc[:5], *units)
    st.download_button(
        "입력 템플릿 · Excel",
        excel_bytes(template, "PQ_Data"),
        file_name="fan_pq_template_v10.xlsx",
        key="pq_template",
        on_click="ignore",
    )
    st.caption("편집값은 단위를 바꿔도 보존됩니다. 소음이 없으면 빈칸으로 두십시오.")
    return edited


def _switch_project():
    for key in list(st.session_state):
        if key.startswith("set2_"):
            del st.session_state[key]


def _show_chart(fig, *, layout_overrides=None, **kwargs):
    """공통 차트 디자인, 원본 좌표 툴팁, 고해상도 내보내기와 데이터 보기를 제공한다."""
    font = "Pretendard Variable, Pretendard, Apple SD Gothic Neo, Noto Sans KR, Malgun Gothic, sans-serif"
    chart_number = st.session_state.get("_chart_counter", 0) + 1
    st.session_state["_chart_counter"] = chart_number
    chart_key = kwargs.pop("key", f"v10_chart_{chart_number}")
    surface = any(trace.type == "surface" for trace in fig.data)
    chart_title = fig.layout.title.text or ""
    has_right_axis = any(
        getattr(fig.layout[name], "overlaying", None) for name in fig.layout if name.startswith("yaxis")
    )
    fig.update_layout(
        template="plotly_white",
        font=dict(family=font, size=11, color="#334155"),
        paper_bgcolor="#FFFFFF",
        plot_bgcolor="#FFFFFF",
        # 모바일: 제목은 차트 위 HTML로 표시하여 툴바와 겹치지 않게 한다.
        title=dict(text=""),
        # 모바일: 좁은 여백, 터치 스크롤을 막지 않도록 드래그 비활성화, 탭 시 가장 가까운 점 표시
        margin=dict(l=50, r=52 if has_right_axis else 14, t=50, b=88),
        height=min(max(fig.layout.height or 400, 360), 460),
        dragmode=False,
        hovermode="closest",
        hoverlabel=dict(
            bgcolor="#0F172A", font=dict(color="#F8FAFC", family=font, size=12), bordercolor="#0F172A"
        ),
        legend=dict(
            orientation="h",
            x=0,
            y=-0.23,
            xanchor="left",
            yanchor="top",
            font=dict(size=10.5),
            bgcolor="rgba(0,0,0,0)",
            borderwidth=0,
            tracegroupgap=8,
        ),
        modebar=dict(bgcolor="rgba(255,255,255,.8)", color="#64748B", activecolor="#0D9488"),
        uirevision=chart_key,
    )
    fig.update_xaxes(
        automargin=True,
        title_standoff=8,
        showgrid=True,
        gridcolor="#F1F5F9",
        linecolor="#CBD5E1",
        zeroline=False,
        ticks="outside",
        tickcolor="#CBD5E1",
        showspikes=True,
        spikemode="across",
        spikethickness=1,
    )
    fig.update_yaxes(
        automargin=True,
        title_standoff=6,
        gridcolor="#E9EEF4",
        linecolor="#CBD5E1",
        zeroline=False,
        ticks="outside",
        tickcolor="#CBD5E1",
        separatethousands=True,
    )
    data = []
    for trace in fig.data:
        if trace.type not in ("scatter", "scattergl"):
            continue
        n = len(trace.x) if trace.x is not None else 0
        if trace.mode and "lines" in trace.mode:
            trace.line.width = max(trace.line.width or 2, 2)
            trace.line.shape = "linear"
        if trace.mode and "markers" in trace.mode and n > 2:
            trace.marker.size = 5
        if trace.x is not None and trace.y is not None:
            xaxis = "xaxis" + (trace.xaxis[1:] if trace.xaxis and trace.xaxis != "x" else "")
            yaxis = "yaxis" + (trace.yaxis[1:] if trace.yaxis and trace.yaxis != "y" else "")
            xlabel = fig.layout[xaxis].title.text or "X"
            ylabel = fig.layout[yaxis].title.text or "Y"
            if not trace.hovertemplate:
                trace.hovertemplate = (
                    f"{xlabel}: %{{x:,.2f}}<br>{ylabel}: %{{y:,.2f}}<extra>%{{fullData.name}}</extra>"
                )
            data.extend(
                {"계열": trace.name or "Data", "X": x, "Y": y, "X 단위/축": xlabel, "Y 단위/축": ylabel}
                for x, y in zip(trace.x, trace.y)
            )
    if layout_overrides:
        fig.update_layout(**layout_overrides)
    if chart_title:
        st.markdown(f'<div class="m-chart-title">{chart_title}</div>', unsafe_allow_html=True)
    st.plotly_chart(
        fig,
        key=chart_key,
        config=dict(
            displaylogo=False,
            scrollZoom=False,
            responsive=True,
            modeBarButtonsToRemove=["select2d", "lasso2d", "pan2d", "zoom2d", "autoScale2d", "toggleSpikelines"],
            toImageButtonOptions=dict(format="png", filename="fan_analysis_mobile", scale=3),
        ),
        **kwargs,
    )
    with st.expander("📊 차트 데이터 · 내보내기", expanded=False):
        st.caption(
            "곡선을 탭하면 값이 표시됩니다. 범례를 탭하면 계열을 숨기고, 차트 우상단 +/− 로 확대·축소합니다. 카메라 버튼은 고해상도 PNG입니다."
        )
        if data:
            frame = pd.DataFrame(data)
            st.dataframe(frame, width="stretch", hide_index=True, height=180)
            st.download_button(
                "차트 데이터 CSV",
                frame.to_csv(index=False).encode("utf-8-sig"),
                file_name="fan_chart_mobile.csv",
                key=f"{chart_key}_csv",
                on_click="ignore",
            )


def mark_extrapolated_target(fig, q, p, target_q, target_p, label, color, *, row=1, col=1):
    if q[0] <= target_q <= q[-1]:
        return
    idx = 0 if target_q < q[0] else -1
    fig.add_trace(
        go.Scatter(
            x=[q[idx], target_q],
            y=[p[idx], target_p],
            mode="lines",
            name=f"{label} PQ 외삽",
            line=dict(color=color, width=2, dash="dash"),
        ),
        row=row,
        col=col,
    )
    fig.add_vrect(
        x0=min(q[idx], target_q),
        x1=max(q[idx], target_q),
        fillcolor=color,
        opacity=0.07,
        line_width=0,
        row=row,
        col=col,
    )


def convert_flow(val, from_unit, to_unit):
    # convert val from `from_unit` to CMM, then to `to_unit`
    val_cmm = val / UNIT_FLOW[from_unit]
    return val_cmm * UNIT_FLOW[to_unit]


def convert_pressure(val, from_unit, to_unit):
    # convert val from `from_unit` to Pa, then to `to_unit`
    val_pa = val / UNIT_PRESSURE[from_unit]
    return val_pa * UNIT_PRESSURE[to_unit]


def convert_power(val, from_unit, to_unit):
    # convert val from `from_unit` to W, then to `to_unit`
    val_w = val / UNIT_POWER[from_unit]
    return val_w * UNIT_POWER[to_unit]


def delta_arrow_class(value, higher_is_bad=False):
    """
    변화량 부호에 따라 (화살표, CSS 클래스)를 반환.
    higher_is_bad=False (풍량/정압): 증가=초록▲, 감소=빨강▼
    higher_is_bad=True  (동력/소음): 증가=빨강▲, 감소=초록▼
    """
    if value > 0:
        arrow = "▲"
        css = "delta-negative" if higher_is_bad else "delta-positive"
    elif value < 0:
        arrow = "▼"
        css = "delta-positive" if higher_is_bad else "delta-negative"
    else:
        arrow = "▬"
        css = "delta-neutral"
    return arrow, css


def safe_pct(new, base):
    """(new/base - 1)*100. base가 0이면 NaN 반환하여 ZeroDivisionError 방지."""
    try:
        if base == 0:
            return float("nan")
    except TypeError:
        return float("nan")
    return ((new / base) - 1.0) * 100.0


def apply_chart_axes(fig, secondary_y=False):
    """공통 차트 축 스타일(그리드/라인 색상)을 일괄 적용."""
    fig.update_xaxes(showgrid=True, gridcolor="#EEF2F6", linecolor="#E2E8F0")
    if secondary_y:
        fig.update_yaxes(showgrid=True, gridcolor="#EEF2F6", linecolor="#E2E8F0", secondary_y=False)
    else:
        fig.update_yaxes(showgrid=True, gridcolor="#EEF2F6", linecolor="#E2E8F0")
    return fig


def make_curve_fig(
    title,
    x_axis,
    y_axis,
    q1,
    y1,
    q2,
    y2,
    scaled_color,
    op_q_base=None,
    op_y_base=None,
    op_q_scaled=None,
    op_y_scaled=None,
    op_fmt="{:.1f}",
    op_unit="",
    height=350,
):
    """기준/스케일링 곡선 + (선택) 운전점 마커 비교 차트를 생성.
    Tab1의 동력/소음 곡선(중복 3블록)을 단일 함수로 통합."""
    fig = go.Figure()
    fig.add_trace(
        go.Scatter(x=q1, y=y1, mode="lines+markers", name="기준", line=dict(color="#64748B", dash="dot"))
    )
    fig.add_trace(
        go.Scatter(x=q2, y=y2, mode="lines+markers", name="스케일링", line=dict(color=scaled_color))
    )
    if (
        op_q_base is not None
        and op_q_scaled is not None
        and op_y_base is not None
        and op_y_scaled is not None
        and np.isfinite(op_y_base)
        and np.isfinite(op_y_scaled)
    ):
        fig.add_trace(
            go.Scatter(
                x=[op_q_base],
                y=[op_y_base],
                mode="markers+text",
                name="기준 운전점",
                text=[f"기준: {op_fmt.format(op_y_base)} {op_unit}"],
                textposition="top center",
                marker=dict(color="#64748B", size=12, symbol="circle", line=dict(color="black", width=1.5)),
            )
        )
        fig.add_trace(
            go.Scatter(
                x=[op_q_scaled],
                y=[op_y_scaled],
                mode="markers+text",
                name="예측 운전점",
                text=[f"예측: {op_fmt.format(op_y_scaled)} {op_unit}"],
                textposition="top center",
                marker=dict(
                    color=scaled_color, size=12, symbol="diamond", line=dict(color="black", width=1.5)
                ),
            )
        )
    fig.update_layout(
        title=dict(
            text=title,
            font=dict(family="Inter, Pretendard, Malgun Gothic, sans-serif", size=16, color="#0F172A"),
        ),
        xaxis_title=x_axis,
        yaxis_title=y_axis,
        plot_bgcolor="#FFFFFF",
        paper_bgcolor="rgba(0,0,0,0)",
        legend=dict(x=0.02, y=0.98, bordercolor="#E2E8F0", borderwidth=1, bgcolor="rgba(255,255,255,0.7)"),
        margin=dict(l=40, r=20, t=55, b=40),
        height=height,
    )
    return apply_chart_axes(fig)


def build_scaling_commentary(
    base_D,
    target_D,
    base_X,
    target_X,
    base_N,
    target_N_calc,
    dim_label,
    use_eff_correction,
    base_eff,
    t_eff,
    md=True,
    fan_category=None,
):
    """직경/폭/회전수/효율 변화에 대한 물리 코멘터리 리스트 생성.
    md=True: 화면용(불릿+볼드), md=False: PPT용(평문). 문구 단일 소스."""
    out = []

    def add(header, body):
        out.append(f"• **{header}**: {body}" if md else f"{header}: {body}")

    if target_D > base_D:
        add(
            f"직경 증가 (+{((target_D/base_D)-1)*100:.1f}%)",
            (
                "축류팬 기하상사에서 유량은 직경의 3승, 동력은 직경의 5승에 비례합니다. 동일 RPM이면 유량과 정압이 함께 증가합니다."
                if is_axial_category(fan_category)
                else "유동 통과 단면적이 직경의 2승에 비례하여 넓어짐에 따라 동일 팁속도 대비 훨씬 큰 풍량을 밀어내며, 팁 주속도가 상승하여 정압이 향상되었습니다."
            ),
        )
    elif target_D < base_D:
        add(
            f"직경 감소 ({((target_D/base_D)-1)*100:.1f}%)",
            "회전 주속도가 떨어지고 유동 유로 폭이 협소해져 풍량과 압력이 크게 감소하므로, 목표 사양을 맞추기 위해 RPM을 크게 높여야 할 수 있습니다.",
        )

    if not is_axial_category(fan_category) and target_X > base_X:
        add(
            f"{dim_label} 증가 (+{((target_X/base_X)-1)*100:.1f}%)",
            "직경 변경 없이 날개의 토출 폭(길이)만 넓어짐으로써 정압에는 큰 변화를 유발하지 않고 풍량만을 선형적으로 비례하여 확장시켰습니다. (동력은 풍량 비례 상승)",
        )
    elif not is_axial_category(fan_category) and target_X < base_X:
        add(
            f"{dim_label} 감소 ({((target_X/base_X)-1)*100:.1f}%)",
            "토출구 좁아짐 현상으로 풍량 유동 면적이 감소하였습니다. 동일 압력 조건 하에서 유량이 저하됩니다.",
        )

    if target_N_calc > base_N:
        add(
            f"회전수 증가 (+{((target_N_calc/base_N)-1)*100:.1f}%)",
            f"회전수가 빨라지면서 유량은 1승에 비례, 압력은 2승에 비례, 동력은 3승에 비례하여 급격히 증가합니다. 특히 소음이 {50*np.log10(target_N_calc/base_N):+.1f} dB(A) 가량 급격히 상승하므로 소음 설계 마진을 고려하십시오.",
        )
    elif target_N_calc < base_N:
        add(
            f"회전수 감소 ({((target_N_calc/base_N)-1)*100:.1f}%)",
            "회전수가 낮아짐에 따라 유량, 압력이 전체적으로 감소하며 동력 및 소음 저감 효과가 있습니다.",
        )

    if use_eff_correction:
        add(
            "효율 보정 작용",
            f"크기 및 RPM 스케일 변화가 반영되어 목표 예측 팬의 정압효율이 기존 {base_eff*100:.1f}%에서 {t_eff*100:.1f}%로 변경 반영되었습니다. 이에 따라 축동력이 조정 계산되었습니다.",
        )
    return out


# -----------------------------------------------------------------------------
# 3. HELPER MATH FUNCTIONS (FAN SIMILARITY & INTERSECTION)
# -----------------------------------------------------------------------------


def require_positive(name, value, *, allow_zero=False):
    values = np.asarray(value, dtype=float)
    invalid = values < 0 if allow_zero else values <= 0
    if not np.all(np.isfinite(values)) or np.any(invalid):
        condition = "0 이상" if allow_zero else "0보다 큰"
        raise ValueError(f"{name}: {condition} 유한한 숫자를 입력하십시오.")


def clean_curve_arrays(*arrays, required=(0, 1)):
    """비유한 필수값 제거, Q 정렬, 중복 Q 평균. 선택 데이터의 결측은 유지한다."""
    values = [np.asarray(a, dtype=float) for a in arrays]
    if not values or any(a.ndim != 1 or a.shape != values[0].shape for a in values):
        raise ValueError("곡선 데이터는 길이가 같은 1차원 배열이어야 합니다.")
    frame = pd.DataFrame({i: a for i, a in enumerate(values)}).replace([np.inf, -np.inf], np.nan)
    frame = frame.dropna(subset=list(required))
    frame = frame[frame[0] >= 0].groupby(0, as_index=False).mean().sort_values(0)
    return tuple(frame[i].to_numpy(float) for i in range(len(values)))


def real_polynomial_roots(a, b, c):
    """2차 이하 방정식의 실근. 작은 근의 소거 오차를 피한다."""
    if a == 0:
        return [-c / b] if b != 0 else []
    scale = max(abs(a), abs(b), abs(c))
    a, b, c = a / scale, b / scale, c / scale
    discriminant = b * b - 4 * a * c
    tolerance = 1e-14 * (b * b + abs(4 * a * c))
    if discriminant < -tolerance:
        return []
    root = np.sqrt(max(discriminant, 0.0))
    q = -0.5 * (b + np.copysign(root, b))
    return [q / a, c / q] if q != 0 else [-b / (2 * a)]


def calc_air_density(temp_c, alt_m):
    if not np.isfinite(temp_c) or temp_c <= -273.15:
        raise ValueError("온도는 절대영도(-273.15°C)보다 높아야 합니다.")
    if not np.isfinite(alt_m) or not -500 <= alt_m <= 11000:
        raise ValueError("간편 고도식의 지원 범위는 -500~11,000 m입니다. 밀도 직접 입력을 사용하십시오.")
    pressure = 101325.0 * (1.0 - 2.25577e-5 * alt_m) ** 5.25588
    return pressure / (R_SPECIFIC_AIR * (temp_c + 273.15))


def calc_air_viscosity(temp_c):
    """건공기 Sutherland 근사 [Pa·s], T[K], β=1.458e-6, S=110.4 K."""
    temperature = np.asarray(temp_c, dtype=float) + 273.15
    require_positive("점성 계산 절대온도", temperature)
    return 1.458e-6 * temperature**1.5 / (temperature + 110.4)


def static_efficiency(q_cmm, p_pa, w_w):
    """동일 운전점 Q·P/W의 정압효율. 동력이 없거나 0 이하면 결측이다."""
    q, p, w = np.broadcast_arrays(
        np.asarray(q_cmm, dtype=float), np.asarray(p_pa, dtype=float), np.asarray(w_w, dtype=float)
    )
    valid = np.isfinite(q) & np.isfinite(p) & np.isfinite(w) & (q >= 0) & (p >= 0) & (w > 0)
    return np.divide(q * p / 60.0, w, out=np.full(q.shape, np.nan), where=valid)


def efficiency_curve_points(q, p, w, extra_q=()):
    """선형 PQ/W 곡선에서 효율을 재계산한다. extra_q 운전점을 곡선에 포함한다."""
    q, p, w = clean_curve_arrays(q, p, w)
    if not len(q):
        return np.array([]), np.array([])
    extra = [v for v in extra_q if v is not None and np.isfinite(v) and q[0] <= v <= q[-1]]
    samples = np.unique(np.r_[q, np.linspace(q[0], q[-1], 201), extra])
    return samples, static_efficiency(samples, np.interp(samples, q, p), np.interp(samples, q, w))


def interpolate_optional_curve(q, values):
    """소음 등 선택 데이터는 측정 범위 안에서만 Q 좌표로 보간한다."""
    return np.array([np.nan if (v := _optional_interp(x, q, values)) is None else v for x in q])


def scale_derived_curve(
    q,
    p,
    w,
    noise,
    d1,
    x1,
    n1,
    d2,
    x2,
    n2,
    rho1,
    rho2,
    use_eff_correction,
    eff_cap,
    c_q=1.0,
    c_p=1.0,
    c_w=1.0,
    fan_category=None,
    *,
    mu1=AIR_VISCOSITY,
    mu2=AIR_VISCOSITY,
):
    """후보점과 표시 곡선에 동일한 점별 효율·제조 상한·보정식을 적용한다."""
    q, p, w, noise = [np.atleast_1d(np.asarray(v, dtype=float)) for v in (q, p, w, noise)]
    eta = np.divide(q * p / 60, w, out=np.zeros_like(q), where=w > 0)
    qs, ps, ws, ns, _ = run_similarity_scaling(
        q,
        p,
        w,
        noise,
        d1,
        x1,
        n1,
        d2,
        x2,
        n2,
        rho1,
        rho2,
        use_eff_correction,
        eta,
        fan_category,
        mu1=mu1,
        mu2=mu2,
    )
    for name, value in [("풍량 보정", c_q), ("정압 보정", c_p), ("동력 보정", c_w), ("효율 상한", eff_cap)]:
        require_positive(name, value)
    if eff_cap > 1:
        raise ValueError("효율 상한은 100% 이하여야 합니다.")
    qs, ps, ws = qs * c_q, ps * c_p, ws * c_q * c_p * c_w
    local_eff = static_efficiency(qs, ps, ws)
    cap_mask = local_eff > eff_cap
    ws = np.where(cap_mask, qs * ps / 60 / eff_cap, ws)
    eta_final = static_efficiency(qs, ps, ws)
    return qs, ps, ws, ns, eta_final


def find_derived_candidates(
    q,
    p,
    w,
    noise,
    d1,
    x1,
    n1,
    target_q,
    target_p,
    rho1,
    rho2,
    max_d,
    max_n,
    use_eff_correction,
    eff_cap,
    c_q=1.0,
    c_p=1.0,
    c_w=1.0,
    fan_category=None,
    *,
    mu1=AIR_VISCOSITY,
    mu2=AIR_VISCOSITY,
):
    """측정 구간 안의 후보 Q를 탐색한다. 유한 격자 탐색이며 전역 최적 보장은 없다."""
    q, p, w, noise = clean_curve_arrays(q, p, w, noise, required=(0, 1, 2))
    if len(q) < 2:
        return []
    samples = np.unique(np.r_[q, np.linspace(q[0], q[-1], 161)])
    results = []
    for qb in samples:
        pb = float(np.interp(qb, q, p))
        wb = float(np.interp(qb, q, w))
        if qb <= 0 or pb <= 0 or wb <= 0:
            continue
        eta = qb * pb / 60 / wb
        if not 0 < eta <= 1:
            continue
        rd = np.sqrt((target_q / c_q / qb) * np.sqrt(pb / (target_p / c_p)) * np.sqrt(rho2 / rho1))
        rn = np.sqrt((target_p / c_p / pb) * (rho1 / rho2)) / rd
        d2, x2, n2 = d1 * rd, x1 * rd, n1 * rn
        if d2 > max_d or n2 > max_n or d2 < 10 or n2 < 10:
            continue
        nb = _optional_interp(qb, q, noise)
        predicted = scale_derived_curve(
            [qb],
            [pb],
            [wb],
            [np.nan if nb is None else nb],
            d1,
            x1,
            n1,
            d2,
            x2,
            n2,
            rho1,
            rho2,
            use_eff_correction,
            eff_cap,
            c_q,
            c_p,
            c_w,
            fan_category,
            mu1=mu1,
            mu2=mu2,
        )
        qf, pf, wf, nf, ef = [float(a[0]) for a in predicted]
        if not 0 < ef <= 1 or not np.isfinite(wf):
            continue
        results.append(
            dict(
                q_ref=float(qb),
                p_ref=pb,
                w_ref=wb,
                noise_ref=nb,
                d2=d2,
                x2=x2,
                n2=n2,
                r_d=rd,
                r_n=rn,
                q=qf,
                p=pf,
                w=wf,
                noise=nf,
                eff=ef,
            )
        )
    return results


def derived_candidate_score(candidate, preference):
    if preference == "최고 효율 순 (Max Efficiency)":
        return -candidate["eff"]
    if preference == "최소 동력 순 (Min Power)":
        return candidate["w"]
    if preference == "최소 소음 순 (Min Noise)":
        return candidate["noise"] if np.isfinite(candidate["noise"]) else np.inf
    return candidate["d2"]


def fit_system_resistance(q_values, p_values):
    q, p = clean_curve_arrays(q_values, p_values)
    valid = (q > 0) & (p > 0)
    q, p = q[valid], p[valid]
    if not len(q):
        raise ValueError("원본 PQ 범위 안의 양수 Q·P 보정점이 필요합니다.")
    exponent = 2.0
    if len(q) >= 2 and np.ptp(np.log(q)) > 1e-8:
        exponent = float(np.clip(np.polyfit(np.log(q), np.log(p), 1)[0], 0.5, 3.0))
    k = float(np.exp(np.mean(np.log(p) - exponent * np.log(q))))
    rmse = float(np.sqrt(np.mean(((k * q**exponent - p) / p * 100) ** 2)))
    return k, exponent, rmse


@st.cache_data(show_spinner=False, max_entries=16)
def excel_bytes(frame, sheet_name="Data"):
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False, sheet_name=sheet_name)
    return buffer.getvalue()


@st.cache_data(show_spinner=False, max_entries=16)
def read_uploaded_table(content, filename):
    stream = io.BytesIO(content)
    return pd.read_csv(stream) if filename.lower().endswith(".csv") else pd.read_excel(stream)


def _prepare_optional_fan_curve(df):
    """Q/P만 필수로 정리하고 W↔정압효율을 가능한 행에서 상호 보완한다."""
    required = ["풍량 Q (CMM)", "정압 P (Pa)"]
    optional = ["축동력 W (W)", "정압효율 ηs (%)", "소음 Lw (dB(A))"]
    if df is None or any(c not in df.columns for c in required):
        return pd.DataFrame(columns=required + optional), 0, 0
    clean = df.copy()
    for c in optional:
        if c not in clean.columns:
            clean[c] = np.nan
    clean = clean[required + optional]
    for c in required + optional:
        clean[c] = pd.to_numeric(clean[c], errors="coerce")
    clean = clean.replace([np.inf, -np.inf], np.nan).dropna(subset=required)
    clean = clean[(clean["풍량 Q (CMM)"] >= 0) & (clean["정압 P (Pa)"] >= 0)]
    if clean.empty:
        return clean, 0, 0
    if (clean["축동력 W (W)"] < 0).any():
        raise ValueError("축동력에는 음수를 입력할 수 없습니다.")
    invalid_eta = clean["정압효율 ηs (%)"].notna() & ~clean["정압효율 ηs (%)"].between(0, 100)
    if invalid_eta.any():
        raise ValueError("정압효율은 0~100% 범위여야 합니다.")

    q_m3s = clean["풍량 Q (CMM)"] / 60.0
    eta = clean["정압효율 ηs (%)"] / 100.0
    can_make_w = clean["축동력 W (W)"].isna() & eta.between(0.001, 1.0) & (q_m3s > 0)
    clean.loc[can_make_w, "축동력 W (W)"] = (
        q_m3s[can_make_w] * clean.loc[can_make_w, "정압 P (Pa)"] / eta[can_make_w]
    )
    can_make_eta = clean["정압효율 ηs (%)"].isna() & (clean["축동력 W (W)"] > 0) & (q_m3s > 0)
    clean.loc[can_make_eta, "정압효율 ηs (%)"] = (
        q_m3s[can_make_eta]
        * clean.loc[can_make_eta, "정압 P (Pa)"]
        / clean.loc[can_make_eta, "축동력 W (W)"]
        * 100.0
    )
    computed = np.divide(q_m3s * clean["정압 P (Pa)"], clean["축동력 W (W)"])
    if (computed > 1.0 + 1e-9).any():
        raise ValueError("단품 Q·P/W 효율이 100%를 초과합니다. 단위와 축동력을 확인하십시오.")
    clean = clean.groupby("풍량 Q (CMM)", as_index=False).mean(numeric_only=True)
    clean = clean.sort_values("풍량 Q (CMM)")
    return clean, int(can_make_w.sum()), int(can_make_eta.sum())


def _optional_interp(x, xp, fp, *, allow_extrapolation=False, max_fraction=0.20):
    """유효 Q 좌표로 보간한다. 범위 밖·단일점의 다른 Q는 기본적으로 결측이다."""
    if x is None or not np.isfinite(x):
        return None
    xp, fp = clean_curve_arrays(xp, fp)
    if len(xp) == 0:
        return None
    exact = np.flatnonzero(np.isclose(xp, x, rtol=1e-10, atol=1e-9))
    if exact.size:
        return float(fp[exact[0]])
    if len(xp) < 2:
        return None
    if xp[0] <= x <= xp[-1]:
        return float(np.interp(x, xp, fp))
    span = xp[-1] - xp[0]
    if (
        not allow_extrapolation
        or not max(0.0, xp[0] - max_fraction * span) <= x <= xp[-1] + max_fraction * span
    ):
        return None
    i, j = (0, 1) if x < xp[0] else (-2, -1)
    return float(fp[i] + (fp[j] - fp[i]) / (xp[j] - xp[i]) * (x - xp[i]))


def _fan_set_curve_at_rpm(
    q,
    p,
    w,
    noise,
    d1,
    x1,
    n1,
    d2,
    x2,
    rpm,
    rho1,
    rho2,
    count,
    arrangement,
    fan_category,
    efficiency_correction=False,
    original_efficiency=0.60,
    *,
    mu1=AIR_VISCOSITY,
    mu2=AIR_VISCOSITY,
):
    """원본 측정 범위를 보존하여 지정 조건의 세트 곡선을 계산한다."""
    q, p, w, noise = clean_curve_arrays(q, p, w, noise)
    sq, sp, sw, slw, _ = run_similarity_scaling(
        q,
        p,
        w,
        noise,
        d1,
        x1,
        n1,
        d2,
        x2,
        rpm,
        rho1,
        rho2,
        efficiency_correction,
        original_efficiency,
        fan_category,
        mu1=mu1,
        mu2=mu2,
    )
    return combine_identical_fans(sq, sp, sw, slw, count, arrangement)


def _solve_required_rpm(
    q_target,
    p_target,
    q,
    p,
    w,
    noise,
    d1,
    x1,
    n1,
    d2,
    x2,
    rho1,
    rho2,
    count,
    arrangement,
    fan_category,
    efficiency_correction=False,
    original_efficiency=0.60,
    preferred_rpm=None,
    allow_limited_extrapolation=False,
    return_details=False,
    *,
    mu1=AIR_VISCOSITY,
    mu2=AIR_VISCOSITY,
):
    """각 원본 PQ 구간에서 b·r²+a·Qtarget·r-Ptarget=0을 풀어 RPM을 구한다.

    원본 범위 내 해를 우선한다. 허용한 경우에만 원본 Q 폭의 20% 이내에서
    양 끝의 기울기로 압력을 외삽한다. 동력·소음은 외삽하지 않는다.
    """

    def finish(rpm, details):
        return (rpm, details) if return_details else rpm

    if not (
        np.isfinite(q_target)
        and q_target > 0
        and np.isfinite(p_target)
        and p_target >= 0
        and np.isfinite(n1)
        and n1 > 0
        and len(q) >= 2
    ):
        return finish(None, {"method": "입력 부족"})
    ref_q, ref_p, _, _ = _fan_set_curve_at_rpm(
        q,
        p,
        w,
        noise,
        d1,
        x1,
        n1,
        d2,
        x2,
        n1,
        rho1,
        rho2,
        count,
        arrangement,
        fan_category,
        efficiency_correction,
        original_efficiency,
        mu1=mu1,
        mu2=mu2,
    )
    if len(ref_q) < 2:
        return finish(None, {"method": "유효 PQ 부족"})
    q_min, q_max = float(ref_q[0]), float(ref_q[-1])
    span = q_max - q_min
    preferred = (
        preferred_rpm
        if preferred_rpm is not None and np.isfinite(preferred_rpm) and preferred_rpm > 0
        else n1
    )
    rpm_min, rpm_max = max(5.0, n1 * 0.02), min(100000.0, n1 * 20.0)
    details = dict(
        reference_q_min=q_min,
        reference_q_max=q_max,
        limited_q_min=max(0.0, q_min - 0.2 * span),
        limited_q_max=q_max + 0.2 * span,
        test_rpm=float(n1),
    )
    segments = list(zip(ref_q[:-1], ref_q[1:], ref_p[:-1], ref_p[1:]))

    def candidates(extrapolate=False):
        candidates_found = []
        for idx, (qa, qb, pa, pb) in enumerate(segments):
            a = (pb - pa) / (qb - qa)
            b = pa - a * qa
            lower, upper = qa, qb
            if extrapolate and idx == 0:
                lower = details["limited_q_min"]
            if extrapolate and idx == len(segments) - 1:
                upper = details["limited_q_max"]
            ratios = real_polynomial_roots(b, a * q_target, -p_target)
            if b == 0 and a == 0 and p_target == 0:
                lo = max(rpm_min / n1, q_target / upper)
                hi = min(rpm_max / n1, q_target / lower if lower > 0 else np.inf)
                ratios = [float(np.clip(preferred / n1, lo, hi))] if lo <= hi else []
            for ratio in ratios:
                if ratio <= 0 or not rpm_min <= ratio * n1 <= rpm_max:
                    continue
                equivalent_q = q_target / ratio
                if lower - 1e-8 <= equivalent_q <= upper + 1e-8:
                    distance = max(q_min - equivalent_q, equivalent_q - q_max, 0.0)
                    candidates_found.append((float(ratio * n1), float(equivalent_q), distance))
        return candidates_found

    options = candidates()
    if not options and allow_limited_extrapolation:
        options = candidates(True)
    if not options:
        return finish(None, {**details, "method": "상사 교점 없음"})
    rpm, equivalent_q, distance = min(options, key=lambda row: abs(row[0] - preferred))
    outside = bool(distance > 1e-7)
    details.update(
        method="제한적 PQ 끝단 외삽 + RPM 상사" if outside else "입력 PQ 범위 내 RPM 상사",
        rpm_ratio=rpm / n1,
        equivalent_q=equivalent_q,
        extrapolation_pct=float(distance / span * 100) if outside else 0.0,
        used_extrapolation=outside,
        solution_count=len(options),
    )
    return finish(rpm, details)


def run_similarity_scaling(
    q,
    p,
    w,
    noise,
    d1,
    x1,
    n1,
    d2,
    x2,
    n2,
    rho1,
    rho2,
    use_eff_correction=False,
    orig_eff=0.6,
    fan_category=None,
    *,
    mu1=AIR_VISCOSITY,
    mu2=AIR_VISCOSITY,
):
    """
    q: flow (CMM)
    p: pressure (Pa)
    w: power (W)
    noise: dB(A)
    x1, x2: centrifugal height (H) or cross-flow length (L).
            Axial fans ignore this independent dimension and use D³/D⁵ laws.
    d1, d2: diameter (D)
    n1, n2: RPM
    mu1, mu2: dynamic viscosity [Pa·s], defaults to equal 20°C approximations.
    orig_eff: retained for call compatibility; actual Q·P/W determines efficiency.
    If Ackeret predicts nonpositive efficiency, retain uncorrected similarity power.
    """
    for name, value in [
        ("기준 직경", d1),
        ("목표 직경", d2),
        ("기준 RPM", n1),
        ("목표 RPM", n2),
        ("기준 밀도", rho1),
        ("목표 밀도", rho2),
        ("기준 점성", mu1),
        ("목표 점성", mu2),
    ]:
        require_positive(name, value)
    if not is_axial_category(fan_category):
        require_positive("기준 폭/길이", x1)
        require_positive("목표 폭/길이", x2)
    # Ratios
    r_n = n2 / n1
    r_d = d2 / d1
    # 원심/횡류는 폭·길이가 독립 변수다. 축류는 독립 높이가 없으므로
    # 기하상사 시 그 효과를 직경비 한 번으로 대체해 D³/D⁵ 법칙을 만든다.
    r_x = r_d if is_axial_category(fan_category) else x2 / x1
    r_rho = rho2 / rho1

    # Flow: 원심/횡류 Q∝N·D²·X, 축류는 effective X_ratio=D_ratio → Q∝N·D³
    q2 = q * r_n * (r_d**2) * r_x

    # Scale Pressure (P2 = P1 * N_ratio^2 * D_ratio^2 * rho_ratio)
    p2 = p * (r_n**2) * (r_d**2) * r_rho

    # Power: 원심/횡류 W∝ρ·N³·D⁴·X, 축류는 effective X_ratio=D_ratio → W∝ρ·N³·D⁵
    w2 = w * (r_n**3) * (r_d**4) * r_x * r_rho

    # orig_eff는 기존 호출 호환용이다. 효율의 기준은 항상 입력 Q·P/W이다.
    eta = static_efficiency(q, p, w)
    actual_w2 = w2
    if use_eff_correction:
        valid_eff = np.isfinite(eta) & (eta > 0) & (eta <= 1)
        r_re = r_rho * (r_d**2) * r_n * (mu1 / mu2)
        corrected = 1.0 - (1.0 - eta) * (0.5 + 0.5 * (1.0 / r_re) ** ACKERET_EXP)
        # 보정식이 비양수 효율을 내는 영역은 모델 적용 불가: 기본 상사를 유지한다.
        applicable = valid_eff & np.isfinite(corrected) & (corrected > 0) & (corrected <= 1)
        eta_b, corrected_b, applicable_b = np.broadcast_arrays(eta, corrected, applicable)
        factor = np.divide(eta_b, corrected_b, out=np.ones(eta_b.shape), where=applicable_b)
        actual_w2 = w2 * factor
    target_eff = static_efficiency(q2, p2, actual_w2)

    # Noise: 원심/횡류 +60logD+10logX, 축류는 +70logD (KN 모델과 일치)
    # Clip ratio inputs to avoid log of <= 0 (np.maximum: 스칼라·배열 공용)
    r_n_c = np.maximum(r_n, 1e-5)
    r_d_c = np.maximum(r_d, 1e-5)
    r_x_c = np.maximum(r_x, 1e-5)
    r_rho_c = np.maximum(r_rho, 1e-5)

    noise2 = (
        noise + 50 * np.log10(r_n_c) + 60 * np.log10(r_d_c) + 10 * np.log10(r_x_c) + 20 * np.log10(r_rho_c)
    )

    return q2, p2, actual_w2, noise2, target_eff


def fit_set_noise_log_rpm(rpm_values, noise_values):
    """세트 raw 소음을 L=a+b·log10(N)으로 회귀한다.

    서로 다른 RPM의 유효점이 2개 이상일 때만 모델을 반환한다. 이 회귀는
    현재 세트의 운전 경로에서 측정된 소음 경향이며 별도의 실험 잔차 보정은
    포함하지 않는다.
    """
    rpm = np.asarray(rpm_values, dtype=float)
    noise = np.asarray(noise_values, dtype=float)
    valid = np.isfinite(rpm) & np.isfinite(noise) & (rpm > 0)
    if np.count_nonzero(valid) < 2:
        return None

    rpm = rpm[valid]
    noise = noise[valid]
    grouped = (
        pd.DataFrame({"RPM": rpm, "소음": noise})
        .groupby("RPM", as_index=False)["소음"]
        .mean()
        .sort_values("RPM")
    )
    if len(grouped) < 2:
        return None

    rpm = grouped["RPM"].to_numpy(float)
    noise = grouped["소음"].to_numpy(float)
    x = np.log10(rpm)
    slope, intercept = np.polyfit(x, noise, 1)
    fitted = intercept + slope * x
    residual = noise - fitted
    rmse = float(np.sqrt(np.mean(residual**2)))
    ss_tot = float(np.sum((noise - np.mean(noise)) ** 2))
    r_squared = 1.0 - float(np.sum(residual**2)) / ss_tot if ss_tot > 0 else 1.0
    return {
        "slope": float(slope),
        "intercept": float(intercept),
        "rmse": rmse,
        "r_squared": float(r_squared),
        "rpm_min": float(rpm.min()),
        "rpm_max": float(rpm.max()),
        "count": int(len(rpm)),
    }


def predict_set_noise_from_rpm(model, rpm):
    """fit_set_noise_log_rpm 결과로 지정 RPM의 세트 소음을 계산한다."""
    if model is None or rpm is None or not np.isfinite(rpm) or rpm <= 0:
        return None
    return float(model["intercept"] + model["slope"] * np.log10(float(rpm)))


def fan_geometry_noise_delta(d1, x1, d2, x2, rho1, rho2, fan_category=None):
    """RPM 효과를 제외한 팬 종류별 형상·밀도 소음 상사 보정값을 반환한다.

    축류팬:       70log(D2/D1)
    원심/횡류팬: 60log(D2/D1) + 10log(X2/X1)
    공통:         20log(rho2/rho1)
    """
    values = (d1, d2, rho1, rho2)
    if any(v is None or not np.isfinite(v) or v <= 0 for v in values):
        return None
    r_d = float(d2) / float(d1)
    r_rho = float(rho2) / float(rho1)
    if is_axial_category(fan_category):
        geometry = 70.0 * np.log10(r_d)
    else:
        if x1 is None or x2 is None or not np.isfinite(x1) or not np.isfinite(x2) or x1 <= 0 or x2 <= 0:
            return None
        geometry = 60.0 * np.log10(r_d) + 10.0 * np.log10(float(x2) / float(x1))
    return float(geometry + 20.0 * np.log10(r_rho))


def find_system_intersection(q_arr, p_arr, k_coef):
    """구간별 선형 팬 곡선과 KQ²의 정확한 교점 중 최대 풍량을 반환한다."""
    require_positive("시스템 저항계수 K", k_coef, allow_zero=True)
    q, p = clean_curve_arrays(q_arr, p_arr)
    if len(q) < 2:
        return None
    roots = []
    for qa, qb, pa, pb in zip(q[:-1], q[1:], p[:-1], p[1:]):
        a = (pb - pa) / (qb - qa)
        b = pa - a * qa
        if k_coef == 0 and a == 0 and b == 0:
            roots.append(float(qb))
        for root in real_polynomial_roots(k_coef, -a, -b):
            if qa - 1e-9 <= root <= qb + 1e-9:
                roots.append(float(np.clip(root, qa, qb)))
    return max(roots) if roots else None


def combine_identical_fans(q, p, w, noise, count=1, arrangement="병렬 (Parallel)"):
    """단품 곡선을 동일 팬 여러 대의 세트 곡선으로 변환한다."""
    count = max(int(count), 1)
    q = np.asarray(q, dtype=float)
    p = np.asarray(p, dtype=float)
    w = np.asarray(w, dtype=float)
    noise = np.asarray(noise, dtype=float)
    noise_add = 10.0 * np.log10(count)
    if arrangement == "직렬 (Series)":
        return q, p * count, w * count, noise + noise_add
    return q * count, p, w * count, noise + noise_add


def find_curve_intersection(q_arr, fan_p_arr, system_p_arr):
    """같은 Q 격자의 두 구간별 선형 곡선에서 최대 풍량 교점을 반환한다."""
    q, fan, system = clean_curve_arrays(q_arr, fan_p_arr, system_p_arr, required=(0, 1, 2))
    if len(q) < 2:
        return None
    diff = fan - system
    roots = list(q[np.isclose(diff, 0.0, atol=1e-9, rtol=0.0)])
    for i in range(len(q) - 1):
        if diff[i] * diff[i + 1] < 0:
            roots.append(q[i] + diff[i] / (diff[i] - diff[i + 1]) * (q[i + 1] - q[i]))
    return float(max(roots)) if roots else None


def extend_fan_curve_to_zero_pressure(q, p, w=None, noise=None, max_extension_ratio=1.5):
    """표시용 끝단 연장. 원본 범위 메타데이터 없이 계산 입력으로 사용하지 않는다."""
    w = np.full_like(np.asarray(q, dtype=float), np.nan) if w is None else w
    noise = np.full_like(np.asarray(q, dtype=float), np.nan) if noise is None else noise
    q, p, w, noise = clean_curve_arrays(q, p, w, noise)
    if len(q) < 2 or p[-1] <= 0:
        return q, p, w, noise, False
    slope = (p[-1] - p[-2]) / (q[-1] - q[-2])
    if slope >= 0:
        return q, p, w, noise, False
    q_zero = q[-1] - p[-1] / slope
    if q_zero > q[-1] * max_extension_ratio:
        return q, p, w, noise, False
    return np.append(q, q_zero), np.append(p, 0.0), np.append(w, np.nan), np.append(noise, np.nan), True


RESISTANCE_POINT_MODEL = "기준점 압손"
RESISTANCE_VELOCITY_MODEL = "면풍속 1·2차식"
RESISTANCE_K_MODEL = "손실계수 K"


def component_pressure_by_model(q, part, rho_old=1.2, rho_new=1.2):
    """부품 한 행의 기존·신규 압손곡선을 선택 모델로 계산한다.

    q는 세트 풍량 [CMM], 면적은 m²이다. 면풍속식은
    ΔP=a1·v+a2·v², 손실계수식은 ΔP=K·ρv²/2를 사용한다.
    """
    q_values = np.maximum(np.asarray(q, dtype=float), 0.0)

    def number(name, default=0.0):
        try:
            value = float(part.get(name, default))
            return value if np.isfinite(value) else float(default)
        except (TypeError, ValueError):
            return float(default)

    def optional_number(name):
        try:
            value = float(part.get(name, np.nan))
            return value if np.isfinite(value) else None
        except (TypeError, ValueError):
            return None

    mode = str(part.get("입력 모델", RESISTANCE_POINT_MODEL))
    old_area = max(number("기존 면적 m²", 0.0), 1e-9)
    new_area = max(number("신규 면적 m²", old_area), 1e-9)
    old_velocity = (q_values / 60.0) / old_area
    new_velocity = (q_values / 60.0) / new_area

    if mode == RESISTANCE_VELOCITY_MODEL:
        a1 = max(number("1차계수 a1", 0.0), 0.0)
        a2 = max(number("2차계수 a2", 0.0), 0.0)
        new_a1_input = optional_number("신규 1차계수 a1")
        new_a2_input = optional_number("신규 2차계수 a2")
        new_a1 = a1 if new_a1_input is None else max(new_a1_input, 0.0)
        new_a2 = a2 if new_a2_input is None else max(new_a2_input, 0.0)
        old_pressure = a1 * old_velocity + a2 * old_velocity**2
        new_pressure = new_a1 * new_velocity + new_a2 * new_velocity**2
    elif mode == RESISTANCE_K_MODEL:
        loss_k = max(number("손실계수 K", 0.0), 0.0)
        new_k_input = optional_number("신규 손실계수 K")
        new_k = loss_k if new_k_input is None else max(new_k_input, 0.0)
        old_pressure = loss_k * max(float(rho_old), 1e-9) * old_velocity**2 / 2.0
        new_pressure = new_k * max(float(rho_new), 1e-9) * new_velocity**2 / 2.0
    else:
        q_ref = max(number("기준 풍량 CMM", 0.0), 1e-9)
        dp_ref = max(number("기준 압손 Pa", 0.0), 0.0)
        exponent = float(np.clip(number("압손지수 m", 2.0), 0.5, 3.0))
        old_pressure = dp_ref * (q_values / q_ref) ** exponent
        # 기준 압손은 Qref, 기존 면적에서의 측정값이므로 면풍속비로 신규를 계산한다.
        new_pressure = dp_ref * ((q_values / new_area) / (q_ref / old_area)) ** exponent

    return np.maximum(old_pressure, 0.0), np.maximum(new_pressure, 0.0)


# -----------------------------------------------------------------------------
# 4. SIDEBAR - GLOBAL CONFIGURATIONS
# -----------------------------------------------------------------------------
def render_sidebar():
    """모바일: 사이드바 대신 화면 상단의 접이식 설정 패널을 사용한다."""
    with st.container(key="m_settings"):
        panel = st.expander("⚙️  설계 조건 · 단위 · 운전 환경", expanded=False)
    with panel:
        basic_tab, unit_tab, env_tab = st.tabs(["팬 계열", "단위", "운전 환경"])

    with basic_tab:
        st.text_input(
            "작업 이름",
            value="default",
            key="input_project",
            on_change=_switch_project,
            help="다른 작업 이름은 별도 입력을 저장합니다. 같은 이름은 같은 저장 작업을 사용합니다.",
        )
        fan_category = st.selectbox(
            "팬 계열", [CENTRIFUGAL_CATEGORY, AXIAL_CATEGORY, CROSSFLOW_CATEGORY]
        )
        is_axial = is_axial_category(fan_category)

        if fan_category == CENTRIFUGAL_CATEGORY:
            st.info(r"""**원심팬** · 고압·중풍량, 날개 외주방향 원심력 이용  
비속도 $\sigma \approx 0.15 \sim 1.80$ · 집진기, 실외기, 환기조화기""")
        elif fan_category == AXIAL_CATEGORY:
            st.info(r"""**축류팬** · 저압·대풍량, 축 방향 유체 통과  
비속도 $\sigma \approx 1.50 \sim 8.00$ · 환기 팬, 냉각탑, 라디에이터""")
        else:  # 횡류팬 계열
            st.info(r"""**횡류팬** · 저압·대면적 균일류, 입출구 직교형  
비속도 $\sigma \approx 0.05 \sim 0.35$ · 에어커튼, 벽걸이 실내기""")

        if is_axial:
            st.caption(
                "축류팬은 독립 높이(H)를 사용하지 않습니다. 직경과 RPM만으로 "
                "Q∝ND³, P∝ρN²D², W∝ρN³D⁵를 적용하며 높이 입력과 높이 스윕은 비활성화됩니다. "
                "허브비·블레이드 피치 등 형상이 기하학적으로 유사하다는 조건입니다."
            )
        elif fan_category == CENTRIFUGAL_CATEGORY:
            st.caption("적용 모델: Q∝ND²H, P∝ρN²D², W∝ρN³D⁴H (높이 H 독립 조정)")
        else:
            st.caption("적용 모델: Q∝ND²L, P∝ρN²D², W∝ρN³D⁴L (길이 L 독립 조정)")

    # Label based on fan type
    if fan_category == CROSSFLOW_CATEGORY:
        dim_label = "길이 (L)"
        dim_key = "L"
    elif is_axial:
        dim_label = "높이 (H, 미적용)"
        dim_key = "H_DISABLED"
    else:
        dim_label = "높이 (H)"
        dim_key = "H"

    with unit_tab:
        flow_unit = st.selectbox("풍량 단위", list(UNIT_FLOW.keys()))
        pressure_unit = st.selectbox("정압 단위", list(UNIT_PRESSURE.keys()))
        power_unit = st.selectbox("동력 단위", list(UNIT_POWER.keys()))
        st.caption("단위 변경 시 입력값을 자동 환산합니다. 세트·실험 데이터 탭은 CMM·Pa·W 고정 단위입니다.")

    with env_tab:
        density_mode = st.radio(
            "밀도 설정 방식", ["간편 온도/고도 설정", "직접 입력"], key="global_density_mode", horizontal=True
        )

        if density_mode == "간편 온도/고도 설정":
            st.markdown("**기준 상태** → **목표 상태**")
            t_col1, t_col2 = st.columns(2)
            t1 = t_col1.number_input(
                "기준 온도 T1 [°C]",
                value=20.0,
                step=1.0,
                min_value=-100.0,
                max_value=200.0,
                key="global_t1",
            )
            t2 = t_col2.number_input(
                "목표 온도 T2 [°C]",
                value=20.0,
                step=1.0,
                min_value=-100.0,
                max_value=200.0,
                key="global_t2",
            )
            h_col1, h_col2 = st.columns(2)
            h1 = h_col1.number_input(
                "기준 고도 h1 [m]",
                value=0.0,
                step=100.0,
                min_value=-500.0,
                max_value=11000.0,
                key="global_h1",
            )
            h2 = h_col2.number_input(
                "목표 고도 h2 [m]",
                value=0.0,
                step=100.0,
                min_value=-500.0,
                max_value=11000.0,
                key="global_h2",
            )

            air_density_1 = calc_air_density(t1, h1)
            air_density_2 = calc_air_density(t2, h2)
            air_viscosity_1 = float(calc_air_viscosity(t1))
            air_viscosity_2 = float(calc_air_viscosity(t2))

            st.info(f"계산된 밀도 · **ρ1 {air_density_1:.3f}** → **ρ2 {air_density_2:.3f}** kg/m³")
            st.caption(
                "ISA 표준대기 기압식 + 건공기 이상기체 기준입니다 (습도 영향 미반영, 고습 환경에서는 실측 밀도 직접 입력을 권장)."
            )
        else:
            r_col1, r_col2 = st.columns(2)
            air_density_1 = r_col1.number_input(
                "기준 밀도 ρ1 [kg/m³]",
                value=1.20,
                min_value=0.1,
                max_value=3.0,
                step=0.01,
                format="%.3f",
                key="global_rho1",
            )
            air_density_2 = r_col2.number_input(
                "목표 밀도 ρ2 [kg/m³]",
                value=1.20,
                min_value=0.1,
                max_value=3.0,
                step=0.01,
                format="%.3f",
                key="global_rho2",
            )
            m_col1, m_col2 = st.columns(2)
            air_viscosity_1 = (
                m_col1.number_input(
                    "기준 점성 μ1 [μPa·s]",
                    min_value=1.0,
                    value=18.1,
                    step=0.1,
                    key="global_mu1",
                    help="건공기 20°C의 근삿값은 약 18.1 μPa·s입니다.",
                )
                * 1e-6
            )
            air_viscosity_2 = (
                m_col2.number_input(
                    "목표 점성 μ2 [μPa·s]", min_value=1.0, value=18.1, step=0.1, key="global_mu2"
                )
                * 1e-6
            )

        st.caption(
            f"동적 점성 μ₁ {air_viscosity_1 * 1e6:.2f} → μ₂ {air_viscosity_2 * 1e6:.2f} μPa·s. "
            "온도 입력 모드에서는 건공기 Sutherland 식으로 자동 계산합니다."
        )

        use_eff_correction = st.toggle(
            "스케일별 효율 보정 적용",
            value=True,
            help="Q·P/W 정압효율과 Re₂/Re₁=(ρ₂/ρ₁)(N₂/N₁)(D₂/D₁)²(μ₁/μ₂)로 Ackeret 보정을 계산합니다.",
        )
        if use_eff_correction:
            st.caption(
                "보정식이 0 이하의 효율을 예측하는 조건에서는 효율 보정을 적용할 수 없어 기본 상사 동력을 사용합니다."
            )

    # Display units for UI (stripping parentheses/spaces)
    flow_unit_disp = flow_unit.split(" ")[0]
    pressure_unit_disp = pressure_unit.split(" ")[0]
    power_unit_disp = power_unit.split(" ")[0]

    return {
        "fan_category": fan_category,
        "is_axial": is_axial,
        "dim_label": dim_label,
        "dim_key": dim_key,
        "flow_unit": flow_unit,
        "pressure_unit": pressure_unit,
        "power_unit": power_unit,
        "flow_unit_disp": flow_unit_disp,
        "pressure_unit_disp": pressure_unit_disp,
        "power_unit_disp": power_unit_disp,
        "air_density_1": air_density_1,
        "air_density_2": air_density_2,
        "air_viscosity_1": air_viscosity_1,
        "air_viscosity_2": air_viscosity_2,
        "use_eff_correction": use_eff_correction,
    }


def main():
    render_theme()
    st.session_state["_chart_counter"] = 0
    st.markdown(
        """
    <div class="m-hero">
      <div class="logo">💨</div>
      <div><div class="eyebrow">FAN STUDIO · MOBILE</div>
      <h1>팬 성능 예측 · 설계</h1>
      <p>상사 예측 · 운전점 분석 · 세트 교체 시뮬레이션</p></div>
    </div>
    """,
        unsafe_allow_html=True,
    )
    settings = render_sidebar()
    fan_category = settings["fan_category"]
    is_axial = settings["is_axial"]
    dim_label = settings["dim_label"]
    dim_key = settings["dim_key"]
    flow_unit = settings["flow_unit"]
    pressure_unit = settings["pressure_unit"]
    power_unit = settings["power_unit"]
    flow_unit_disp = settings["flow_unit_disp"]
    pressure_unit_disp = settings["pressure_unit_disp"]
    power_unit_disp = settings["power_unit_disp"]
    air_density_1 = settings["air_density_1"]
    air_density_2 = settings["air_density_2"]
    air_viscosity_1 = settings["air_viscosity_1"]
    air_viscosity_2 = settings["air_viscosity_2"]
    use_eff_correction = settings["use_eff_correction"]
    st.markdown(
        f"""
    <div class="m-chips">
      <span>{fan_category.split(" (")[0]}</span><span>{flow_unit_disp} · {pressure_unit_disp} · {power_unit_disp}</span>
      <span>ρ <b>{air_density_1:.3f} → {air_density_2:.3f}</b></span>
      <span>효율 보정 <b>{'ON' if use_eff_correction else 'OFF'}</b></span>
    </div>
    """,
        unsafe_allow_html=True,
    )
    containers = st.tabs(
        ["📈 단품 PQ", "🏠 세트 예측", "📐 경험식", "🧪 실험 기반", "⚖️ 1:1 비교", "🔍 설계 탐색", "🌀 덕트 손실"]
    )
    renderers = [
        render_tab_pq,
        render_tab_set,
        render_tab_empirical,
        render_tab_experiments,
        render_tab_comparison,
        render_tab_sweep,
        render_tab_duct,
    ]
    for container, renderer in zip(containers, renderers):
        with container:
            try:
                renderer(settings)
            except (ValueError, TypeError, ZeroDivisionError) as error:
                st.error(f"입력값을 확인하십시오: {error}")


# =============================================================================
# TAB 1: 단품 PQ선도 스케일링
# =============================================================================
def render_tab_pq(settings):
    """단품 PQ 화면. 전역 입력은 settings로 전달한다."""
    fan_category = settings["fan_category"]
    is_axial = settings["is_axial"]
    dim_label = settings["dim_label"]
    dim_key = settings["dim_key"]
    flow_unit = settings["flow_unit"]
    pressure_unit = settings["pressure_unit"]
    power_unit = settings["power_unit"]
    flow_unit_disp = settings["flow_unit_disp"]
    pressure_unit_disp = settings["pressure_unit_disp"]
    power_unit_disp = settings["power_unit_disp"]
    air_density_1 = settings["air_density_1"]
    air_density_2 = settings["air_density_2"]
    air_viscosity_1 = settings["air_viscosity_1"]
    air_viscosity_2 = settings["air_viscosity_2"]
    use_eff_correction = settings["use_eff_correction"]

    st.markdown("### 단품 PQ선도 기반 상사 예측 및 덕트 저항 곡선 매칭")
    if is_axial:
        st.markdown(
            "실제 축류팬의 운전 데이터를 입력하고 직경과 RPM 변화에 따른 전체 성능 곡선을 표준 축류팬 상사법칙으로 예측합니다."
        )
    else:
        st.markdown(
            "실제 가동되는 단품 팬의 운전 데이터를 입력하고, 직경, 높이/길이 및 RPM이 변경되었을 때의 전체 성능 곡선을 예측합니다."
        )

    col1, col2 = st.columns([1.3, 2])

    with col1:
        with st.container(border=True, key="panel_1917"):
            st.markdown("#### 1. 팬 물리 스펙 입력")

            # Dimensions
            st.markdown("##### 기준 사양 (Baseline)")
            subcol1, subcol2, subcol3 = st.columns(3)
            with subcol1:
                d1 = st.number_input(
                    "직경 D1 (mm)", value=300.0, min_value=10.0, step=10.0, key="d1_tab1", format="%.0f"
                )
            with subcol2:
                if is_axial:
                    st.text_input(
                        "독립 높이 H1", value="미적용 (직경 D로 기하상사)", disabled=True, key="x1_tab1_axial"
                    )
                    x1 = d1
                else:
                    x1 = st.number_input(
                        f"{dim_label}1 (mm)",
                        value=100.0,
                        min_value=5.0,
                        step=5.0,
                        key="x1_tab1",
                        format="%.0f",
                    )
            with subcol3:
                n1 = st.number_input(
                    "회전수 N1 (RPM)", value=1500.0, min_value=10.0, step=50.0, key="n1_tab1", format="%.0f"
                )

            st.markdown("##### 목표 사양 (Target)")
            subcol4, subcol5, subcol6 = st.columns(3)
            with subcol4:
                d2 = st.number_input(
                    "직경 D2 (mm)", value=330.0, min_value=10.0, step=10.0, key="d2_tab1", format="%.0f"
                )
            with subcol5:
                if is_axial:
                    st.text_input(
                        "독립 높이 H2", value="미적용 (직경 D로 기하상사)", disabled=True, key="x2_tab1_axial"
                    )
                    x2 = d2
                else:
                    x2 = st.number_input(
                        f"{dim_label}2 (mm)",
                        value=110.0,
                        min_value=5.0,
                        step=5.0,
                        key="x2_tab1",
                        format="%.0f",
                    )
            with subcol6:
                n2 = st.number_input(
                    "회전수 N2 (RPM)", value=1600.0, min_value=10.0, step=50.0, key="n2_tab1", format="%.0f"
                )

            st.markdown("#### 2. 시스템 저항선 & 목표 운전점")
            has_system_curve = st.checkbox("덕트 계통 저항선 플로팅", value=True)

            # Default system target
            sys_q = unit_number_input(
                f"정격 요구 유량 ({flow_unit_disp})",
                value=35.0,
                min_value=0.1,
                step=1.0,
                quantity="flow",
                unit=flow_unit,
            )

            sys_p = unit_number_input(
                f"정격 요구 압력 ({pressure_unit_disp})",
                value=15.0,
                min_value=0.1,
                step=1.0,
                quantity="pressure",
                unit=pressure_unit,
            )

            st.markdown("---")
            st.markdown("##### 차트 보기 옵션 (축 범위)")
            custom_range = st.checkbox(
                "축 범위 수동 설정",
                value=False,
                help="정압성능 비교 곡선의 X축(풍량) 및 Y축(정압) 최대값을 직접 지정합니다.",
            )
            if custom_range:
                subcol_range_x, subcol_range_y = st.columns(2)
                with subcol_range_x:
                    user_x_max = unit_number_input(
                        f"풍량 최대값 ({flow_unit_disp})",
                        value=100.0,
                        min_value=1.0,
                        step=10.0,
                        quantity="flow",
                        unit=flow_unit,
                    )
                with subcol_range_y:
                    user_y_max = unit_number_input(
                        f"정압 최대값 ({pressure_unit_disp})",
                        value=30.0,
                        min_value=1.0,
                        step=5.0,
                        quantity="pressure",
                        unit=pressure_unit,
                    )
            else:
                user_x_max = None
                user_y_max = None

        with st.container(border=True, key="panel_1968"):
            st.markdown("#### 3. 기준 PQ 데이터 입력")
            st.markdown("데이터를 아래 표에서 직접 편집할 수 있습니다.")

            edited_pq = render_pq_editor(flow_unit, pressure_unit, power_unit)

    with col2:
        # Perform calculations
        try:
            # Clean dataframe: drop rows with any null/None/NaN values in required columns
            if edited_pq is None:
                return
            req_cols = [
                f"풍량 ({flow_unit_disp})",
                f"정압 ({pressure_unit_disp})",
                f"동력 ({power_unit_disp})",
            ]
            noise_col = "소음 (dB(A))"

            cleaned_df = _normalize_editor_dataframe(edited_pq, req_cols + [noise_col])
            invalid = (
                cleaned_df[req_cols].isna().any(axis=1)
                | (cleaned_df[req_cols[0]] < 0)
                | (cleaned_df[req_cols[1]] < 0)
                | (cleaned_df[req_cols[2]] <= 0)
            )
            if invalid.any():
                st.warning(
                    f"필수값 누락·음수 Q/P·0 이하 동력이 있는 {int(invalid.sum())}개 행을 제외했습니다."
                )
            cleaned_df = (
                cleaned_df.loc[~invalid]
                .groupby(req_cols[0], as_index=False)
                .mean(numeric_only=True)
                .sort_values(req_cols[0])
            )
            if len(cleaned_df) < 2:
                raise ValueError("서로 다른 풍량의 유효 데이터가 2개 이상 필요합니다.")
            has_noise = np.isfinite(cleaned_df[noise_col]).any()
            if has_noise:
                cleaned_df[noise_col] = interpolate_optional_curve(
                    cleaned_df[req_cols[0]].to_numpy(), cleaned_df[noise_col].to_numpy()
                )

            # Extract lists
            q1_vals = cleaned_df[f"풍량 ({flow_unit_disp})"].values
            p1_vals = cleaned_df[f"정압 ({pressure_unit_disp})"].values
            w1_vals = cleaned_df[f"동력 ({power_unit_disp})"].values

            if has_noise:
                db1_vals = cleaned_df[noise_col].values
            else:
                db1_vals = np.full_like(q1_vals, np.nan)

            # Sort arrays by Q
            sort_idx = np.argsort(q1_vals)
            q1_vals = q1_vals[sort_idx]
            p1_vals = p1_vals[sort_idx]
            w1_vals = w1_vals[sort_idx]
            db1_vals = db1_vals[sort_idx]

            # 원본 각 점의 Q·P/W 정압효율 계산 및 물리적 입력 검증
            peak_eff = 0.5
            eff_list = []
            for i in range(len(q1_vals)):
                q_m3s = convert_flow(q1_vals[i], flow_unit, "CMM (m³/min)") / 60.0
                p_pa = convert_pressure(p1_vals[i], pressure_unit, "Pa")
                w_w = convert_power(w1_vals[i], power_unit, "W")
                if w_w > 0:
                    eff = (q_m3s * p_pa) / w_w
                    eff_list.append(eff)
                else:
                    eff_list.append(0.0)
            if len(eff_list) > 0:
                peak_eff = max(eff_list)
                if peak_eff > 1.0:
                    raise ValueError(
                        "Q·P/W로 계산한 정압효율이 100%를 초과합니다. 단위와 동력 데이터를 확인하십시오."
                    )

            # Scale each point
            q2_vals = []
            p2_vals = []
            w2_vals = []
            db2_vals = []
            target_effs = []
            eff1_vals = []
            eff2_vals = []

            for i in range(len(q1_vals)):
                # Convert to base internal units
                q_base = convert_flow(q1_vals[i], flow_unit, "CMM (m³/min)")
                p_base = convert_pressure(p1_vals[i], pressure_unit, "Pa")
                w_base = convert_power(w1_vals[i], power_unit, "W")
                db_base = db1_vals[i]

                # Similarity scaling
                q_sc, p_sc, w_sc, db_sc, t_eff = run_similarity_scaling(
                    q_base,
                    p_base,
                    w_base,
                    db_base,
                    d1,
                    x1,
                    n1,
                    d2,
                    x2,
                    n2,
                    air_density_1,
                    air_density_2,
                    use_eff_correction,
                    eff_list[i],
                    fan_category,
                    mu1=air_viscosity_1,
                    mu2=air_viscosity_2,
                )

                # Convert back to display units
                q2_vals.append(convert_flow(q_sc, "CMM (m³/min)", flow_unit))
                p2_vals.append(convert_pressure(p_sc, "Pa", pressure_unit))
                w2_vals.append(convert_power(w_sc, "W", power_unit))
                db2_vals.append(db_sc)
                target_effs.append(t_eff)

                # Calculate local static efficiencies in %
                # eta_s = Q(m3/s) * P(Pa) / W(W) * 100
                q_m3s = q_base / 60.0
                local_eff1 = (q_m3s * p_base) / w_base * 100.0 if w_base > 0 else 0.0
                local_eff1 = np.clip(local_eff1, 0.0, 100.0)

                q_sc_m3s = q_sc / 60.0
                local_eff2 = (q_sc_m3s * p_sc) / w_sc * 100.0 if w_sc > 0 else 0.0
                local_eff2 = np.clip(local_eff2, 0.0, 100.0)

                eff1_vals.append(local_eff1)
                eff2_vals.append(local_eff2)

            q2_vals = np.array(q2_vals)
            p2_vals = np.array(p2_vals)
            w2_vals = np.array(w2_vals)
            db2_vals = np.array(db2_vals)
            eff1_vals = np.array(eff1_vals)
            eff2_vals = np.array(eff2_vals)

            # Build system curve
            # System resistance coefficient: k = P_req / Q_req^2
            k_coef = 0.0
            if sys_q > 0:
                q_req_cmm = convert_flow(sys_q, flow_unit, "CMM (m³/min)")
                p_req_pa = convert_pressure(sys_p, pressure_unit, "Pa")
                k_coef = p_req_pa / (q_req_cmm**2) if q_req_cmm > 0 else 0

            q_sys_plot = np.linspace(0, max(q2_vals.max(), q1_vals.max()) * 1.1, 100)
            # System curve in Pa
            p_sys_plot_pa = k_coef * (convert_flow(q_sys_plot, flow_unit, "CMM (m³/min)") ** 2)
            p_sys_plot = convert_pressure(p_sys_plot_pa, "Pa", pressure_unit)

            # Find operating points
            # Base intersection
            k_coef_base = (
                convert_pressure(sys_p, pressure_unit, "Pa")
                / (convert_flow(sys_q, flow_unit, "CMM (m³/min)") ** 2)
                if sys_q > 0
                else 0
            )
            q1_base_arr = np.array([convert_flow(v, flow_unit, "CMM (m³/min)") for v in q1_vals])
            p1_base_arr = np.array([convert_pressure(v, pressure_unit, "Pa") for v in p1_vals])

            q2_base_arr = np.array([convert_flow(v, flow_unit, "CMM (m³/min)") for v in q2_vals])
            p2_base_arr = np.array([convert_pressure(v, pressure_unit, "Pa") for v in p2_vals])

            q_int_base_cmm = find_system_intersection(q1_base_arr, p1_base_arr, k_coef_base)
            q_int_scaled_cmm = find_system_intersection(q2_base_arr, p2_base_arr, k_coef_base)

            # Precalculate operating points parameters (Power, Noise, Efficiency, Pressure)
            w_int_base, db_int_base, eff_int_base = None, None, None
            w_int_scaled, db_int_scaled, eff_int_scaled = None, None, None
            p_int_base_disp, p_int_scaled_disp = 0.0, 0.0

            if q_int_base_cmm is not None:
                q_int_base_disp = convert_flow(q_int_base_cmm, "CMM (m³/min)", flow_unit)
                p_int_base_pa = k_coef_base * (q_int_base_cmm**2)
                p_int_base_disp = convert_pressure(p_int_base_pa, "Pa", pressure_unit)

                w_int_base = np.interp(q_int_base_disp, q1_vals, w1_vals)
                db_int_base = np.interp(q_int_base_disp, q1_vals, db1_vals)
                eff_int_base = (
                    float(
                        static_efficiency(
                            q_int_base_cmm, p_int_base_pa, convert_power(w_int_base, power_unit, "W")
                        )
                    )
                    * 100.0
                )

            if q_int_scaled_cmm is not None:
                q_int_scaled_disp = convert_flow(q_int_scaled_cmm, "CMM (m³/min)", flow_unit)
                p_int_scaled_pa = k_coef_base * (q_int_scaled_cmm**2)
                p_int_scaled_disp = convert_pressure(p_int_scaled_pa, "Pa", pressure_unit)

                w_int_scaled = np.interp(q_int_scaled_disp, q2_vals, w2_vals)
                db_int_scaled = np.interp(q_int_scaled_disp, q2_vals, db2_vals)
                eff_int_scaled = (
                    float(
                        static_efficiency(
                            q_int_scaled_cmm, p_int_scaled_pa, convert_power(w_int_scaled, power_unit, "W")
                        )
                    )
                    * 100.0
                )

            eff1_q, eff1_curve = efficiency_curve_points(
                q1_base_arr, p1_base_arr, convert_power(w1_vals, power_unit, "W"), [q_int_base_cmm]
            )
            eff2_q, eff2_curve = efficiency_curve_points(
                q2_base_arr, p2_base_arr, convert_power(w2_vals, power_unit, "W"), [q_int_scaled_cmm]
            )

            # Plotly Visualization with Secondary Y-axis (Static Efficiency on Right Axis)
            fig = make_subplots(specs=[[{"secondary_y": True}]])

            # 1. Base PQ (Left Y-axis)
            fig.add_trace(
                go.Scatter(
                    x=q1_vals,
                    y=p1_vals,
                    mode="lines+markers",
                    name="기준 팬 PQ",
                    line=dict(color="#64748B", width=2, dash="dot"),
                ),
                secondary_y=False,
            )
            # 2. Scaled PQ (Left Y-axis)
            fig.add_trace(
                go.Scatter(
                    x=q2_vals,
                    y=p2_vals,
                    mode="lines+markers",
                    name="스케일링 팬 PQ",
                    line=dict(color="#0D9488", width=3),
                ),
                secondary_y=False,
            )

            # 3. Static Efficiency (Right Y-axis)
            fig.add_trace(
                go.Scatter(
                    x=convert_flow(eff1_q, "CMM (m³/min)", flow_unit),
                    y=eff1_curve * 100.0,
                    mode="lines",
                    name="기준 정압효율",
                    line=dict(color="#A7C4BC", width=1.5, dash="dot"),
                ),
                secondary_y=True,
            )
            fig.add_trace(
                go.Scatter(
                    x=convert_flow(eff2_q, "CMM (m³/min)", flow_unit),
                    y=eff2_curve * 100.0,
                    mode="lines",
                    name="예측 정압효율",
                    line=dict(color="#0F766E", width=2),
                ),
                secondary_y=True,
            )

            # 4. System Curve (Left Y-axis)
            if has_system_curve:
                fig.add_trace(
                    go.Scatter(
                        x=q_sys_plot,
                        y=p_sys_plot,
                        mode="lines",
                        name="덕트 계통 저항선",
                        line=dict(color="#2563EB", width=1.5, dash="dash"),
                    ),
                    secondary_y=False,
                )
                # Plot operating point markers
                if q_int_base_cmm is not None:
                    q_int_base_disp = convert_flow(q_int_base_cmm, "CMM (m³/min)", flow_unit)
                    p_int_base_pa = k_coef_base * (q_int_base_cmm**2)
                    p_int_base_disp = convert_pressure(p_int_base_pa, "Pa", pressure_unit)
                    # Pressure marker (Left Axis)
                    fig.add_trace(
                        go.Scatter(
                            x=[q_int_base_disp],
                            y=[p_int_base_disp],
                            mode="markers",
                            marker=dict(color="#4A4744", size=10, symbol="circle"),
                            name="기준 운전점(정압)",
                        ),
                        secondary_y=False,
                    )
                    # Efficiency marker (Right Axis)
                    fig.add_trace(
                        go.Scatter(
                            x=[q_int_base_disp],
                            y=[eff_int_base],
                            mode="markers",
                            marker=dict(color="#7CA380", size=9, symbol="circle-open", line=dict(width=2)),
                            name="기준 운전점(효율)",
                        ),
                        secondary_y=True,
                    )

                if q_int_scaled_cmm is not None:
                    q_int_scaled_disp = convert_flow(q_int_scaled_cmm, "CMM (m³/min)", flow_unit)
                    p_int_scaled_pa = k_coef_base * (q_int_scaled_cmm**2)
                    p_int_scaled_disp = convert_pressure(p_int_scaled_pa, "Pa", pressure_unit)
                    # Pressure marker (Left Axis)
                    fig.add_trace(
                        go.Scatter(
                            x=[q_int_scaled_disp],
                            y=[p_int_scaled_disp],
                            mode="markers",
                            marker=dict(color="#0D9488", size=12, symbol="star"),
                            name="스케일 예측 운전점(정압)",
                        ),
                        secondary_y=False,
                    )
                    # Efficiency marker (Right Axis)
                    fig.add_trace(
                        go.Scatter(
                            x=[q_int_scaled_disp],
                            y=[eff_int_scaled],
                            mode="markers",
                            marker=dict(color="#0F766E", size=11, symbol="diamond-open", line=dict(width=2)),
                            name="스케일 예측 운전점(효율)",
                        ),
                        secondary_y=True,
                    )

            # Calculate Y-axis and X-axis limits (use user custom range if checked)
            y_max_limit = (
                user_y_max
                if (custom_range and user_y_max)
                else float(max(p1_vals.max(), p2_vals.max(), sys_p) * 1.15)
            )
            x_max_limit = (
                user_x_max
                if (custom_range and user_x_max)
                else float(max(q1_vals.max(), q2_vals.max()) * 1.1)
            )

            fig.update_layout(
                title=dict(
                    text="정압 성능 곡선 및 시스템 저항선 비교 (정압효율 이중 축 포함)",
                    font=dict(
                        family="Inter, Pretendard, Malgun Gothic, sans-serif", size=18, color="#0F172A"
                    ),
                ),
                xaxis_title=f"풍량 ({flow_unit_disp})",
                plot_bgcolor="#FFFFFF",
                paper_bgcolor="rgba(0,0,0,0)",
                legend=dict(
                    x=0.02, y=0.98, bordercolor="#E2E8F0", borderwidth=1, bgcolor="rgba(255,255,255,0.7)"
                ),
                margin=dict(l=40, r=40, t=55, b=40),
                height=450,
            )
            fig.update_xaxes(range=[0, x_max_limit], showgrid=True, gridcolor="#EEF2F6", linecolor="#E2E8F0")
            fig.update_yaxes(
                title_text=f"정압 ({pressure_unit_disp})",
                range=[0, y_max_limit],
                showgrid=True,
                gridcolor="#EEF2F6",
                linecolor="#E2E8F0",
                secondary_y=False,
            )
            fig.update_yaxes(title_text="정압효율 (%)", range=[0, 100], showgrid=False, secondary_y=True)
            show_efficiency = st.toggle("정압효율 보조축 표시", value=False, key="pq_show_efficiency")
            for trace in fig.data:
                if trace.yaxis == "y2":
                    trace.visible = show_efficiency
            fig.update_yaxes(visible=show_efficiency, secondary_y=True)
            fig.update_layout(title_text="PQ 성능 · 시스템 운전점")
            if q_int_scaled_cmm is not None:
                fig.add_vline(x=q_int_scaled_disp, line_width=1, line_dash="dot", line_color="#0D9488")
                fig.add_annotation(
                    x=q_int_scaled_disp,
                    y=p_int_scaled_disp,
                    text=f"예측 운전점<br>{q_int_scaled_disp:.2f} {flow_unit_disp} · {p_int_scaled_disp:.2f} {pressure_unit_disp}",
                    showarrow=True,
                    arrowhead=2,
                    ax=35,
                    ay=-50,
                    bgcolor="#ECFDF5",
                    bordercolor="#99F6E4",
                    font=dict(size=11, color="#115E59"),
                )
            _show_chart(fig, width="stretch")

            # Detailed operational metrics comparison
            st.markdown("#### 실 운전점 이동 예측 요약")
            if has_noise:
                col_m1, col_m2, col_m3, col_m4 = st.columns(4)
            else:
                col_m1, col_m2, col_m3 = st.columns(3)

            if q_int_base_cmm is not None and q_int_scaled_cmm is not None:

                # Tab 1 AI Diagnostics
                st.markdown(" ")
                with st.container(border=True, key="panel_2207"):
                    st.markdown("#### 운전점 및 계통 안전성 진단 리포트")

                    # 1. Tip Speed calculation (Target Fan)
                    # d2 is in mm, n2 is in RPM
                    v_tip = (np.pi * d2 * n2) / (60.0 * 1000.0)

                    # 2. Surge/Stall Margin evaluation — 정압 피크점 유량 기준
                    # FC(시로코) 원심팬은 정압 곡선 피크 좌측이 실속/서징 영역이므로,
                    # 운전점 유량이 피크 유량 대비 얼마나 우측(안전)인지로 판정한다.
                    # surge_margin = (Q_op - Q_peak) / Q_peak  [+: 피크 우측(안전), -: 피크 좌측(실속영역)]
                    peak_idx = int(np.argmax(p2_vals))
                    q_peak_cmm = convert_flow(float(q2_vals[peak_idx]), flow_unit, "CMM (m³/min)")
                    if q_peak_cmm > 0:
                        surge_margin = (q_int_scaled_cmm - q_peak_cmm) / q_peak_cmm
                    else:
                        surge_margin = 1.0
                    # 피크가 데이터 첫 점(폐쇄점)이면 곡선 전체가 우하향 → 실속 영역 미포함으로 간주
                    peak_at_shutoff = peak_idx == 0

                    # 3. Efficiency/Cordier match
                    q_int_m3s = q_int_scaled_cmm / 60.0
                    p_int_pa = k_coef_base * (q_int_scaled_cmm**2)
                    y_head = p_int_pa / air_density_2 if air_density_2 > 0 else 0.0
                    omega = n2 * 2.0 * np.pi / 60.0
                    if y_head > 0 and q_int_m3s > 0:
                        sigma = omega * np.sqrt(q_int_m3s) / (y_head**0.75)
                    else:
                        sigma = 0.0

                    # Diagnostic display in 3 columns
                    diag_col1, diag_col2, diag_col3 = st.columns(3)

                    with diag_col1:
                        st.markdown(f"**임펠러 깃 끝 속도 (Tip Speed)**")
                        if v_tip < 35.0:
                            st.success(f"🟢 {v_tip:.1f} m/s (저소음 양호)")
                            st.caption(
                                "날개 끝 선속도가 낮아 기류 마찰 소음이 적고 구조 진동 위험이 매우 낮습니다."
                            )
                        elif v_tip < 55.0:
                            st.warning(f"🟡 {v_tip:.1f} m/s (표준 산업 운전)")
                            st.caption(
                                "일반적인 산업용 운전 범위이나, 실내 가동 시 방음 대책 수립을 권장합니다."
                            )
                        else:
                            st.error(f"🔴 {v_tip:.1f} m/s (초고속 소음 위험)")
                            st.caption(
                                "임펠러 끝 강한 유동 전단력으로 고주파 소음 및 허브 진동 위험이 있어 재설계를 권장합니다."
                            )

                    with diag_col2:
                        st.markdown(f"**서징/실속(Surge·Stall) 위험도**")
                        if peak_at_shutoff:
                            st.success(f"🟢 안전 (곡선 전체 우하향)")
                            st.caption(
                                "정압 곡선에 좌측 상승 구간(실속 영역)이 없어 계통 저항 변동에 대한 안정성이 높습니다."
                            )
                        elif surge_margin >= 0.15:
                            st.success(f"🟢 안전 (피크 대비 +{surge_margin*100:.1f}%)")
                            st.caption(
                                f"운전점이 정압 피크점(유량 {convert_flow(q_peak_cmm, 'CMM (m³/min)', flow_unit):.1f} {flow_unit_disp}) 우측에 충분한 여유를 두고 위치하여 맥동 위험이 낮습니다."
                            )
                        elif surge_margin >= 0.0:
                            st.warning(f"🟡 경계 (피크 대비 +{surge_margin*100:.1f}%)")
                            st.caption(
                                "운전점이 정압 피크점 직우측에 근접해 있습니다. 필터 막힘 등 계통 저항 증가 시 실속 영역으로 밀려날 수 있으니 저항 마진을 확보하십시오."
                            )
                        else:
                            st.error(f"🔴 위험 (피크 좌측 {surge_margin*100:.1f}%)")
                            st.caption(
                                "운전점이 정압 피크 좌측(실속/서징 영역)에 있습니다. 계통 저항 저감, 팬 사이즈 축소 또는 회전수 하향으로 운전점을 피크 우측으로 이동시키십시오."
                            )

                    with diag_col3:
                        st.markdown(f"**유체역학 설계 적합도 (σ)**")
                        if fan_category == CENTRIFUGAL_CATEGORY:
                            lo, hi = 0.15, 1.80
                            cat_disp = "원심팬 계열 (터보/시로코)"
                        elif fan_category == AXIAL_CATEGORY:
                            lo, hi = 1.50, 8.00
                            cat_disp = "축류팬 계열"
                        else:
                            lo, hi = 0.05, 0.35
                            cat_disp = "횡류팬 계열"

                        if sigma == 0.0:
                            st.info("⚪ 계산 대기 중")
                            st.caption("운전점이 식별되지 않아 비속도를 분석할 수 없습니다.")
                        elif lo <= sigma <= hi:
                            st.success(f"🟢 σ = {sigma:.3f} (매칭 양호)")
                            st.caption(
                                f"운전 비속도가 {cat_disp}의 최적 설계 영역({lo}~{hi}) 내에 균형 있게 위치합니다."
                            )
                        else:
                            st.error(f"🔴 σ = {sigma:.3f} (비매칭)")
                            st.caption(
                                f"현재 운전 요구점은 {cat_disp}의 대표 사양을 벗어납니다. 타 기종(예: 원심↔축류) 변경을 권장합니다."
                            )

                dq_pct = safe_pct(q_int_scaled_cmm, q_int_base_cmm)
                dp_pct = safe_pct(p_int_scaled_disp, p_int_base_disp)
                dw_pct = safe_pct(w_int_scaled, w_int_base)
                ddb = db_int_scaled - db_int_base

                q_arrow, q_cls = delta_arrow_class(dq_pct, higher_is_bad=False)
                p_arrow, p_cls = delta_arrow_class(dp_pct, higher_is_bad=False)
                w_arrow, w_cls = delta_arrow_class(dw_pct, higher_is_bad=True)
                db_arrow, db_cls = delta_arrow_class(ddb, higher_is_bad=True)

                with col_m1:
                    st.markdown(
                        f"""
                    <div class="metric-box">
                        <div class="title">운전 풍량</div>
                        <div class="value">{convert_flow(q_int_scaled_cmm, "CMM (m³/min)", flow_unit):.2f} <span style="font-size: 1.0rem; font-weight: normal; color: #64748B;">{flow_unit}</span></div>
                        <div class="delta {q_cls}">{q_arrow} {dq_pct:+.1f}% <span style="font-size: 0.8rem; font-weight: normal; color: #64748B;">(기존: {convert_flow(q_int_base_cmm, "CMM (m³/min)", flow_unit):.2f})</span></div>
                    </div>
                    """,
                        unsafe_allow_html=True,
                    )
                with col_m2:
                    st.markdown(
                        f"""
                    <div class="metric-box">
                        <div class="title">운전 정압</div>
                        <div class="value">{p_int_scaled_disp:.2f} <span style="font-size: 1.0rem; font-weight: normal; color: #64748B;">{pressure_unit}</span></div>
                        <div class="delta {p_cls}">{p_arrow} {dp_pct:+.1f}% <span style="font-size: 0.8rem; font-weight: normal; color: #64748B;">(기존: {p_int_base_disp:.2f})</span></div>
                    </div>
                    """,
                        unsafe_allow_html=True,
                    )
                with col_m3:
                    st.markdown(
                        f"""
                    <div class="metric-box">
                        <div class="title">운전 축동력</div>
                        <div class="value">{w_int_scaled:.1f} <span style="font-size: 1.0rem; font-weight: normal; color: #64748B;">{power_unit}</span></div>
                        <div class="delta {w_cls}">{w_arrow} {dw_pct:+.1f}% <span style="font-size: 0.8rem; font-weight: normal; color: #64748B;">(기존: {w_int_base:.1f})</span></div>
                    </div>
                    """,
                        unsafe_allow_html=True,
                    )
                if has_noise:
                    with col_m4:
                        st.markdown(
                            f"""
                        <div class="metric-box">
                            <div class="title">예측 소음</div>
                            <div class="value">{f'{db_int_scaled:.1f}' if np.isfinite(db_int_scaled) else '—'} <span style="font-size: 1.0rem; font-weight: normal; color: #64748B;">dB(A)</span></div>
                            <div class="delta {db_cls}">{f'{ddb:+.1f} dB(A)' if np.isfinite(ddb) else '목표점 소음 데이터 부족'}</div>
                        </div>
                        """,
                            unsafe_allow_html=True,
                        )
            else:
                st.warning(
                    "선택하신 계통 저항선과 정압 성능 곡선이 만나지 않습니다. 유량 범위 혹은 정압 사양을 검토하십시오."
                )

            # Sub-charts: Power & Noise Curves (make_curve_fig로 통합)
            both_int = q_int_base_cmm is not None and q_int_scaled_cmm is not None
            op_qb = convert_flow(q_int_base_cmm, "CMM (m³/min)", flow_unit) if both_int else None
            op_qs = convert_flow(q_int_scaled_cmm, "CMM (m³/min)", flow_unit) if both_int else None

            if has_noise:
                col_c1, col_c2 = st.columns(2)
                with col_c1:
                    fig_w = make_curve_fig(
                        "풍량 대 축동력 곡선",
                        f"풍량 ({flow_unit_disp})",
                        f"동력 ({power_unit_disp})",
                        q1_vals,
                        w1_vals,
                        q2_vals,
                        w2_vals,
                        "#0D9488",
                        op_q_base=op_qb,
                        op_y_base=w_int_base,
                        op_q_scaled=op_qs,
                        op_y_scaled=w_int_scaled,
                        op_unit=power_unit_disp,
                        height=350,
                    )
                    _show_chart(fig_w, width="stretch")
                with col_c2:
                    fig_db = make_curve_fig(
                        "풍량 대 소음 레벨 곡선",
                        f"풍량 ({flow_unit_disp})",
                        "소음 레벨 dB(A)",
                        q1_vals,
                        db1_vals,
                        q2_vals,
                        db2_vals,
                        "#D97706",
                        op_q_base=op_qb,
                        op_y_base=db_int_base,
                        op_q_scaled=op_qs,
                        op_y_scaled=db_int_scaled,
                        op_unit="dB(A)",
                        height=350,
                    )
                    _show_chart(fig_db, width="stretch")
            else:
                # No noise data, render Power curve only at full width
                fig_w = make_curve_fig(
                    "풍량 대 축동력 곡선",
                    f"풍량 ({flow_unit_disp})",
                    f"동력 ({power_unit_disp})",
                    q1_vals,
                    w1_vals,
                    q2_vals,
                    w2_vals,
                    "#0D9488",
                    op_q_base=op_qb,
                    op_y_base=w_int_base,
                    op_q_scaled=op_qs,
                    op_y_scaled=w_int_scaled,
                    op_unit=power_unit_disp,
                    height=400,
                )
                _show_chart(fig_w, width="stretch")

            # Export Predicted Curves Section
            st.markdown("---")
            st.markdown("##### 📥 예측 성능 곡선 데이터 내보내기")
            st.markdown(
                "상사 예측이 완료된 성능 곡선(Q-P-W-Eff-Lw) 전구간 데이터를 다운로드하고 모니터링 테이블을 확인하십시오."
            )

            # Excel & CSV data preparation
            results_export_df = pd.DataFrame(
                {
                    f"기준 풍량 ({flow_unit_disp})": q1_vals,
                    f"기준 정압 ({pressure_unit_disp})": p1_vals,
                    f"기준 축동력 ({power_unit_disp})": w1_vals,
                    "기준 정압효율 (%)": np.round(eff1_vals, 2),
                    f"예측 풍량 ({flow_unit_disp})": q2_vals,
                    f"예측 정압 ({pressure_unit_disp})": p2_vals,
                    f"예측 축동력 ({power_unit_disp})": w2_vals,
                    "예측 정압효율 (%)": np.round(eff2_vals, 2),
                }
            )
            if has_noise:
                results_export_df["기준 소음 (dB(A))"] = db1_vals
                results_export_df["예측 소음 (dB(A))"] = db2_vals

            with st.expander("예측 결과 데이터 테이블 전체 보기", expanded=False):
                st.dataframe(results_export_df, width="stretch", hide_index=True)

            col_dl1, col_dl2 = st.columns(2)
            with col_dl1:
                # Excel export
                export_excel = excel_bytes(results_export_df, "Scaled_Performance")

                st.download_button(
                    label="📥 예측 데이터 엑셀(.xlsx) 다운로드",
                    data=export_excel,
                    file_name=f"scaled_fan_performance_{datetime.date.today().strftime('%Y%m%d')}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="dl_pq_xlsx",
                    on_click="ignore",
                )
            with col_dl2:
                # CSV export
                csv_bytes = results_export_df.to_csv(index=False).encode("utf-8-sig")
                st.download_button(
                    label="📥 예측 데이터 CSV(.csv) 다운로드",
                    data=csv_bytes,
                    file_name=f"scaled_fan_performance_{datetime.date.today().strftime('%Y%m%d')}.csv",
                    mime="text/csv",
                    key="dl_pq_csv",
                    on_click="ignore",
                )

        except Exception as e:
            st.error(f"예측 도중 에러가 발생했습니다: {str(e)}. 입력된 PQ 데이터를 다시 확인해 주세요.")


# =============================================================================
# ACTIVE SET TAB: 다중 raw data 및 부분 계산 지원 UI
# =============================================================================
def render_tab_set(settings):
    """세트 예측 화면. 전역 입력은 settings로 전달한다."""
    fan_category = settings["fan_category"]
    is_axial = settings["is_axial"]
    dim_label = settings["dim_label"]
    dim_key = settings["dim_key"]
    flow_unit = settings["flow_unit"]
    pressure_unit = settings["pressure_unit"]
    power_unit = settings["power_unit"]
    flow_unit_disp = settings["flow_unit_disp"]
    pressure_unit_disp = settings["pressure_unit_disp"]
    power_unit_disp = settings["power_unit_disp"]
    air_density_1 = settings["air_density_1"]
    air_density_2 = settings["air_density_2"]
    air_viscosity_1 = settings["air_viscosity_1"]
    air_viscosity_2 = settings["air_viscosity_2"]
    use_eff_correction = settings["use_eff_correction"]

    set2_saved_inputs = _read_input_state()
    st.markdown("### 세트 실측 보정 및 팬·부품 변경 예측")
    st.markdown(
        "단품 PQ·효율 데이터와 여러 RPM의 세트 raw data를 결합해 현재 세트 저항을 추정하고, "
        "변경 부품과 신규 팬을 적용한 운전점을 예측합니다. **CMM·Pa·W·dB(A)** 고정 단위입니다."
    )
    st.caption(
        "이 탭의 현재 단품표·세트 raw 표·신규 직접입력표와 주요 형상/시험 조건은 로컬에 "
        "자동 저장되어 브라우저를 새로고침해도 복원됩니다."
    )

    # 계산 결과는 입력 열의 폭에 제한되지 않는 상단 전체 너비에 표시한다.
    with st.container(key="set2_chart_deck"):
        chart_slot = st.empty()
    input_col, result_col = st.columns([1.22, 2.0])

    with input_col:
        with st.container(border=True, key="panel_2456"):
            st.markdown("#### 1. 현재 단품 팬 데이터")
            g1, g2, g3 = st.columns(3)
            with g1:
                set2_cur_d = st.number_input(
                    "현재 D1 (mm)",
                    10.0,
                    value=float(set2_saved_inputs.get("set2_cur_d", 300.0)),
                    step=10.0,
                    key="set2_cur_d",
                    format="%.0f",
                )
            with g2:
                if is_axial:
                    st.text_input("현재 독립 높이", "미적용", disabled=True, key="set2_cur_x_ax")
                    set2_cur_x = set2_cur_d
                else:
                    set2_cur_x = st.number_input(
                        f"현재 {dim_label}1 (mm)",
                        5.0,
                        value=float(set2_saved_inputs.get("set2_cur_x", 100.0)),
                        step=5.0,
                        key="set2_cur_x",
                        format="%.0f",
                    )
            with g3:
                set2_cur_n = st.number_input(
                    "단품 곡선 시험 RPM",
                    10.0,
                    value=float(set2_saved_inputs.get("set2_cur_n", 1500.0)),
                    step=10.0,
                    key="set2_cur_n",
                )

            g4, g5 = st.columns(2)
            with g4:
                set2_cur_count = st.number_input(
                    "현재 팬 수량",
                    1,
                    value=int(set2_saved_inputs.get("set2_cur_count", 1)),
                    step=1,
                    key="set2_cur_count",
                )
            with g5:
                arrangement_options = ["병렬 (Parallel)", "직렬 (Series)"]
                saved_arrangement = set2_saved_inputs.get("set2_cur_arr", arrangement_options[0])
                set2_cur_arr = st.selectbox(
                    "현재 팬 구성",
                    arrangement_options,
                    index=(
                        arrangement_options.index(saved_arrangement)
                        if saved_arrangement in arrangement_options
                        else 0
                    ),
                    key="set2_cur_arr",
                )
            _persist_input_state(
                set2_cur_d=float(set2_cur_d),
                set2_cur_x=float(set2_cur_x),
                set2_cur_n=float(set2_cur_n),
                set2_cur_count=int(set2_cur_count),
                set2_cur_arr=set2_cur_arr,
            )
            st.caption(
                "팬 수량은 입력 PQ가 **팬 1대 기준**일 때만 실제 수량을 입력합니다. PQ 데이터가 이미 팬 2대 조합 결과라면 수량은 1로 두어 중복 합산을 피하십시오."
            )

            default_unit_curve = pd.DataFrame(
                {
                    "풍량 Q (CMM)": [0, 10, 20, 30, 40, 50, 60, 70, 80],
                    "정압 P (Pa)": [220, 215, 200, 180, 150, 110, 70, 30, 0],
                    "축동력 W (W)": [70.0, None, None, None, None, None, None, None, 105.0],
                    "정압효율 ηs (%)": [None, 30.0, 45.0, 55.0, 62.0, 65.0, 60.0, 45.0, None],
                    "소음 Lw (dB(A))": [62, 61, 60, 59, 60, 62, 64, 67, 70],
                }
            )
            default_unit_curve = _state_dataframe(
                set2_saved_inputs,
                "set2_unit_curve",
                default_unit_curve,
                numeric_columns=default_unit_curve.columns,
            )
            set2_unit_df = st.data_editor(
                default_unit_curve,
                num_rows="dynamic",
                hide_index=True,
                width="stretch",
                key="set2_unit_curve",
                column_config=_numeric_editor_column_config(default_unit_curve.columns),
            )
            _persist_input_state(
                set2_unit_curve=_dataframe_records(set2_unit_df, numeric_columns=default_unit_curve.columns)
            )
            st.caption(
                "Q·P만 입력해도 압력곡선은 계산됩니다. 축동력 또는 정압효율 중 하나를 입력하면 나머지를 자동 계산하며, 소음은 선택 입력입니다."
            )

            pe1, pe2, pe3 = st.columns(3)
            with pe1:
                set2_cur_motor_eff = st.slider(
                    "현재 모터 효율", 0.30, 1.00, 0.82, 0.01, key="set2_cur_motor_eff"
                )
            with pe2:
                set2_cur_ctrl_eff = st.slider(
                    "현재 제어기 효율", 0.30, 1.00, 0.95, 0.01, key="set2_cur_ctrl_eff"
                )
            with pe3:
                set2_cur_aux = st.number_input(
                    "팬 외 소비전력 (W)", 0.0, value=30.0, step=5.0, key="set2_cur_aux"
                )

        with st.container(border=True, key="panel_2536"):
            st.markdown("#### 2. 현재 세트 raw data")
            st.caption(
                "RPM과 풍량이 있는 각 행에서 단품 곡선을 해당 RPM으로 환산해 운전 정압을 추정하고, 전체 raw 점에 가장 잘 맞는 세트 저항식 P=K·Qᵐ을 계산합니다. 소전·소음은 비워도 됩니다."
            )
            default_raw = pd.DataFrame(
                {
                    "RPM": [900, 1100, 1300, 1500],
                    "풍량 Q (CMM)": [21.0, 26.0, 31.0, 35.0],
                    "소비전력 (W)": [95.0, 130.0, 180.0, 240.0],
                    "세트 소음 (dB(A))": [56.5, 58.5, 61.0, 63.0],
                }
            )
            default_raw = _state_dataframe(
                set2_saved_inputs,
                "set2_raw_data",
                default_raw,
                numeric_columns=default_raw.columns,
            )
            set2_raw_df = st.data_editor(
                default_raw,
                num_rows="dynamic",
                hide_index=True,
                width="stretch",
                key="set2_raw_data",
                column_config=_numeric_editor_column_config(default_raw.columns),
            )
            _persist_input_state(
                set2_raw_data=_dataframe_records(set2_raw_df, numeric_columns=default_raw.columns)
            )
            set2_allow_extrapolation = st.checkbox(
                "원본 PQ 끝단의 제한 외삽 허용 (풍량 폭의 최대 20%)",
                value=False,
                key="set2_allow_extrapolation",
            )
            st.caption(
                "기본값은 측정 범위 내 계산입니다. 외삽 허용 시 압력에만 적용하고 동력·소음 결측은 유지합니다."
            )
            set2_effective_cur_n_preview = float(set2_cur_n)
            st.info(
                f"직접 입력한 단품 곡선 시험 RPM **{set2_effective_cur_n_preview:.0f} RPM**을 계산에 적용합니다."
            )

        raw_q_preview = pd.to_numeric(
            set2_raw_df.get("풍량 Q (CMM)", pd.Series(dtype=float)), errors="coerce"
        ).dropna()
        raw_q_preview = raw_q_preview[raw_q_preview > 0]
        query_q_default = float(raw_q_preview.max()) if len(raw_q_preview) else 35.0

        with st.container(border=True, key="panel_2572"):
            st.markdown("#### 3. 목표 풍량")
            set2_target_q = st.number_input(
                "목표 세트 풍량 (CMM)",
                0.1,
                value=float(set2_saved_inputs.get("set2_target_q", query_q_default)),
                step=1.0,
                key="set2_target_q",
            )
            st.caption(
                "현재 세트와 변경 세트가 이 풍량을 내는 데 필요한 RPM을 각각 자동 계산합니다. "
                "기준점 압손 모델의 기본 기준 풍량에도 이 값을 사용합니다."
            )
            _persist_input_state(set2_target_q=float(set2_target_q))

        with st.container(border=True, key="panel_2585"):
            st.markdown("#### 4. 세트 부품별 저항 변경")
            st.caption(
                "열교환기·필터·그릴·루버 등 통과면적이 있는 부품을 행으로 추가하십시오. "
                "‘변경’한 부품만 현재 세트 저항에서 기존 압손을 빼고 신규 압손으로 교체합니다."
            )
            set2_link_ref_q = st.checkbox(
                "압손 기준 풍량을 목표 풍량과 연동",
                bool(set2_saved_inputs.get("set2_link_ref_q", True)),
                key="set2_link_ref_q",
            )
            if set2_link_ref_q:
                set2_basic_ref_q = float(set2_target_q)
                st.caption(f"기준 풍량: **{set2_basic_ref_q:.2f} CMM**")
            else:
                set2_basic_ref_q = st.number_input(
                    "선정 기준 풍량 (CMM)",
                    min_value=0.1,
                    value=float(set2_saved_inputs.get("set2_basic_ref_q", set2_target_q)),
                    step=1.0,
                    key="set2_basic_ref_q",
                )

            # 이전의 분리 표 입력을 단일 부품 카드 목록으로 한 번만 마이그레이션한다.
            saved_cards = set2_saved_inputs.get("set2_parts_cards")
            if not isinstance(saved_cards, list) or not saved_cards:
                saved_cards = []
                for row in set2_saved_inputs.get("set2_parts_point", []):
                    saved_cards.append(
                        {
                            "name": row.get("부품명", "부품"),
                            "enabled": row.get("적용", False),
                            "model": RESISTANCE_POINT_MODEL,
                            "old_area": row.get("기존 면적 m²", 0.20),
                            "new_area": row.get("신규 면적 m²", 0.25),
                            "ref_q": set2_saved_inputs.get("set2_basic_ref_q", set2_basic_ref_q),
                            "old_dp": row.get("기존 압손 Pa", 20.0),
                            "exponent": row.get("지수 m", 2.0),
                        }
                    )
                for row in set2_saved_inputs.get("set2_parts_velocity", []):
                    saved_cards.append(
                        {
                            "name": row.get("부품명", "부품"),
                            "enabled": row.get("적용", False),
                            "model": RESISTANCE_VELOCITY_MODEL,
                            "old_area": row.get("기존 면적 m²", 0.20),
                            "new_area": row.get("신규 면적 m²", 0.25),
                            "a1": row.get("기존 a1", 0.0),
                            "a2": row.get("기존 a2", 0.0),
                            "change_coeff": pd.notna(row.get("신규 a1")) or pd.notna(row.get("신규 a2")),
                            "new_a1": row.get("신규 a1"),
                            "new_a2": row.get("신규 a2"),
                        }
                    )
                for row in set2_saved_inputs.get("set2_parts_k", []):
                    saved_cards.append(
                        {
                            "name": row.get("부품명", "부품"),
                            "enabled": row.get("적용", False),
                            "model": RESISTANCE_K_MODEL,
                            "old_area": row.get("기존 면적 m²", 0.20),
                            "new_area": row.get("신규 면적 m²", 0.25),
                            "loss_k": row.get("기존 K", 0.0),
                            "change_coeff": pd.notna(row.get("신규 K")),
                            "new_k": row.get("신규 K"),
                        }
                    )
            if not saved_cards:
                saved_cards = [
                    {
                        "name": "열교환기",
                        "enabled": True,
                        "model": RESISTANCE_POINT_MODEL,
                        "old_area": 1.84,
                        "new_area": 2.0,
                        "ref_q": set2_basic_ref_q,
                        "old_dp": 46.0,
                        "exponent": 2.0,
                    }
                ]

            set2_part_count = st.number_input(
                "저항 변경 부품 수",
                min_value=1,
                max_value=8,
                value=int(set2_saved_inputs.get("set2_part_count", len(saved_cards))),
                step=1,
                key="set2_part_count",
            )
            normalized_parts = []
            saved_card_updates = []
            for card_index in range(int(set2_part_count)):
                base_card = (
                    saved_cards[card_index]
                    if card_index < len(saved_cards)
                    else {
                        "name": f"부품 {card_index+1}",
                        "enabled": False,
                        "model": RESISTANCE_POINT_MODEL,
                        "old_area": 0.20,
                        "new_area": 0.20,
                        "ref_q": set2_basic_ref_q,
                        "old_dp": 0.0,
                        "exponent": 2.0,
                    }
                )
                with st.container(border=True):
                    title_col, apply_col = st.columns([3.0, 1.0])
                    with title_col:
                        part_name = st.text_input(
                            "부품명",
                            value=str(base_card.get("name", f"부품 {card_index+1}")),
                            key=f"set2_card_name_{card_index}",
                        )
                    with apply_col:
                        part_enabled = st.checkbox(
                            "변경 적용",
                            value=bool(base_card.get("enabled", False)),
                            key=f"set2_card_enabled_{card_index}",
                        )
                    saved_model = base_card.get("model", RESISTANCE_POINT_MODEL)
                    model_options = [RESISTANCE_POINT_MODEL, RESISTANCE_VELOCITY_MODEL, RESISTANCE_K_MODEL]
                    part_model = st.radio(
                        "압손 입력 방식",
                        model_options,
                        index=model_options.index(saved_model) if saved_model in model_options else 0,
                        horizontal=True,
                        key=f"set2_card_model_{card_index}",
                    )
                    area1_col, area2_col = st.columns(2)
                    with area1_col:
                        old_area = st.number_input(
                            "기존 통과면적 (m²)",
                            min_value=0.000001,
                            value=float(base_card.get("old_area", 0.20)),
                            format="%.4f",
                            key=f"set2_card_old_area_{card_index}",
                        )
                    with area2_col:
                        new_area = st.number_input(
                            "신규 통과면적 (m²)",
                            min_value=0.000001,
                            value=float(base_card.get("new_area", old_area)),
                            format="%.4f",
                            key=f"set2_card_new_area_{card_index}",
                        )

                    part_record = {
                        "부품명": part_name,
                        "변경": part_enabled,
                        "입력 모델": part_model,
                        "기존 면적 m²": old_area,
                        "신규 면적 m²": new_area,
                    }
                    saved_record = {
                        **base_card,
                        "name": part_name,
                        "enabled": part_enabled,
                        "model": part_model,
                        "old_area": old_area,
                        "new_area": new_area,
                    }
                    if part_model == RESISTANCE_POINT_MODEL:
                        if set2_link_ref_q:
                            ref_q = float(set2_target_q)
                            st.caption(f"압손 기준 풍량: 목표 풍량 {ref_q:.2f} CMM")
                        else:
                            ref_q = float(set2_basic_ref_q)
                            st.caption(f"압손 기준 풍량: 선정 풍량 {ref_q:.2f} CMM")
                        dp_col, exponent_col = st.columns(2)
                        with dp_col:
                            old_dp = st.number_input(
                                "기준 풍량의 기존 압손 (Pa)",
                                min_value=0.0,
                                value=float(base_card.get("old_dp", 0.0)),
                                key=f"set2_card_dp_{card_index}",
                            )
                        with exponent_col:
                            exponent = st.number_input(
                                "압손지수 m",
                                min_value=0.5,
                                max_value=3.0,
                                value=float(base_card.get("exponent", 2.0)),
                                step=0.1,
                                help="모르면 일반적인 제곱 저항인 2.0을 사용합니다.",
                                key=f"set2_card_exp_{card_index}",
                            )
                        part_record.update(
                            {
                                "기준 풍량 CMM": ref_q,
                                "기준 압손 Pa": old_dp,
                                "압손지수 m": exponent,
                            }
                        )
                        saved_record.update({"ref_q": ref_q, "old_dp": old_dp, "exponent": exponent})
                    elif part_model == RESISTANCE_VELOCITY_MODEL:
                        a1_col, a2_col = st.columns(2)
                        with a1_col:
                            a1 = st.number_input(
                                "기존 a₁ [Pa/(m/s)]",
                                min_value=0.0,
                                value=float(base_card.get("a1", 0.0)),
                                key=f"set2_card_a1_{card_index}",
                            )
                        with a2_col:
                            a2 = st.number_input(
                                "기존 a₂ [Pa/(m/s)²]",
                                min_value=0.0,
                                value=float(base_card.get("a2", 0.0)),
                                key=f"set2_card_a2_{card_index}",
                            )
                        change_coeff = st.checkbox(
                            "신규 부품의 a₁·a₂가 기존과 다름",
                            value=bool(base_card.get("change_coeff", False)),
                            key=f"set2_card_change_a_{card_index}",
                        )
                        new_a1 = new_a2 = None
                        if change_coeff:
                            na1_col, na2_col = st.columns(2)
                            with na1_col:
                                new_a1 = st.number_input(
                                    "신규 a₁",
                                    min_value=0.0,
                                    value=float(base_card.get("new_a1") or a1),
                                    key=f"set2_card_new_a1_{card_index}",
                                )
                            with na2_col:
                                new_a2 = st.number_input(
                                    "신규 a₂",
                                    min_value=0.0,
                                    value=float(base_card.get("new_a2") or a2),
                                    key=f"set2_card_new_a2_{card_index}",
                                )
                        part_record.update(
                            {
                                "1차계수 a1": a1,
                                "2차계수 a2": a2,
                                "신규 1차계수 a1": new_a1,
                                "신규 2차계수 a2": new_a2,
                            }
                        )
                        saved_record.update(
                            {
                                "a1": a1,
                                "a2": a2,
                                "change_coeff": change_coeff,
                                "new_a1": new_a1,
                                "new_a2": new_a2,
                            }
                        )
                    else:
                        loss_k = st.number_input(
                            "기존 손실계수 K",
                            min_value=0.0,
                            value=float(base_card.get("loss_k", 0.0)),
                            key=f"set2_card_k_{card_index}",
                        )
                        change_coeff = st.checkbox(
                            "신규 부품의 K가 기존과 다름",
                            value=bool(base_card.get("change_coeff", False)),
                            key=f"set2_card_change_k_{card_index}",
                        )
                        new_k = None
                        if change_coeff:
                            new_k = st.number_input(
                                "신규 손실계수 K",
                                min_value=0.0,
                                value=float(base_card.get("new_k") or loss_k),
                                key=f"set2_card_new_k_{card_index}",
                            )
                        part_record.update({"손실계수 K": loss_k, "신규 손실계수 K": new_k})
                        saved_record.update(
                            {
                                "loss_k": loss_k,
                                "change_coeff": change_coeff,
                                "new_k": new_k,
                            }
                        )

                    if part_enabled:
                        preview_old, preview_new = component_pressure_by_model(
                            np.array([set2_target_q]), part_record, air_density_1, air_density_2
                        )
                        old_velocity = (set2_target_q / 60.0) / old_area
                        new_velocity = (set2_target_q / 60.0) / new_area
                        st.info(
                            f"목표 풍량 {set2_target_q:.2f} CMM · 면풍속 "
                            f"{old_velocity:.2f}→{new_velocity:.2f} m/s · 압손 "
                            f"{float(preview_old[0]):.2f}→{float(preview_new[0]):.2f} Pa"
                        )
                    normalized_parts.append(part_record)
                    saved_card_updates.append(saved_record)

            with st.expander("공통 고급 설정", expanded=False):
                set2_remainder_factor = st.slider(
                    "입력하지 않은 나머지 세트 저항비",
                    0.50,
                    1.50,
                    float(set2_saved_inputs.get("set2_remainder_factor", 1.00)),
                    0.01,
                    key="set2_remainder",
                )
                set2_uncertainty = st.slider(
                    "신규 세트 저항 불확실성 (±%)",
                    0,
                    40,
                    int(set2_saved_inputs.get("set2_uncertainty", 10)),
                    1,
                    key="set2_uncertainty",
                )
            set2_parts_df = pd.DataFrame(normalized_parts)
            _persist_input_state(
                set2_link_ref_q=bool(set2_link_ref_q),
                set2_basic_ref_q=float(set2_basic_ref_q),
                set2_part_count=int(set2_part_count),
                set2_parts_cards=saved_card_updates,
                set2_remainder_factor=float(set2_remainder_factor),
                set2_uncertainty=int(set2_uncertainty),
            )

        with st.container(border=True, key="panel_2833"):
            st.markdown("#### 5. 신규 팬 조건")
            new_mode_options = ["현재 팬에서 상사 예측", "신규 단품 곡선 직접 입력"]
            saved_new_mode = set2_saved_inputs.get("set2_new_mode", new_mode_options[0])
            set2_new_mode = st.radio(
                "신규 팬 입력 방식",
                new_mode_options,
                index=(new_mode_options.index(saved_new_mode) if saved_new_mode in new_mode_options else 0),
                horizontal=True,
                key="set2_new_mode",
            )
            _persist_input_state(set2_new_mode=set2_new_mode)
            set2_same_fan = set2_new_mode == "현재 팬에서 상사 예측" and st.checkbox(
                "현재 팬 그대로 사용: 직경·높이/길이·수량·구성 동일", value=False, key="set2_same_fan"
            )
            if set2_new_mode == "현재 팬에서 상사 예측" and not set2_same_fan:
                set2_new_change_scope = st.radio(
                    "신규 팬 변경 범위",
                    ["직경·높이/길이 유지", "직경·높이/길이 변경"],
                    horizontal=True,
                    key="set2_new_change_scope",
                )
            else:
                set2_new_change_scope = None
            if set2_same_fan:
                set2_new_d = set2_cur_d
                set2_new_x = set2_cur_x
                set2_new_direct_df = None
                st.info(
                    f"동일 팬 조건: D={set2_new_d:.1f} mm, 수량={set2_cur_count}, 구성={set2_cur_arr}. "
                    "RPM은 목표 풍량과 변경 저항으로 다시 계산합니다."
                )
            elif set2_new_change_scope == "직경·높이/길이 유지":
                set2_new_d = set2_cur_d
                set2_new_x = set2_cur_x
                set2_new_direct_df = None
                st.caption(
                    f"형상은 D={set2_new_d:.1f} mm, {dim_label}={set2_new_x:.1f} mm로 현재 팬과 동일하게 고정됩니다."
                )
            elif set2_new_mode == "현재 팬에서 상사 예측":
                n1c, n2c = st.columns(2)
                with n1c:
                    set2_new_d = st.number_input(
                        "신규 직경 D2 (mm)",
                        10.0,
                        value=float(set2_cur_d),
                        step=10.0,
                        key="set2_new_d_v3",
                        format="%.0f",
                    )
                with n2c:
                    if is_axial:
                        st.text_input("신규 독립 높이", "미적용", disabled=True, key="set2_new_x_ax_v3")
                        set2_new_x = set2_new_d
                    else:
                        set2_new_x = st.number_input(
                            f"신규 {dim_label}2 (mm)",
                            5.0,
                            value=float(set2_cur_x),
                            step=5.0,
                            key="set2_new_x_v3",
                            format="%.0f",
                        )
                set2_new_direct_df = None
            else:
                nd1, nd2, nd3 = st.columns(3)
                with nd1:
                    set2_new_d = st.number_input(
                        "신규 실제 직경 D2 (mm)",
                        10.0,
                        value=float(set2_saved_inputs.get("set2_new_direct_d", set2_cur_d)),
                        step=10.0,
                        key="set2_new_direct_d",
                        format="%.0f",
                    )
                with nd2:
                    if is_axial:
                        st.text_input("신규 독립 높이", "미적용", disabled=True, key="set2_new_direct_x_ax")
                        set2_new_x = set2_new_d
                    else:
                        set2_new_x = st.number_input(
                            f"신규 실제 {dim_label}2 (mm)",
                            5.0,
                            value=float(set2_saved_inputs.get("set2_new_direct_x", set2_cur_x)),
                            step=5.0,
                            key="set2_new_direct_x",
                            format="%.0f",
                        )
                with nd3:
                    set2_new_curve_test_n = st.number_input(
                        "신규 단품 곡선 시험 RPM",
                        10.0,
                        value=float(
                            set2_saved_inputs.get("set2_new_curve_test_n", set2_effective_cur_n_preview)
                        ),
                        step=10.0,
                        key="set2_new_curve_test_n",
                    )
                default_new_curve2 = pd.DataFrame(
                    {
                        "풍량 Q (CMM)": [0, 12, 24, 36, 48, 60, 72, 84, 96],
                        "정압 P (Pa)": [260, 255, 238, 210, 175, 130, 82, 35, 0],
                        "축동력 W (W)": [80.0, None, None, None, None, None, None, None, 115.0],
                        "정압효율 ηs (%)": [None, 32.0, 48.0, 58.0, 65.0, 68.0, 62.0, 45.0, None],
                        # 신규 단품 소음은 선택 입력이다. 예시 숫자를 넣으면 실제 측정값으로
                        # 오인되므로 기본값은 반드시 비워 둔다.
                        "소음 Lw (dB(A))": [None] * 9,
                    }
                )
                default_new_curve2 = _state_dataframe(
                    set2_saved_inputs,
                    "set2_new_direct_curve",
                    default_new_curve2,
                    numeric_columns=default_new_curve2.columns,
                )
                set2_new_direct_df = st.data_editor(
                    default_new_curve2,
                    num_rows="dynamic",
                    hide_index=True,
                    width="stretch",
                    key="set2_new_direct",
                    column_config=_numeric_editor_column_config(default_new_curve2.columns),
                )
                _persist_input_state(
                    set2_new_direct_d=float(set2_new_d),
                    set2_new_direct_x=float(set2_new_x),
                    set2_new_curve_test_n=float(set2_new_curve_test_n),
                    set2_new_direct_curve=_dataframe_records(
                        set2_new_direct_df,
                        numeric_columns=default_new_curve2.columns,
                    ),
                )
                st.caption(
                    "입력 PQ는 위 시험 RPM의 신규 팬 1대 기준곡선으로 사용합니다. 목표 RPM에서는 "
                    "Q∝N, P∝N², W∝N³으로 이동하며 신규 직경비를 PQ에 다시 중복 적용하지 않습니다. "
                    "소음은 실제 신규 단품 측정값이 있을 때만 입력하고, 비워 두면 현재 세트의 "
                    "RPM–소음 회귀와 팬 종류별 형상 상사로 예측합니다."
                )

            if set2_same_fan:
                set2_new_count = set2_cur_count
                set2_new_arr = set2_cur_arr
            else:
                set2_change_fan_config = st.checkbox(
                    "신규 팬 수량·직렬/병렬 구성도 변경",
                    bool(set2_saved_inputs.get("set2_change_fan_config", False)),
                    key="set2_change_fan_config",
                )
                if set2_change_fan_config:
                    ns1, ns2 = st.columns(2)
                    with ns1:
                        set2_new_count = st.number_input(
                            "신규 팬 수량",
                            1,
                            value=int(set2_saved_inputs.get("set2_new_count", set2_cur_count)),
                            step=1,
                            key="set2_new_count_v3",
                        )
                    with ns2:
                        new_arrangement_options = ["병렬 (Parallel)", "직렬 (Series)"]
                        saved_new_arrangement = set2_saved_inputs.get("set2_new_arr", set2_cur_arr)
                        set2_new_arr = st.selectbox(
                            "신규 팬 구성",
                            new_arrangement_options,
                            index=(
                                new_arrangement_options.index(saved_new_arrangement)
                                if saved_new_arrangement in new_arrangement_options
                                else 0
                            ),
                            key="set2_new_arr_v3",
                        )
                else:
                    set2_new_count = set2_cur_count
                    set2_new_arr = set2_cur_arr
                    st.caption(
                        f"신규 팬 수량·구성은 현재와 동일하게 자동 적용: "
                        f"{set2_new_count}대, {set2_new_arr}"
                    )
                _persist_input_state(
                    set2_change_fan_config=bool(set2_change_fan_config),
                    set2_new_count=int(set2_new_count),
                    set2_new_arr=set2_new_arr,
                )
            if set2_same_fan:
                set2_new_motor_eff = set2_cur_motor_eff
                set2_new_ctrl_eff = set2_cur_ctrl_eff
                set2_new_aux = set2_cur_aux
                set2_apply_power_cal = True
                _persist_input_state(set2_reuse_current_drive=True)
                set2_acoustic_delta = 0.0
                st.caption(
                    "동일 조건 검증은 raw 결과값을 신규 결과에 복사하지 않습니다. 현재 모터·제어기 효율과 "
                    "팬 외 전력을 적용하고, raw에서 산출한 RPM별 보정계수와 팬 상사식으로 다시 계산합니다."
                )
            else:
                reuse_drive_default = bool(
                    set2_saved_inputs.get(
                        "set2_reuse_current_drive",
                        st.session_state.get("set2_apply_power_cal", False),
                    )
                )
                set2_reuse_current_drive = st.checkbox(
                    "현재 모터·제어기 및 전력 보정조건 그대로 사용",
                    reuse_drive_default,
                    key="set2_reuse_current_drive",
                )
                _persist_input_state(set2_reuse_current_drive=bool(set2_reuse_current_drive))
                if set2_reuse_current_drive:
                    set2_new_motor_eff = set2_cur_motor_eff
                    set2_new_ctrl_eff = set2_cur_ctrl_eff
                    set2_new_aux = set2_cur_aux
                    set2_apply_power_cal = True
                    ne1, ne2, ne3 = st.columns(3)
                    ne1.metric("승계 모터 효율", f"{set2_new_motor_eff:.2f}")
                    ne2.metric("승계 제어기 효율", f"{set2_new_ctrl_eff:.2f}")
                    ne3.metric("승계 팬 외 전력", f"{set2_new_aux:.2f} W")
                    st.caption(
                        "현재 모터·제어기 효율과 팬 외 전력을 그대로 적용하고, 현재 세트 raw에서 "
                        "산출한 전력 보정계수도 신규 조건에 전달합니다."
                    )
                else:
                    ne1, ne2, ne3 = st.columns(3)
                    with ne1:
                        set2_new_motor_eff = st.slider(
                            "신규 모터 효율", 0.30, 1.00, 0.88, 0.01, key="set2_new_motor_eff"
                        )
                    with ne2:
                        set2_new_ctrl_eff = st.slider(
                            "신규 제어기 효율", 0.30, 1.00, 0.96, 0.01, key="set2_new_ctrl_eff"
                        )
                    with ne3:
                        set2_new_aux = st.number_input(
                            "신규 팬 외 전력 (W)", 0.0, value=30.0, step=5.0, key="set2_new_aux"
                        )
                    set2_apply_power_cal = False
                    st.caption(
                        "신규 모터·제어기 효율과 팬 외 전력을 독립 입력하며, 현재 raw 전력 "
                        "보정계수는 신규 조건에 전달하지 않습니다."
                    )
                set2_acoustic_delta = 0.0
                st.caption(
                    "소음에는 특정 실험 결과를 맞추기 위한 고정 잔차 보정을 적용하지 않습니다. "
                    "세트 raw RPM–소음 회귀와 팬 종류별 형상 상사만 사용합니다."
                )

    with result_col:
        unit_curve, made_w_count, made_eta_count = _prepare_optional_fan_curve(set2_unit_df)
        if len(unit_curve) < 2:
            st.info("단품 팬의 풍량 Q와 정압 P를 서로 다른 2개 지점 이상 입력하면 분석이 시작됩니다.")
        else:
            if made_w_count or made_eta_count:
                st.success(
                    f"입력 데이터에서 축동력 {made_w_count}개, 정압효율 {made_eta_count}개를 자동 보완했습니다."
                )
            with st.expander("자동 보완된 현재 단품 데이터", expanded=False):
                st.dataframe(
                    unit_curve.style.format(precision=2, na_rep="—"), width="stretch", hide_index=True
                )

            try:
                uq = unit_curve["풍량 Q (CMM)"].to_numpy(float)
                up = unit_curve["정압 P (Pa)"].to_numpy(float)
                uw = unit_curve["축동력 W (W)"].to_numpy(float)
                ulw = unit_curve["소음 Lw (dB(A))"].to_numpy(float)
                ueta = unit_curve["정압효율 ηs (%)"].to_numpy(float)

                raw = set2_raw_df.copy()
                raw_cols = ["RPM", "풍량 Q (CMM)", "소비전력 (W)", "세트 소음 (dB(A))"]
                for c in raw_cols:
                    if c not in raw.columns:
                        raw[c] = np.nan
                    raw[c] = pd.to_numeric(raw[c], errors="coerce")

                inferred_rows = []
                excluded_raw_rows = []
                for pos, row in raw.iterrows():
                    rpm = row["RPM"]
                    q_meas = row["풍량 Q (CMM)"]
                    if not (np.isfinite(rpm) and rpm > 0 and np.isfinite(q_meas) and q_meas > 0):
                        continue
                    sq, sp, sw, slw, _ = run_similarity_scaling(
                        uq,
                        up,
                        uw,
                        ulw,
                        set2_cur_d,
                        set2_cur_x,
                        set2_effective_cur_n_preview,
                        set2_cur_d,
                        set2_cur_x,
                        rpm,
                        air_density_1,
                        air_density_1,
                        use_eff_correction,
                        np.nanmax(ueta) / 100.0 if np.any(np.isfinite(ueta)) else 0.60,
                        fan_category,
                        mu1=air_viscosity_1,
                        mu2=air_viscosity_1,
                    )
                    cq, cp, cw, clw = combine_identical_fans(sq, sp, sw, slw, set2_cur_count, set2_cur_arr)
                    order = np.argsort(cq)
                    cq, cp, cw, clw = cq[order], cp[order], cw[order], clw[order]
                    if not cq.min() <= q_meas <= cq.max():
                        excluded_raw_rows.append(str(pos + 1))
                        continue
                    p_at_q = _optional_interp(q_meas, cq, cp)
                    if p_at_q is None or p_at_q <= 0:
                        continue
                    w_at_q = _optional_interp(q_meas, cq, cw)
                    lw_at_q = _optional_interp(q_meas, cq, clw)
                    input_model = (
                        w_at_q / max(set2_cur_motor_eff * set2_cur_ctrl_eff, 1e-9)
                        if w_at_q is not None
                        else None
                    )
                    measured_power = row["소비전력 (W)"] if np.isfinite(row["소비전력 (W)"]) else None
                    measured_noise = (
                        row["세트 소음 (dB(A))"] if np.isfinite(row["세트 소음 (dB(A))"]) else None
                    )
                    power_factor = (
                        (measured_power - set2_cur_aux) / input_model
                        if measured_power is not None and input_model and measured_power > set2_cur_aux
                        else None
                    )
                    acoustic_offset = (
                        measured_noise - lw_at_q
                        if measured_noise is not None and lw_at_q is not None
                        else None
                    )
                    measured_fan_input = (
                        (measured_power - set2_cur_aux)
                        if measured_power is not None and measured_power > set2_cur_aux
                        else None
                    )
                    set_static_efficiency = (
                        (q_meas / 60.0) * p_at_q / measured_fan_input
                        if measured_fan_input and measured_fan_input > 0
                        else None
                    )
                    if set2_cur_arr == "직렬 (Series)":
                        q_per_fan = q_meas
                        p_per_fan = p_at_q / max(float(set2_cur_count), 1.0)
                    else:
                        q_per_fan = q_meas / max(float(set2_cur_count), 1.0)
                        p_per_fan = p_at_q
                    acoustic_kn = (
                        measured_noise
                        - 10.0 * np.log10(max(float(set2_cur_count), 1.0))
                        - 10.0 * np.log10(q_per_fan)
                        - 20.0 * np.log10(p_per_fan)
                        if measured_noise is not None and q_per_fan > 0 and p_per_fan > 0
                        else None
                    )
                    inferred_rows.append(
                        {
                            "원본행": int(pos),
                            "RPM": rpm,
                            "실측 풍량 (CMM)": q_meas,
                            "팬곡선 추정 정압 (Pa)": p_at_q,
                            "Q² 기준 K(참고)": p_at_q / q_meas**2,
                            "팬 축동력 모델 (W)": w_at_q,
                            "실측 소비전력 (W)": measured_power,
                            "전력 보정계수": power_factor,
                            "세트 총정압효율": set_static_efficiency,
                            "팬 합성 Lw": lw_at_q,
                            "실측 세트 소음": measured_noise,
                            "음향 보정 (dB)": acoustic_offset,
                            "음향 특성계수 KN": acoustic_kn,
                            "곡선범위내": bool(cq.min() <= q_meas <= cq.max()),
                        }
                    )
                inferred = pd.DataFrame(inferred_rows)
                if excluded_raw_rows:
                    st.warning(
                        "원본 PQ 범위 밖이어서 저항·전력·음향 보정에서 제외한 raw 행: "
                        + ", ".join(excluded_raw_rows)
                    )

                st.markdown("#### 목표 풍량 자동 RPM 비교")
                if inferred.empty:
                    st.info(
                        "세트 raw data에 RPM과 풍량을 입력하면 목표 풍량의 현재·신규 필요 RPM을 계산합니다."
                    )
                else:
                    q_fit = inferred["실측 풍량 (CMM)"].to_numpy()
                    y_fit = inferred["팬곡선 추정 정압 (Pa)"].to_numpy()
                    current_k, system_exponent, fit_rmse_pct = fit_system_resistance(q_fit, y_fit)

                    def _coefficient_at_rpm(column_name, target_rpm):
                        candidates = inferred.dropna(subset=[column_name, "RPM"])
                        if candidates.empty or target_rpm is None:
                            return None, None
                        by_rpm = (
                            candidates.groupby("RPM", as_index=False)[column_name].mean().sort_values("RPM")
                        )
                        rpms = by_rpm["RPM"].to_numpy(float)
                        values = by_rpm[column_name].to_numpy(float)
                        value = float(np.interp(float(target_rpm), rpms, values))
                        if np.any(np.isclose(rpms, target_rpm, rtol=0.0, atol=1e-6)):
                            source = f"{target_rpm:.0f} RPM 실측점"
                        elif rpms.min() < target_rpm < rpms.max():
                            source = f"RPM별 계수 보간({target_rpm:.0f} RPM)"
                        else:
                            anchor_rpm = rpms[0] if target_rpm <= rpms.min() else rpms[-1]
                            source = f"{anchor_rpm:.0f} RPM 경계계수"
                        return value, source

                    with st.expander("세트 보정 데이터와 저항식", expanded=False):
                        st.caption(
                            f"현재 세트 저항식: P={current_k:.5g}·Q^{system_exponent:.3f}, "
                            f"raw 적합 RMSE={fit_rmse_pct:.2f}%"
                        )
                        st.dataframe(
                            inferred.drop(columns=["원본행"]).style.format(precision=3, na_rep="—"),
                            width="stretch",
                            hide_index=True,
                        )
                    if (~inferred["곡선범위내"]).any():
                        st.warning(
                            "일부 raw 운전점이 해당 RPM의 팬 곡선 범위를 벗어났습니다. "
                            "이 구간의 저항·보정계수 신뢰도가 낮을 수 있습니다."
                        )

                    parts2 = set2_parts_df.copy()
                    numeric_part_cols = [
                        "기준 풍량 CMM",
                        "기준 압손 Pa",
                        "기존 면적 m²",
                        "신규 면적 m²",
                        "압손지수 m",
                        "1차계수 a1",
                        "2차계수 a2",
                        "손실계수 K",
                        "신규 1차계수 a1",
                        "신규 2차계수 a2",
                        "신규 손실계수 K",
                    ]
                    for column in numeric_part_cols:
                        if column not in parts2.columns:
                            parts2[column] = np.nan
                        parts2[column] = pd.to_numeric(parts2.get(column), errors="coerce")
                    if "입력 모델" not in parts2.columns:
                        parts2["입력 모델"] = RESISTANCE_POINT_MODEL
                    parts2["입력 모델"] = parts2["입력 모델"].fillna(RESISTANCE_POINT_MODEL)
                    parts2["변경"] = parts2.get("변경", False).fillna(False).astype(bool)
                    requested_parts = parts2.loc[parts2["변경"]].copy()
                    area_valid = (requested_parts["기존 면적 m²"] > 0) & (requested_parts["신규 면적 m²"] > 0)
                    point_valid = (
                        (requested_parts["입력 모델"] == RESISTANCE_POINT_MODEL)
                        & (requested_parts["기준 풍량 CMM"] > 0)
                        & (requested_parts["기준 압손 Pa"] >= 0)
                    )
                    velocity_valid = (requested_parts["입력 모델"] == RESISTANCE_VELOCITY_MODEL) & (
                        (requested_parts["1차계수 a1"] > 0) | (requested_parts["2차계수 a2"] > 0)
                    )
                    k_valid = (requested_parts["입력 모델"] == RESISTANCE_K_MODEL) & (
                        requested_parts["손실계수 K"] > 0
                    )
                    changed2 = requested_parts[area_valid & (point_valid | velocity_valid | k_valid)]
                    if len(changed2) < len(requested_parts):
                        st.warning(
                            f"변경 체크한 부품 중 {len(requested_parts)-len(changed2)}개는 면적 또는 "
                            "선택 모델의 필수 압손·계수가 없어 저항 변경에서 제외했습니다."
                        )
                    target_q = float(set2_target_q)
                    component_scale = 1.0
                    old_target_sum = 0.0
                    for _, part in changed2.iterrows():
                        old_part, _ = component_pressure_by_model(
                            np.array([target_q]), part, air_density_1, air_density_2
                        )
                        old_target_sum += float(old_part[0])
                    total_target_pressure = current_k * target_q**system_exponent
                    if old_target_sum > total_target_pressure and old_target_sum > 0:
                        component_scale = total_target_pressure / old_target_sum
                        st.warning(
                            f"목표 풍량에서 변경 부품의 기존 압손 합({old_target_sum:.1f} Pa)이 "
                            f"세트 전체 압손({total_target_pressure:.1f} Pa)을 초과하여 부품 기여도를 "
                            f"{component_scale:.3f}배로 제한했습니다."
                        )

                    def _system2(q_values):
                        q_values = np.asarray(q_values, dtype=float)
                        current_total = current_k * q_values**system_exponent
                        old_sum = np.zeros_like(q_values)
                        new_sum = np.zeros_like(q_values)
                        for _, part in changed2.iterrows():
                            old_part, new_part = component_pressure_by_model(
                                q_values, part, air_density_1, air_density_2
                            )
                            old_sum += component_scale * old_part
                            new_sum += component_scale * new_part
                        remainder = np.maximum(current_total - old_sum, 0.0) * float(set2_remainder_factor)
                        return current_total, remainder + new_sum, old_sum, new_sum

                    current_target_p = float(current_k * target_q**system_exponent)
                    _, new_target_p_arr, _, _ = _system2(np.array([target_q]))
                    new_target_p = float(new_target_p_arr[0])
                    original_eff = float(np.nanmax(ueta) / 100.0) if np.any(np.isfinite(ueta)) else 0.60

                    raw_by_q = raw.dropna(subset=["풍량 Q (CMM)", "RPM"]).sort_values("풍량 Q (CMM)")
                    preferred_current_rpm = _optional_interp(
                        target_q, raw_by_q["풍량 Q (CMM)"], raw_by_q["RPM"]
                    )
                    current_required_rpm, current_rpm_details = _solve_required_rpm(
                        target_q,
                        current_target_p,
                        uq,
                        up,
                        uw,
                        ulw,
                        set2_cur_d,
                        set2_cur_x,
                        set2_effective_cur_n_preview,
                        set2_cur_d,
                        set2_cur_x,
                        air_density_1,
                        air_density_1,
                        set2_cur_count,
                        set2_cur_arr,
                        fan_category,
                        use_eff_correction,
                        original_eff,
                        preferred_current_rpm,
                        allow_limited_extrapolation=set2_allow_extrapolation,
                        return_details=True,
                        mu1=air_viscosity_1,
                        mu2=air_viscosity_1,
                    )

                    if set2_new_mode == "현재 팬에서 상사 예측":
                        nbq, nbp, nbw, nblw = uq, up, uw, ulw
                        new_base_d, new_base_x = set2_cur_d, set2_cur_x
                        new_base_rpm = set2_effective_cur_n_preview
                        new_original_eff = original_eff
                    else:
                        new_completed, _, _ = _prepare_optional_fan_curve(set2_new_direct_df)
                        if len(new_completed) >= 2:
                            nbq = new_completed["풍량 Q (CMM)"].to_numpy(float)
                            nbp = new_completed["정압 P (Pa)"].to_numpy(float)
                            nbw = new_completed["축동력 W (W)"].to_numpy(float)
                            nblw = new_completed["소음 Lw (dB(A))"].to_numpy(float)
                            new_eta_values = new_completed["정압효율 ηs (%)"].to_numpy(float)
                            new_original_eff = (
                                float(np.nanmax(new_eta_values) / 100.0)
                                if np.any(np.isfinite(new_eta_values))
                                else original_eff
                            )
                        else:
                            nbq = nbp = nbw = nblw = np.array([])
                            new_original_eff = original_eff
                        new_base_d, new_base_x = set2_new_d, set2_new_x
                        new_base_rpm = float(set2_new_curve_test_n)

                    if set2_new_mode == "현재 팬에서 상사 예측" and current_required_rpm is not None:
                        current_series_factor = (
                            float(set2_cur_count) if set2_cur_arr == "직렬 (Series)" else 1.0
                        )
                        new_series_factor = float(set2_new_count) if set2_new_arr == "직렬 (Series)" else 1.0
                        preferred_new_rpm = (
                            float(current_required_rpm)
                            * float(set2_cur_d)
                            / float(set2_new_d)
                            * np.sqrt(
                                max(new_target_p, 1e-9)
                                / max(current_target_p, 1e-9)
                                * float(air_density_1)
                                / float(air_density_2)
                                * current_series_factor
                                / max(new_series_factor, 1e-9)
                            )
                        )
                    else:
                        preferred_new_rpm = new_base_rpm

                    new_solution = (
                        _solve_required_rpm(
                            target_q,
                            new_target_p,
                            nbq,
                            nbp,
                            nbw,
                            nblw,
                            new_base_d,
                            new_base_x,
                            new_base_rpm,
                            set2_new_d,
                            set2_new_x,
                            air_density_1,
                            air_density_2,
                            set2_new_count,
                            set2_new_arr,
                            fan_category,
                            use_eff_correction,
                            new_original_eff,
                            preferred_new_rpm,
                            allow_limited_extrapolation=set2_allow_extrapolation,
                            return_details=True,
                            mu1=air_viscosity_1,
                            mu2=air_viscosity_2,
                        )
                        if len(nbq) >= 2
                        else (None, {"method": "신규 PQ 입력 부족"})
                    )
                    new_required_rpm, new_rpm_details = new_solution

                    # 실측 보정은 현재 목표 운전점에서 한 번만 고정한다. 신규 RPM에서
                    # 다시 보간하면 RPM 변화와 보정계수 변화가 중첩되어 대구경·저속화
                    # 효과를 왜곡할 수 있다.
                    calibration_rpm = current_required_rpm

                    raw_noise_model = fit_set_noise_log_rpm(
                        raw["RPM"].to_numpy(float),
                        raw["세트 소음 (dB(A))"].to_numpy(float),
                    )
                    geometry_noise_delta = fan_geometry_noise_delta(
                        set2_cur_d,
                        set2_cur_x,
                        set2_new_d,
                        set2_new_x,
                        air_density_1,
                        air_density_2,
                        fan_category,
                    )
                    if raw_noise_model is not None:
                        geometry_formula = (
                            "70log(D₂/D₁)+20log(ρ₂/ρ₁)"
                            if is_axial
                            else f"60log(D₂/D₁)+10log({dim_key}₂/{dim_key}₁)+20log(ρ₂/ρ₁)"
                        )
                        st.caption(
                            f"세트 소음 회귀: L={raw_noise_model['intercept']:.2f}"
                            f"{raw_noise_model['slope']:+.2f}·log₁₀(N), "
                            f"R²={raw_noise_model['r_squared']:.4f}, "
                            f"RMSE={raw_noise_model['rmse']:.2f} dB. "
                            f"신규 팬 형상 보정: {geometry_formula}. "
                            "별도 실험 잔차 보정은 적용하지 않습니다."
                        )

                    def _predict_power(
                        model_w, q_value, p_value, rpm, motor_eff, ctrl_eff, aux_power, use_raw_factor, is_new
                    ):
                        factor, factor_source = _coefficient_at_rpm("전력 보정계수", calibration_rpm)
                        if model_w is not None:
                            applied_factor = factor if use_raw_factor and factor is not None else 1.0
                            value = model_w / max(motor_eff * ctrl_eff, 1e-9) * applied_factor + aux_power
                            source = (
                                f"단품 축동력 × 현재 목표점 고정계수({factor_source})"
                                if use_raw_factor and factor is not None
                                else "단품 축동력·전기효율"
                            )
                            return float(value), source

                        overall_eff, eff_source = _coefficient_at_rpm("세트 총정압효율", calibration_rpm)
                        if overall_eff is None or overall_eff <= 0:
                            return None, "계산 불가"
                        if is_new:
                            electrical_ratio = (
                                motor_eff * ctrl_eff / max(set2_cur_motor_eff * set2_cur_ctrl_eff, 1e-9)
                            )
                            overall_eff *= electrical_ratio
                        value = (q_value / 60.0) * p_value / max(overall_eff, 1e-9) + aux_power
                        return float(value), f"현재 목표점 {eff_source} 총정압효율(Q·P/η)"

                    def _predict_noise(
                        model_lw, q_value, p_value, rpm, count, arrangement, acoustic_delta, is_new=False
                    ):
                        offset, offset_source = _coefficient_at_rpm("음향 보정 (dB)", calibration_rpm)
                        if model_lw is not None:
                            # 신규 단품곡선을 직접 입력한 경우 그 Lw는 신규 팬의 실제
                            # 소음 데이터다. 현재 팬 단품↔세트 차이로 구한 offset을 더하면
                            # 서로 다른 팬의 시험 조건을 중복 보정하므로 전달하지 않는다.
                            direct_new_lw = is_new and set2_new_mode == "신규 단품 곡선 직접 입력"
                            applied_offset = 0.0 if direct_new_lw else (offset if offset is not None else 0.0)
                            if direct_new_lw:
                                source = "신규 단품 실측 Lw·RPM 상사(현재 팬 음향보정 미전달)"
                            else:
                                source = (
                                    f"단품 Lw + 현재 목표점 고정 음향보정({offset_source})"
                                    if offset is not None
                                    else "단품 Lw 상사식(세트 보정 없음)"
                                )
                            return float(model_lw + applied_offset + acoustic_delta), source

                        # 단품 소음이 없으면 현재 세트의 raw RPM–소음 경향을 직접 사용한다.
                        # 신규 팬은 PQ·저항 교점에서 구한 RPM을 회귀식에 넣고, RPM 회귀에
                        # 포함되지 않은 형상·밀도·팬 수량 변화만 팬 종류별로 더한다.
                        # 특정 직경 실험값을 맞추는 고정 잔차 보정은 사용하지 않는다.
                        rpm_noise = predict_set_noise_from_rpm(raw_noise_model, rpm)
                        can_transfer_raw_model = True
                        if rpm_noise is not None and can_transfer_raw_model:
                            geometry_delta = (
                                geometry_noise_delta if is_new and geometry_noise_delta is not None else 0.0
                            )
                            count_delta = (
                                10.0 * np.log10(max(float(count), 1.0) / max(float(set2_cur_count), 1.0))
                                if is_new
                                else 0.0
                            )
                            value = rpm_noise + geometry_delta + count_delta
                            if is_new:
                                shape_term = "70logD" if is_axial else f"60logD+10log{dim_key}"
                                source = f"세트 raw log(RPM) 회귀 + {shape_term} 형상상사" + (
                                    " + 팬 수량 보정" if abs(count_delta) > 1e-9 else ""
                                )
                            else:
                                source = "세트 raw log(RPM) 회귀"
                            return float(value), source

                        kn, kn_source = _coefficient_at_rpm("음향 특성계수 KN", calibration_rpm)
                        if kn is None or q_value <= 0 or p_value <= 0:
                            return None, "계산 불가"
                        count_safe = max(float(count), 1.0)
                        if arrangement == "직렬 (Series)":
                            q_per_fan = q_value
                            p_per_fan = p_value / count_safe
                        else:
                            q_per_fan = q_value / count_safe
                            p_per_fan = p_value
                        value = (
                            kn
                            + 10.0 * np.log10(q_per_fan)
                            + 20.0 * np.log10(p_per_fan)
                            + 10.0 * np.log10(count_safe)
                            + acoustic_delta
                        )
                        source_prefix = "신규" if is_new else "현재"
                        return float(value), (
                            f"{source_prefix} 운전점·현재 목표점 고정 KN({kn_source})·Q·P식"
                        )

                    if current_required_rpm is None or new_required_rpm is None:
                        missing = []
                        if current_required_rpm is None:
                            missing.append("현재 팬")
                        if new_required_rpm is None:
                            missing.append("신규 팬")
                        st.warning(
                            f"{', '.join(missing)}에서 목표 풍량·필요정압을 만족하는 RPM을 "
                            "선택한 PQ 범위/외삽 설정에서 찾지 못했습니다. "
                            "단품/세트 팬 수량, 직렬·병렬 구성, 시험 RPM과 PQ 단위를 먼저 확인하십시오."
                        )
                        failed_details = new_rpm_details if new_required_rpm is None else current_rpm_details
                        if failed_details.get("reference_q_min") is not None:
                            st.caption(
                                f"시험 RPM 환산 세트 PQ 풍량 범위: "
                                f"{failed_details['reference_q_min']:.2f}~"
                                f"{failed_details['reference_q_max']:.2f} CMM, "
                                f"제한 외삽 허용범위: {failed_details['limited_q_min']:.2f}~"
                                f"{failed_details['limited_q_max']:.2f} CMM. "
                                "입력 PQ가 팬 1대 기준이면 세트 팬 수량을 실제 수량으로 설정해야 합니다."
                            )
                    else:
                        for label, details in [("현재", current_rpm_details), ("신규", new_rpm_details)]:
                            if details.get("used_extrapolation"):
                                st.warning(
                                    f"{label} 팬: 원본 PQ 풍량 폭의 {details['extrapolation_pct']:.1f}% 외삽. "
                                    f"시험 RPM 환산 풍량 {details['equivalent_q']:.2f} CMM. 추가 측정 권장."
                                )
                            else:
                                st.caption(
                                    f"{label} 팬 · 원본 PQ 범위 내 · 시험 RPM 환산 풍량 {details['equivalent_q']:.2f} CMM"
                                )
                        cq, cp, cw, clw = _fan_set_curve_at_rpm(
                            uq,
                            up,
                            uw,
                            ulw,
                            set2_cur_d,
                            set2_cur_x,
                            set2_effective_cur_n_preview,
                            set2_cur_d,
                            set2_cur_x,
                            current_required_rpm,
                            air_density_1,
                            air_density_1,
                            set2_cur_count,
                            set2_cur_arr,
                            fan_category,
                            use_eff_correction,
                            original_eff,
                            mu1=air_viscosity_1,
                            mu2=air_viscosity_1,
                        )
                        nq, np_, nw, nlw = _fan_set_curve_at_rpm(
                            nbq,
                            nbp,
                            nbw,
                            nblw,
                            new_base_d,
                            new_base_x,
                            new_base_rpm,
                            set2_new_d,
                            set2_new_x,
                            new_required_rpm,
                            air_density_1,
                            air_density_2,
                            set2_new_count,
                            set2_new_arr,
                            fan_category,
                            use_eff_correction,
                            new_original_eff,
                            mu1=air_viscosity_1,
                            mu2=air_viscosity_2,
                        )
                        current_fan_p = current_target_p
                        current_model_w = _optional_interp(target_q, cq, cw)
                        current_model_lw = _optional_interp(target_q, cq, clw)
                        new_fan_p = new_target_p
                        new_model_w = _optional_interp(target_q, nq, nw)
                        new_model_lw = _optional_interp(target_q, nq, nlw)

                        current_power, current_power_source = _predict_power(
                            current_model_w,
                            target_q,
                            current_target_p,
                            current_required_rpm,
                            set2_cur_motor_eff,
                            set2_cur_ctrl_eff,
                            set2_cur_aux,
                            True,
                            False,
                        )
                        current_noise, current_noise_source = _predict_noise(
                            current_model_lw,
                            target_q,
                            current_target_p,
                            current_required_rpm,
                            set2_cur_count,
                            set2_cur_arr,
                            0.0,
                        )
                        new_power, new_power_source = _predict_power(
                            new_model_w,
                            target_q,
                            new_target_p,
                            new_required_rpm,
                            set2_new_motor_eff,
                            set2_new_ctrl_eff,
                            set2_new_aux,
                            set2_apply_power_cal,
                            True,
                        )
                        new_noise, new_noise_source = _predict_noise(
                            new_model_lw,
                            target_q,
                            new_target_p,
                            new_required_rpm,
                            set2_new_count,
                            set2_new_arr,
                            set2_acoustic_delta,
                            True,
                        )

                        if set2_new_mode == "현재 팬에서 상사 예측":
                            r_d = set2_new_d / set2_cur_d
                            r_n = new_required_rpm / set2_effective_cur_n_preview
                            r_x = r_d if is_axial else set2_new_x / set2_cur_x
                            r_rho = air_density_2 / air_density_1
                            st.caption(
                                f"신규 상사비: D₂/D₁={r_d:.3f}, N₂/N시험={r_n:.3f}, "
                                f"{'D 기하상사비' if is_axial else '높이/길이비'}={r_x:.3f} → "
                                f"Q×{r_n*r_d**2*r_x:.3f}, "
                                f"P×{r_n**2*r_d**2*r_rho:.3f}, "
                                f"W×{r_n**3*r_d**4*r_x*r_rho:.3f}"
                            )
                            if not (0.80 <= r_d <= 1.20 and 0.80 <= r_x <= 1.20):
                                st.warning(
                                    "직경 또는 높이/길이 변화가 기준 대비 ±20%를 넘습니다. "
                                    "상사법칙은 계산되지만 형상·효율·톤소음 변화가 커질 수 있으므로 "
                                    "신규 단품시험이나 CFD로 확인하는 것이 좋습니다."
                                )

                        raw_rpm_min = float(inferred["RPM"].min())
                        raw_rpm_max = float(inferred["RPM"].max())
                        if not (raw_rpm_min <= current_required_rpm <= raw_rpm_max):
                            st.warning(
                                f"현재 팬 필요 RPM {current_required_rpm:.0f}은 raw 보정 범위 "
                                f"{raw_rpm_min:.0f}~{raw_rpm_max:.0f} RPM 밖입니다. "
                                "전력 계수는 가장 가까운 경계값을 사용하고, 소음은 log(RPM) 회귀를 "
                                "외삽하므로 불확실성이 커집니다."
                            )
                        if not (raw_rpm_min <= new_required_rpm <= raw_rpm_max):
                            st.warning(
                                f"신규 필요 RPM {new_required_rpm:.0f}은 현재 raw 운전 범위 "
                                f"{raw_rpm_min:.0f}~{raw_rpm_max:.0f} RPM 밖입니다. "
                                "전력 보정계수는 현재 목표점 값으로 고정되고, 소음 회귀는 외삽되므로 "
                                "팬 상사 예측의 불확실성이 커집니다."
                            )
                        if new_model_w is None:
                            st.caption(
                                "신규 단품 W/정압효율이 없어 Q·P와 raw 총정압효율로 소비전력을 "
                                "유추했습니다. 신규 팬 자체 효율이 달라지면 오차가 커질 수 있습니다."
                            )
                        if new_model_lw is None:
                            if raw_noise_model is not None:
                                st.caption(
                                    "신규 단품 Lw가 없어 신규 PQ·변경 저항의 교점 RPM을 현재 세트 raw "
                                    "log(RPM)–소음 회귀식에 대입하고, 팬 종류별 형상 상사를 더했습니다. "
                                    "축류팬은 70logD, 원심·횡류팬은 60logD+10logX를 적용하며 특정 실험을 "
                                    "맞추는 고정 잔차 보정은 사용하지 않습니다. 신규 직접입력 팬은 현재 팬과 "
                                    "동일 계열의 RPM–소음 경향을 갖는다는 가정입니다."
                                )
                            else:
                                st.caption(
                                    "신규 단품 Lw와 전달 가능한 RPM–소음 회귀가 없어 현재 세트 raw의 "
                                    "KN·Q·P식을 보조 모델로 사용했습니다. 신규 팬 단품 소음 데이터가 있으면 "
                                    "그 데이터를 우선 사용합니다."
                                )

                        comparison = pd.DataFrame(
                            [
                                {
                                    "구분": "현재 팬·현재 저항",
                                    "자동 계산 RPM": current_required_rpm,
                                    "목표 풍량 (CMM)": target_q,
                                    "팬 정압 (Pa)": current_fan_p,
                                    "세트 필요정압 (Pa)": current_target_p,
                                    "소비전력 (W)": current_power,
                                    "세트 소음 (dB(A))": current_noise,
                                    "전력 산출": current_power_source,
                                    "소음 산출": current_noise_source,
                                    "PQ 계산 근거": current_rpm_details["method"],
                                    "원본 범위 외삽 (%)": current_rpm_details["extrapolation_pct"],
                                },
                                {
                                    "구분": "신규 팬·변경 저항",
                                    "자동 계산 RPM": new_required_rpm,
                                    "목표 풍량 (CMM)": target_q,
                                    "팬 정압 (Pa)": new_fan_p,
                                    "세트 필요정압 (Pa)": new_target_p,
                                    "소비전력 (W)": new_power,
                                    "세트 소음 (dB(A))": new_noise,
                                    "전력 산출": new_power_source,
                                    "소음 산출": new_noise_source,
                                    "PQ 계산 근거": new_rpm_details["method"],
                                    "원본 범위 외삽 (%)": new_rpm_details["extrapolation_pct"],
                                },
                            ]
                        )
                        st.markdown("##### 교체 후 예상 성능")
                        kpi_columns = st.columns(4)
                        for column, label, new_value, old_value, unit in zip(
                            kpi_columns,
                            ["필요 RPM", "필요 정압", "소비전력", "세트 소음"],
                            [new_required_rpm, new_target_p, new_power, new_noise],
                            [current_required_rpm, current_target_p, current_power, current_noise],
                            ["RPM", "Pa", "W", "dB(A)"],
                        ):
                            with column:
                                display = "—" if new_value is None else f"{new_value:,.1f}"
                                delta = (
                                    None
                                    if new_value is None or old_value is None
                                    else f"{new_value-old_value:+.1f} {unit}"
                                )
                                st.metric(
                                    label,
                                    display,
                                    delta,
                                    delta_color="inverse",
                                    help=f"현재 팬 대비 변화 · {unit}",
                                )
                        with st.expander("현재·신규 수치 비교 · 산출 근거", expanded=False):
                            st.dataframe(
                                comparison.style.format(precision=2, na_rep="—"),
                                width="stretch",
                                hide_index=True,
                            )
                        with st.expander("현재·신규 팬 소비전력·소음 계산식", expanded=False):
                            st.markdown("### 현재 팬")
                            current_formula_r_n = float(current_required_rpm) / max(
                                float(set2_effective_cur_n_preview), 1e-9
                            )
                            if is_axial:
                                current_shaft_formula = (
                                    "W_shaft,cur = W_shaft,test · (Ncur/Ntest)³ "
                                    "· (Dcur/Dtest)⁵ · (ρcur/ρtest)"
                                )
                            else:
                                current_shaft_formula = (
                                    f"W_shaft,cur = W_shaft,test · (Ncur/Ntest)³ "
                                    f"· (Dcur/Dtest)⁴ · ({dim_key}cur/{dim_key}test) · (ρcur/ρtest)"
                                )
                            st.markdown("**소비전력**")
                            st.code(current_shaft_formula, language=None)
                            st.caption(
                                f"단품곡선 시험 {set2_effective_cur_n_preview:.2f} RPM → "
                                f"현재 운전 {current_required_rpm:.2f} RPM · "
                                f"동일 형상·밀도 축동력 상사배율={current_formula_r_n**3:.4f}"
                            )
                            current_display_factor, current_factor_source = _coefficient_at_rpm(
                                "전력 보정계수", calibration_rpm
                            )
                            if current_model_w is not None:
                                current_applied_factor = (
                                    current_display_factor if current_display_factor is not None else 1.0
                                )
                                st.latex(
                                    r"P_{in,cur}=\frac{W_{shaft,cur}}{\eta_{motor,cur}\eta_{ctrl,cur}}"
                                    r"\,C_{raw}+P_{aux,cur}"
                                )
                                st.markdown(
                                    f"수치 대입: `{current_model_w:.2f} / "
                                    f"({set2_cur_motor_eff:.3f}×{set2_cur_ctrl_eff:.3f}) "
                                    f"× {current_applied_factor:.4f} + {set2_cur_aux:.2f} "
                                    f"= {current_power:.2f} W`"
                                )
                                if current_display_factor is not None:
                                    st.caption(f"C_raw 출처: {current_factor_source}")
                            else:
                                current_display_eff, current_eff_source = _coefficient_at_rpm(
                                    "세트 총정압효율", calibration_rpm
                                )
                                if current_display_eff is not None and current_display_eff > 0:
                                    st.latex(
                                        r"P_{in,cur}=\frac{(Q_{cur}/60)\Delta P_{cur}}"
                                        r"{\eta_{set,cur}}+P_{aux,cur}"
                                    )
                                    st.markdown(
                                        f"수치 대입: `({target_q:.2f}/60)×{current_target_p:.2f} "
                                        f"/ {current_display_eff:.4f} + {set2_cur_aux:.2f} "
                                        f"= {current_power:.2f} W`"
                                    )
                                    st.caption(f"세트 종합효율 출처: {current_eff_source}")

                            st.markdown("**소음**")
                            if current_model_lw is None and raw_noise_model is not None:
                                current_rpm_noise = predict_set_noise_from_rpm(
                                    raw_noise_model, current_required_rpm
                                )
                                st.latex(r"L_{set,cur}=a+b\log_{10}(N_{cur})")
                                st.markdown(
                                    f"회귀식: `L={raw_noise_model['intercept']:.3f} "
                                    f"{raw_noise_model['slope']:+.3f}·log10(N)`"
                                )
                                st.markdown(
                                    f"수치 대입: `{raw_noise_model['intercept']:.3f} "
                                    f"{raw_noise_model['slope']:+.3f}×log10({current_required_rpm:.2f}) "
                                    f"= {current_rpm_noise:.2f} dB(A)`"
                                )
                            elif current_model_lw is not None:
                                current_offset, current_offset_source = _coefficient_at_rpm(
                                    "음향 보정 (dB)", calibration_rpm
                                )
                                current_applied_offset = current_offset if current_offset is not None else 0.0
                                st.latex(r"L_{set,cur}=L_{fan,cur}+\Delta L_{set}")
                                st.markdown(
                                    f"수치 대입: `{current_model_lw:.2f} "
                                    f"{current_applied_offset:+.2f} = {current_noise:.2f} dB(A)`"
                                )
                                if current_offset is not None:
                                    st.caption(f"세트 음향보정 출처: {current_offset_source}")
                            else:
                                st.latex(
                                    r"L_{set}=K_N+10\log_{10}(Q_{fan})+20\log_{10}(P_{fan})"
                                    r"+10\log_{10}(n_{fan})"
                                )
                                st.caption(f"보조 KN 모델 적용 · {current_noise_source}")

                            st.divider()
                            st.markdown("### 신규 팬")
                            formula_r_n = float(new_required_rpm) / max(float(new_base_rpm), 1e-9)
                            formula_r_d = float(set2_new_d) / max(float(new_base_d), 1e-9)
                            formula_r_x = (
                                formula_r_d if is_axial else float(set2_new_x) / max(float(new_base_x), 1e-9)
                            )
                            formula_r_rho = float(air_density_2) / max(float(air_density_1), 1e-9)
                            if is_axial:
                                shaft_formula = (
                                    "W_shaft,2 = W_shaft,ref · (N₂/Nref)³ · " "(D₂/Dref)⁵ · (ρ₂/ρref)"
                                )
                                shaft_multiplier = formula_r_n**3 * formula_r_d**5 * formula_r_rho
                            else:
                                shaft_formula = (
                                    f"W_shaft,2 = W_shaft,ref · (N₂/Nref)³ · "
                                    f"(D₂/Dref)⁴ · ({dim_key}₂/{dim_key}ref) · (ρ₂/ρref)"
                                )
                                shaft_multiplier = (
                                    formula_r_n**3 * formula_r_d**4 * formula_r_x * formula_r_rho
                                )
                            st.markdown("**소비전력**")
                            st.code(shaft_formula, language=None)
                            st.caption(
                                f"곡선 기준: {new_base_rpm:.2f} RPM, D={new_base_d:.2f} mm → "
                                f"신규: {new_required_rpm:.2f} RPM, D={set2_new_d:.2f} mm · "
                                f"축동력 상사배율={shaft_multiplier:.4f}"
                            )
                            if new_model_w is not None:
                                display_power_factor, display_factor_source = _coefficient_at_rpm(
                                    "전력 보정계수", calibration_rpm
                                )
                                applied_power_factor = (
                                    display_power_factor
                                    if set2_apply_power_cal and display_power_factor is not None
                                    else 1.0
                                )
                                st.latex(
                                    r"P_{in,2}=\frac{W_{shaft,2}}{\eta_{motor,2}\eta_{ctrl,2}}"
                                    r"\,C_{raw}+P_{aux,2}"
                                )
                                st.markdown(
                                    f"수치 대입: `{new_model_w:.2f} / "
                                    f"({set2_new_motor_eff:.3f}×{set2_new_ctrl_eff:.3f}) "
                                    f"× {applied_power_factor:.4f} + {set2_new_aux:.2f} "
                                    f"= {new_power:.2f} W`"
                                )
                                if set2_apply_power_cal and display_power_factor is not None:
                                    st.caption(f"C_raw 출처: {display_factor_source}")
                                elif not set2_apply_power_cal:
                                    st.caption("현재 raw 전력 보정계수 미전달: C_raw=1.0")
                            else:
                                display_eff, display_eff_source = _coefficient_at_rpm(
                                    "세트 총정압효율", calibration_rpm
                                )
                                if display_eff is not None and display_eff > 0:
                                    adjusted_eff = float(display_eff)
                                    adjusted_eff *= (
                                        set2_new_motor_eff
                                        * set2_new_ctrl_eff
                                        / max(set2_cur_motor_eff * set2_cur_ctrl_eff, 1e-9)
                                    )
                                    st.latex(r"P_{in,2}=\frac{(Q_2/60)\Delta P_2}{\eta_{set,2}}+P_{aux,2}")
                                    st.markdown(
                                        f"수치 대입: `({target_q:.2f}/60)×{new_target_p:.2f} "
                                        f"/ {adjusted_eff:.4f} + {set2_new_aux:.2f} "
                                        f"= {new_power:.2f} W`"
                                    )
                                    st.caption(f"세트 종합효율 출처: {display_eff_source}")

                            st.markdown("**소음**")
                            count_noise_delta = 10.0 * np.log10(
                                max(float(set2_new_count), 1.0) / max(float(set2_cur_count), 1.0)
                            )
                            if new_model_lw is None and raw_noise_model is not None:
                                rpm_noise_value = predict_set_noise_from_rpm(
                                    raw_noise_model, new_required_rpm
                                )
                                shape_noise_delta = (
                                    float(geometry_noise_delta) if geometry_noise_delta is not None else 0.0
                                )
                                st.latex(
                                    r"L_{set,2}=a+b\log_{10}(N_2)+\Delta L_{geometry}"
                                    r"+10\log_{10}(n_2/n_1)"
                                )
                                st.markdown(
                                    f"회귀식: `L={raw_noise_model['intercept']:.3f} "
                                    f"{raw_noise_model['slope']:+.3f}·log10(N)`"
                                )
                                st.markdown(
                                    f"수치 대입: `{rpm_noise_value:.2f} "
                                    f"{shape_noise_delta:+.2f} {count_noise_delta:+.2f} "
                                    f"= {new_noise:.2f} dB(A)`"
                                )
                                if is_axial:
                                    st.caption("형상 보정: 70log₁₀(D₂/D₁)+20log₁₀(ρ₂/ρ₁)")
                                else:
                                    st.caption(
                                        f"형상 보정: 60log₁₀(D₂/D₁)+10log₁₀({dim_key}₂/{dim_key}₁)"
                                        "+20log₁₀(ρ₂/ρ₁)"
                                    )
                            elif new_model_lw is not None:
                                if is_axial:
                                    noise_scale_delta = (
                                        50.0 * np.log10(formula_r_n)
                                        + 70.0 * np.log10(formula_r_d)
                                        + 20.0 * np.log10(formula_r_rho)
                                    )
                                    st.latex(
                                        r"L_2=L_{ref}+50\log_{10}(N_2/N_{ref})"
                                        r"+70\log_{10}(D_2/D_{ref})+20\log_{10}(\rho_2/\rho_{ref})"
                                    )
                                else:
                                    noise_scale_delta = (
                                        50.0 * np.log10(formula_r_n)
                                        + 60.0 * np.log10(formula_r_d)
                                        + 10.0 * np.log10(formula_r_x)
                                        + 20.0 * np.log10(formula_r_rho)
                                    )
                                    st.latex(
                                        r"L_2=L_{ref}+50\log_{10}(N_2/N_{ref})"
                                        r"+60\log_{10}(D_2/D_{ref})+10\log_{10}(X_2/X_{ref})"
                                        r"+20\log_{10}(\rho_2/\rho_{ref})"
                                    )
                                st.markdown(
                                    f"상사 보정량: `{noise_scale_delta:+.2f} dB` · "
                                    f"목표점 예측: `{new_noise:.2f} dB(A)`"
                                )
                                if set2_new_mode == "신규 단품 곡선 직접 입력":
                                    st.caption(
                                        "신규 단품 실측 Lw를 사용하므로 현재 팬의 세트 음향보정은 전달하지 않습니다."
                                    )
                            else:
                                st.latex(
                                    r"L_{set}=K_N+10\log_{10}(Q_{fan})+20\log_{10}(P_{fan})"
                                    r"+10\log_{10}(n_{fan})"
                                )
                                st.caption(f"보조 KN 모델 적용 · {new_noise_source}")
                            if (
                                current_power is not None
                                and new_power is not None
                                and current_noise is not None
                                and new_noise is not None
                            ):
                                st.divider()
                                st.markdown(
                                    f"**최종 변화:** 소비전력 `{current_power:.2f} → {new_power:.2f} W` "
                                    f"(**{new_power-current_power:+.2f} W**), 소음 "
                                    f"`{current_noise:.2f} → {new_noise:.2f} dB(A)` "
                                    f"(**{new_noise-current_noise:+.2f} dB**)"
                                )
                        if set2_apply_power_cal:
                            st.caption(
                                "전력: 현재 목표점의 raw 보정계수를 신규에도 전달했습니다. 따라서 "
                                "현재 모터·제어기 효율 가정이 신규 소전에 영향을 줍니다. "
                                "소음: 전력 보정과 독립적으로 raw RPM–소음 회귀와 팬 종류별 형상 상사를 적용했습니다."
                            )
                        else:
                            st.caption(
                                "전력: 현재 raw 보정은 현재 팬에만 사용하고, 신규 소전은 신규 축동력과 "
                                "신규 모터·제어기 효율로 독립 계산했습니다. 현재 효율 입력은 신규 소전에 "
                                "전달되지 않습니다. 소음은 raw RPM–소음 회귀와 팬 종류별 형상 상사로 독립 계산했습니다."
                            )

                        current_system_on_curve = current_k * cq**system_exponent
                        _, new_system_on_curve, _, _ = _system2(nq)

                        def _build_noise_curve(
                            q_values,
                            p_values,
                            lw_values,
                            rpm,
                            count,
                            arrangement,
                            acoustic_delta,
                            is_new=False,
                        ):
                            if raw_noise_model is not None and not np.any(np.isfinite(lw_values)):
                                # RPM 회귀는 시스템 운전점 소음 모델이지, 고정 RPM의 PQ 전 영역
                                # 소음곡선이 아니다. 수평선으로 오해되지 않도록 목표점만 표시한다.
                                return np.full_like(np.asarray(q_values, dtype=float), np.nan)
                            results = []
                            for q_point, p_point, lw_point in zip(q_values, p_values, lw_values):
                                model_lw_point = float(lw_point) if np.isfinite(lw_point) else None
                                noise_point, _ = _predict_noise(
                                    model_lw_point,
                                    float(q_point),
                                    float(p_point),
                                    rpm,
                                    count,
                                    arrangement,
                                    acoustic_delta,
                                    is_new,
                                )
                                results.append(np.nan if noise_point is None else noise_point)
                            return np.asarray(results, dtype=float)

                        current_noise_on_curve = _build_noise_curve(
                            cq, cp, clw, current_required_rpm, set2_cur_count, set2_cur_arr, 0.0
                        )
                        new_noise_on_curve = _build_noise_curve(
                            nq,
                            np_,
                            nlw,
                            new_required_rpm,
                            set2_new_count,
                            set2_new_arr,
                            set2_acoustic_delta,
                            True,
                        )

                        fig_compare = make_subplots(
                            rows=2,
                            cols=1,
                            row_heights=[0.58, 0.42],
                            vertical_spacing=0.14,
                            subplot_titles=("PQ · 세트 저항", "소음 비교"),
                        )
                        fig_compare.add_trace(
                            go.Scatter(
                                x=cq,
                                y=cp,
                                mode="lines",
                                name=f"현재 팬 @ {current_required_rpm:.0f} RPM",
                                line=dict(color="#64748B", width=3),
                            ),
                            row=1,
                            col=1,
                        )
                        fig_compare.add_trace(
                            go.Scatter(
                                x=cq,
                                y=current_system_on_curve,
                                mode="lines",
                                name="현재 세트 저항",
                                line=dict(color="#64748B", width=2, dash="dot"),
                            ),
                            row=1,
                            col=1,
                        )
                        fig_compare.add_trace(
                            go.Scatter(
                                x=nq,
                                y=np_,
                                mode="lines",
                                name=f"신규 팬 @ {new_required_rpm:.0f} RPM",
                                line=dict(color="#0D9488", width=3),
                            ),
                            row=1,
                            col=1,
                        )
                        fig_compare.add_trace(
                            go.Scatter(
                                x=nq,
                                y=new_system_on_curve,
                                mode="lines",
                                name="신규 세트 저항",
                                line=dict(color="#0D9488", width=2, dash="dot"),
                            ),
                            row=1,
                            col=1,
                        )
                        fig_compare.add_trace(
                            go.Scatter(
                                x=[target_q],
                                y=[current_target_p],
                                mode="markers",
                                showlegend=False,
                                name="현재 목표 운전점",
                                marker=dict(
                                    size=12,
                                    color="#64748B",
                                    symbol="circle",
                                    line=dict(color="white", width=2),
                                ),
                            ),
                            row=1,
                            col=1,
                        )
                        fig_compare.add_trace(
                            go.Scatter(
                                x=[target_q],
                                y=[new_target_p],
                                mode="markers",
                                showlegend=False,
                                name="신규 목표 운전점",
                                marker=dict(
                                    size=14,
                                    color="#0D9488",
                                    symbol="diamond",
                                    line=dict(color="white", width=2),
                                ),
                            ),
                            row=1,
                            col=1,
                        )
                        fig_compare.add_trace(
                            go.Scatter(
                                x=cq,
                                y=current_noise_on_curve,
                                mode="lines",
                                name="현재 세트 소음",
                                showlegend=False,
                                line=dict(color="#64748B", width=2.5),
                            ),
                            row=2,
                            col=1,
                        )
                        fig_compare.add_trace(
                            go.Scatter(
                                x=nq,
                                y=new_noise_on_curve,
                                mode="lines",
                                name="신규 세트 소음",
                                showlegend=False,
                                line=dict(color="#0D9488", width=2.5),
                            ),
                            row=2,
                            col=1,
                        )
                        raw_noise_points = raw.dropna(subset=["풍량 Q (CMM)", "세트 소음 (dB(A))"])
                        raw_noise_points = raw_noise_points[
                            (raw_noise_points["풍량 Q (CMM)"] > 0)
                            & np.isfinite(raw_noise_points["세트 소음 (dB(A))"])
                        ]
                        if len(raw_noise_points):
                            fig_compare.add_trace(
                                go.Scatter(
                                    x=raw_noise_points["풍량 Q (CMM)"],
                                    y=raw_noise_points["세트 소음 (dB(A))"],
                                    mode="markers",
                                    name="실측 소음",
                                    marker=dict(size=7, color="#94A3B8", symbol="circle-open"),
                                ),
                                row=2,
                                col=1,
                            )
                        if current_noise is not None:
                            fig_compare.add_trace(
                                go.Scatter(
                                    x=[target_q],
                                    y=[current_noise],
                                    mode="markers",
                                    name="현재 목표 소음",
                                    showlegend=False,
                                    marker=dict(
                                        size=12,
                                        color="#64748B",
                                        symbol="circle",
                                        line=dict(color="white", width=2),
                                    ),
                                ),
                                row=2,
                                col=1,
                            )
                        if new_noise is not None:
                            fig_compare.add_trace(
                                go.Scatter(
                                    x=[target_q],
                                    y=[new_noise],
                                    mode="markers",
                                    name="신규 목표 소음",
                                    showlegend=False,
                                    marker=dict(
                                        size=14,
                                        color="#0D9488",
                                        symbol="diamond",
                                        line=dict(color="white", width=2),
                                    ),
                                ),
                                row=2,
                                col=1,
                            )
                        fig_compare.update_layout(
                            title="현재·신규 세트 비교",
                            height=760,
                        )
                        fig_compare.update_annotations(xanchor="left", font=dict(size=13, color="#334155"))
                        fig_compare.layout.annotations[0].x = 0
                        fig_compare.layout.annotations[1].x = 0
                        pressure_focus = np.r_[cp, np_, current_target_p, new_target_p]
                        pressure_focus = pressure_focus[np.isfinite(pressure_focus)]
                        pressure_ceiling = max(float(pressure_focus.max()), 1.0) * 1.18
                        fig_compare.update_yaxes(
                            title_text="정압 ΔP (Pa)", range=[0, pressure_ceiling], row=1, col=1
                        )
                        fig_compare.update_yaxes(title_text="소음 (dB(A))", row=2, col=1)
                        fig_compare.update_xaxes(title_text="세트 풍량 Q (CMM)", rangemode="tozero")
                        fig_compare.update_xaxes(matches="x", row=2, col=1)
                        for row_index in (1, 2):
                            fig_compare.add_vline(
                                x=target_q,
                                line=dict(color="#CBD5E1", width=1.5, dash="dash"),
                                layer="below",
                                row=row_index,
                                col=1,
                            )
                        # No synthetic sound curve when only target-point predictions exist.
                        if (
                            current_noise is None
                            and new_noise is None
                            and not np.any(np.isfinite(current_noise_on_curve))
                            and not np.any(np.isfinite(new_noise_on_curve))
                            and raw_noise_points.empty
                        ):
                            fig_compare.add_annotation(
                                text="소음 데이터가 없습니다",
                                x=0.5,
                                y=0.5,
                                xref="x2 domain",
                                yref="y2 domain",
                                showarrow=False,
                                font=dict(size=14, color="#94A3B8"),
                            )
                        mark_extrapolated_target(
                            fig_compare, cq, cp, target_q, current_target_p, "현재", "#64748B"
                        )
                        mark_extrapolated_target(
                            fig_compare, nq, np_, target_q, new_target_p, "신규", "#0D9488"
                        )
                        with chart_slot.container():
                            # Explicit numeric summaries avoid overlapping labels at nearby operating points.
                            noise_current_text = f"{current_noise:.1f}" if current_noise is not None else "—"
                            noise_new_text = f"{new_noise:.1f}" if new_noise is not None else "—"
                            st.markdown(
                                f"""<div class="set-chart-summary">
                                <div><span>목표 풍량</span><strong>{target_q:,.2f}<small> CMM</small></strong></div>
                                <div><span>필요 정압 · 현재 → 신규</span><strong>{current_target_p:,.1f} <em>→</em> <b>{new_target_p:,.1f}</b><small> Pa</small></strong></div>
                                <div><span>목표 소음 · 현재 → 신규</span><strong>{noise_current_text} <em>→</em> <b>{noise_new_text}</b><small> dB(A)</small></strong></div>
                                </div>""",
                                unsafe_allow_html=True,
                            )
                            _show_chart(
                                fig_compare,
                                width="stretch",
                                layout_overrides=dict(
                                    height=760,
                                    margin=dict(l=50, r=14, t=56, b=90),
                                    legend=dict(y=-0.1, font=dict(size=10.5)),
                                ),
                            )
                            st.caption(
                                "회색: 현재 · 청록색: 신규 · 실선: 팬 PQ/소음 · 점선: 세트 저항 · "
                                "세로 가이드: 목표 풍량. 정압축은 팬 곡선·목표점 중심이며, 위아래 풍량 축은 함께 확대됩니다."
                            )

                        uncertainty = float(set2_uncertainty) / 100.0
                        if uncertainty > 0:
                            low_rpm = _solve_required_rpm(
                                target_q,
                                new_target_p * (1.0 - uncertainty),
                                nbq,
                                nbp,
                                nbw,
                                nblw,
                                new_base_d,
                                new_base_x,
                                new_base_rpm,
                                set2_new_d,
                                set2_new_x,
                                air_density_1,
                                air_density_2,
                                set2_new_count,
                                set2_new_arr,
                                fan_category,
                                use_eff_correction,
                                new_original_eff,
                                new_required_rpm,
                                allow_limited_extrapolation=set2_allow_extrapolation,
                                mu1=air_viscosity_1,
                                mu2=air_viscosity_2,
                            )
                            high_rpm = _solve_required_rpm(
                                target_q,
                                new_target_p * (1.0 + uncertainty),
                                nbq,
                                nbp,
                                nbw,
                                nblw,
                                new_base_d,
                                new_base_x,
                                new_base_rpm,
                                set2_new_d,
                                set2_new_x,
                                air_density_1,
                                air_density_2,
                                set2_new_count,
                                set2_new_arr,
                                fan_category,
                                use_eff_correction,
                                new_original_eff,
                                new_required_rpm,
                                allow_limited_extrapolation=set2_allow_extrapolation,
                                mu1=air_viscosity_1,
                                mu2=air_viscosity_2,
                            )
                            rpm_bounds = [rpm for rpm in (low_rpm, high_rpm) if rpm is not None]
                            if rpm_bounds:
                                st.info(
                                    f"변경 세트 저항 ±{set2_uncertainty}% 고려 시 신규 필요 RPM: "
                                    f"**{min(rpm_bounds):.0f} ~ {max(rpm_bounds):.0f} RPM**"
                                )

                        component_note = "변경 부품 없음"
                        if len(changed2):
                            component_rows = []
                            for _, part in changed2.iterrows():
                                old_curve, new_curve = component_pressure_by_model(
                                    np.array([target_q]), part, air_density_1, air_density_2
                                )
                                old_dp = component_scale * float(old_curve[0])
                                new_dp = component_scale * float(new_curve[0])
                                old_velocity = (target_q / 60.0) / float(part["기존 면적 m²"])
                                new_velocity = (target_q / 60.0) / float(part["신규 면적 m²"])
                                component_rows.append(
                                    {
                                        "변경 부품": part.get("부품명", "부품"),
                                        "입력 모델": part.get("입력 모델", RESISTANCE_POINT_MODEL),
                                        "기존 면풍속 (m/s)": old_velocity,
                                        "신규 면풍속 (m/s)": new_velocity,
                                        "기존 압손 (Pa)": old_dp,
                                        "신규 압손 (Pa)": new_dp,
                                        "압손 변화 (Pa)": new_dp - old_dp,
                                    }
                                )
                            component_note = " / ".join(
                                f"{row['변경 부품']} {row['기존 압손 (Pa)']:.1f}→{row['신규 압손 (Pa)']:.1f} Pa"
                                for row in component_rows
                            )
                            st.markdown("##### 목표 풍량의 변경 부품 압손")
                            st.dataframe(
                                pd.DataFrame(component_rows).style.format(precision=2, na_rep="—"),
                                width="stretch",
                                hide_index=True,
                            )

                        try:
                            powerpoint_report = None
                            if st.button("세트 보고서 생성", key="set2_generate_report", type="primary"):
                                powerpoint_report = create_set_prediction_powerpoint_report(
                                    fan_category=fan_category,
                                    dim_label=dim_label,
                                    target_q=target_q,
                                    current_rpm=current_required_rpm,
                                    new_rpm=new_required_rpm,
                                    current_p=current_target_p,
                                    new_p=new_target_p,
                                    current_power=current_power,
                                    new_power=new_power,
                                    current_noise=current_noise,
                                    new_noise=new_noise,
                                    current_d=set2_cur_d,
                                    current_x=set2_cur_x,
                                    new_d=set2_new_d,
                                    new_x=set2_new_x,
                                    current_k=current_k,
                                    system_exponent=system_exponent,
                                    fit_rmse_pct=fit_rmse_pct,
                                    current_q_curve=cq,
                                    current_p_curve=cp,
                                    current_system_curve=current_system_on_curve,
                                    new_q_curve=nq,
                                    new_p_curve=np_,
                                    new_system_curve=new_system_on_curve,
                                    current_noise_curve=current_noise_on_curve,
                                    new_noise_curve=new_noise_on_curve,
                                    current_power_source=current_power_source,
                                    new_power_source=new_power_source,
                                    current_noise_source=current_noise_source,
                                    new_noise_source=new_noise_source,
                                    component_note=component_note
                                    + " · "
                                    + " / ".join(
                                        f"{name}: {details['method']} ({details['extrapolation_pct']:.1f}%)"
                                        for name, details in [
                                            ("현재", current_rpm_details),
                                            ("신규", new_rpm_details),
                                        ]
                                    ),
                                )
                            st.markdown("##### 팬 시스템 성능 보고서")
                            st.caption(
                                "PQ·저항 및 소음 원데이터가 PowerPoint 차트에 내장됩니다. "
                                "PowerPoint의 ‘데이터 편집’과 ‘축 서식’에서 데이터, 축 범위, "
                                "선 색상과 두께를 직접 수정할 수 있습니다."
                            )
                            report_col, csv_col = st.columns(2)
                            with report_col:
                                if powerpoint_report is not None:
                                    st.download_button(
                                        "팬 시스템 성능 보고서 (PPTX)",
                                        powerpoint_report,
                                        file_name=(
                                            f"fan_system_report_v10_"
                                            f"{datetime.date.today().strftime('%Y%m%d')}.pptx"
                                        ),
                                        mime=(
                                            "application/vnd.openxmlformats-officedocument."
                                            "presentationml.presentation"
                                        ),
                                        key="set2_powerpoint_report",
                                        width="stretch",
                                        on_click="ignore",
                                    )
                            with csv_col:
                                st.download_button(
                                    "📥 현재·신규 비교 CSV",
                                    comparison.to_csv(index=False).encode("utf-8-sig"),
                                    file_name=(
                                        f"set_target_prediction_v10_"
                                        f"{datetime.date.today().strftime('%Y%m%d')}.csv"
                                    ),
                                    mime="text/csv",
                                    key="set2_download",
                                    width="stretch",
                                    on_click="ignore",
                                )
                        except Exception as report_err:
                            st.warning(f"PowerPoint 보고서를 생성하지 못했습니다: {report_err}")
                        if new_power is None:
                            st.caption(
                                "단품 축동력/효율과 세트 소비전력 보정자료가 없어 "
                                "신규 소비전력은 계산하지 못했습니다."
                            )
                        if new_noise is None:
                            st.caption(
                                "단품 소음과 세트 소음 raw data가 없어 "
                                "신규 세트 소음은 계산하지 못했습니다."
                            )
                        st.caption(
                            "직경·RPM 변화는 팬 상사법칙으로, 저항 변화는 목표 풍량의 변경 필요정압으로 "
                            "반영합니다. 단품 W가 없을 때 전력은 raw 총정압효율을 사용합니다. 단품 Lw가 "
                            "없고 세트 소음점이 2개 이상이면 log(RPM) 회귀와 팬 종류별 형상 상사를 사용하며, "
                            "회귀가 불가능할 때만 KN=Ls−10logQ−20logP를 보조 모델로 사용합니다."
                        )
            except Exception as analysis_err:
                st.warning(
                    f"입력된 값 중 계산 가능한 구간까지 처리했습니다. 확인이 필요한 항목: {analysis_err}"
                )


# =============================================================================
# TAB 2: 문헌/경험식 기반 설계
# =============================================================================
def render_tab_empirical(settings):
    """경험식 설계 화면. 전역 입력은 settings로 전달한다."""
    fan_category = settings["fan_category"]
    is_axial = settings["is_axial"]
    dim_label = settings["dim_label"]
    dim_key = settings["dim_key"]
    flow_unit = settings["flow_unit"]
    pressure_unit = settings["pressure_unit"]
    power_unit = settings["power_unit"]
    flow_unit_disp = settings["flow_unit_disp"]
    pressure_unit_disp = settings["pressure_unit_disp"]
    power_unit_disp = settings["power_unit_disp"]
    air_density_1 = settings["air_density_1"]
    air_density_2 = settings["air_density_2"]
    air_viscosity_1 = settings["air_viscosity_1"]
    air_viscosity_2 = settings["air_viscosity_2"]
    use_eff_correction = settings["use_eff_correction"]

    st.markdown("### 📖 문헌 및 경험식 기반 성능/소음 설계")
    st.markdown(
        "기존의 기준 데이터가 없을 때, 목표 운전점과 형상 수치를 기반으로 전형적인 효율 및 소음 계수를 활용하여 사양을 추정합니다."
    )

    col2_1, col2_2 = st.columns([1.2, 2.0])

    # Pre-sets for different fans
    presets = {
        "시로코팬 (Sirocco)": {
            "eff": 0.50,
            "kn": 38.0,
            "desc": "다익송풍기: 낮은 비속도, 고정압 및 소형화 가능, 날개가 전향식.",
        },
        "터보팬 (Turbo)": {
            "eff": 0.78,
            "kn": 28.0,
            "desc": "후향식 블레이드: 높은 효율 및 저소음 설계 가능, 중/고정압 계통 적합.",
        },
        "축류팬 (Axial)": {
            "eff": 0.70,
            "kn": 42.0,
            "desc": "축류 송풍기: 고비속도, 대풍량/저정압 가동용. 소음이 비교적 큼.",
        },
        "횡류팬 (Cross-flow)": {
            "eff": 0.38,
            "kn": 36.0,
            "desc": "관류 송풍기: 넓고 균일한 평면 유량 형성, 저속 운전에 적합.",
        },
    }
    preset_options = {
        CENTRIFUGAL_CATEGORY: ["시로코팬 (Sirocco)", "터보팬 (Turbo)"],
        AXIAL_CATEGORY: ["축류팬 (Axial)"],
        CROSSFLOW_CATEGORY: ["횡류팬 (Cross-flow)"],
    }[fan_category]

    with col2_1:
        with st.container(border=True, key="panel_4225"):
            st.markdown("#### 1. 팬 타입 설정")
            selected_preset = st.selectbox(
                "팬 세부 타입 프리셋", preset_options, key=f"tab2_preset_{dim_key}"
            )

            # Display description
            st.caption(f"**특징**: {presets[selected_preset]['desc']}")

            # Setup pre-filled inputs
            default_eff = presets[selected_preset]["eff"]
            default_kn = presets[selected_preset]["kn"]

            st.markdown("#### 2. 사양 및 치수 입력")
            target_q_val = unit_number_input(
                f"목표 작동 유량 ({flow_unit_disp})",
                value=20.0,
                min_value=0.1,
                step=1.0,
                quantity="flow",
                unit=flow_unit,
            )
            target_p_val = unit_number_input(
                f"목표 계통 정압 ({pressure_unit_disp})",
                value=20.0,
                min_value=0.1,
                step=1.0,
                quantity="pressure",
                unit=pressure_unit,
            )

            target_rpm = st.number_input(
                "설계 회전수 (RPM)", value=1800.0, min_value=50.0, step=100.0, format="%.0f"
            )
            target_d = st.number_input(
                "팬 설계 외경 D (mm)", value=250.0, min_value=10.0, step=10.0, format="%.0f"
            )
            if is_axial:
                st.text_input("팬 설계 독립 높이 H", value="미적용", disabled=True, key="tab2_x_axial")
                target_x = target_d
            else:
                target_x = st.number_input(
                    f"팬 설계 {dim_label} (mm)", value=150.0, min_value=5.0, step=5.0, format="%.0f"
                )

            st.markdown("#### 3. 문헌 경험 계수 수정")
            user_eff = st.slider(
                "가정 효율 (η)", min_value=0.10, max_value=0.95, value=default_eff, step=0.01
            )
            user_kn = st.slider(
                "특성 소음 계수 (KN)", min_value=20.0, max_value=50.0, value=default_kn, step=1.0
            )

    with col2_2:
        # Internal computations
        # Convert Q to CMM, P to Pa/mmAq
        q_cmm = convert_flow(target_q_val, flow_unit, "CMM (m³/min)")
        q_m3s = q_cmm / 60.0
        p_pa = convert_pressure(target_p_val, pressure_unit, "Pa")
        p_mmaq = convert_pressure(target_p_val, pressure_unit, "mmAq")

        # Shaft power computation (W_shaft = Q_m3s * P_pa / η)
        w_shaft = (q_m3s * p_pa) / user_eff

        # --- 무차원 비속도 σ / 비직경 δ (Cordier 좌표계) ---
        # σ = ω · Q^0.5 / (Δp/ρ)^0.75   [ω: rad/s, Q: m³/s, Δp/ρ: J/kg = m²/s²]
        # δ = D · (Δp/ρ)^0.25 / Q^0.5
        # Cordier 곡선/팬 타입 판정 범위 모두 이 무차원 정의를 기준으로 한다.
        omega = target_rpm * 2.0 * np.pi / 60.0  # rad/s
        y_head = p_pa / air_density_1 if air_density_1 > 0 else 0.0  # 비일 (J/kg)
        d_m = target_d / 1000.0
        if y_head > 0 and q_m3s > 0:
            sigma = omega * np.sqrt(q_m3s) / (y_head**0.75)
            delta = d_m * (y_head**0.25) / np.sqrt(q_m3s)
        else:
            sigma, delta = 0.0, 0.0
        # 참고용 한국식 비속도 Ns (rpm, m³/min, mmAq) — 국내 실무 관용값 병기
        if p_mmaq > 0 and q_cmm > 0:
            ns_kr = target_rpm * np.sqrt(q_cmm) / (p_mmaq**0.75)
        else:
            ns_kr = 0.0
        # 표시용 별칭 (기존 변수명 유지 목적)
        ns, ds = sigma, delta

        # Noise estimation: Lw = KN + 10 log10(Q_CMM) + 20 log10(P_mmaq)
        if q_cmm > 0 and p_mmaq > 0:
            est_lw = user_kn + 10 * np.log10(q_cmm) + 20 * np.log10(p_mmaq)
        else:
            est_lw = 0.0

        with st.container(border=True, key="panel_4293"):
            st.markdown("#### 예측 사양 및 설계 타당성 평가")

            col_res1, col_res2, col_res3 = st.columns(3)
            with col_res1:
                st.markdown(
                    f"""
                <div class="metric-box">
                    <div class="title">예측 축동력</div>
                    <div class="value">{convert_power(w_shaft, "W", power_unit):.2f} <span style="font-size: 1.0rem; font-weight: normal; color: #64748B;">{power_unit_disp}</span></div>
                    <div class="delta" style="color: #64748B;">({w_shaft/1000.0:.3f} kW)</div>
                </div>
                """,
                    unsafe_allow_html=True,
                )
            with col_res2:
                st.markdown(
                    f"""
                <div class="metric-box">
                    <div class="title">예측 소음레벨</div>
                    <div class="value">{est_lw:.1f} dB(A)</div>
                    <div class="delta" style="color: #64748B;">Lw (특성 KN={user_kn})</div>
                </div>
                """,
                    unsafe_allow_html=True,
                )
            with col_res3:
                st.markdown(
                    f"""
                <div class="metric-box">
                    <div class="title">무차원 비속도 σ</div>
                    <div class="value">{ns:.3f}</div>
                    <div class="delta" style="color: #64748B;">δ = {ds:.2f} · Ns(KR)≈{ns_kr:.0f}</div>
                </div>
                """,
                    unsafe_allow_html=True,
                )

            st.markdown("---")
            st.markdown("##### Cordier 다이어그램 상의 위치 분석")
            st.markdown(
                "설계점의 무차원 비속도 σ와 비직경 δ를 Cordier 참고 곡선에 표시하고, 선택한 팬 형식의 σ 범위와 비교합니다."
            )

            # 기존 σ 판정 범위를 차트와 진단 문구에 함께 사용한다.
            sigma_bands = {
                "시로코팬 (Sirocco)": (0.15, 0.60),
                "터보팬 (Turbo)": (0.40, 1.80),
                "축류팬 (Axial)": (1.50, 8.00),
                "횡류팬 (Cross-flow)": (0.05, 0.35),
            }
            lo, hi = sigma_bands[selected_preset]
            is_normal = lo <= ns <= hi
            status_label = "범위 내" if is_normal else "범위 밖"
            status_color = "#0D9488" if is_normal else "#D97706"

            # Cordier 최적 효율 곡선 (무차원 σ, δ) — 대표 근사값
            cordier_ns = np.array([0.10, 0.20, 0.30, 0.50, 0.70, 1.00, 1.50, 2.00, 3.00, 5.00, 7.00, 10.0])
            cordier_ds = np.array([9.00, 6.00, 4.50, 3.20, 2.60, 2.10, 1.70, 1.50, 1.25, 1.00, 0.90, 0.80])

            # Plot Cordier Diagram
            fig_cor = go.Figure()
            fig_cor.add_vrect(
                x0=lo,
                x1=hi,
                fillcolor="rgba(13,148,136,0.12)",
                line_width=0,
                layer="below",
                exclude_empty_subplots=False,
            )
            for boundary in (lo, hi):
                fig_cor.add_vline(
                    x=boundary,
                    line=dict(color="#0D9488", width=1, dash="dot"),
                    layer="below",
                    exclude_empty_subplots=False,
                )
            fig_cor.add_annotation(
                x=0,
                y=1.07,
                xref="paper",
                yref="paper",
                xanchor="left",
                yanchor="bottom",
                text=f"녹색 영역 · σ {lo:g}–{hi:g} | 설계점: {status_label}",
                showarrow=False,
                font=dict(size=12, color=status_color),
            )
            fig_cor.add_trace(
                go.Scatter(
                    x=cordier_ns,
                    y=cordier_ds,
                    mode="lines+markers",
                    name="Cordier 최적 효율 곡선",
                    line=dict(color="#64748B", width=1.5, dash="dash"),
                    marker=dict(color="#64748B", size=6),
                    hovertemplate="비속도 σ: %{x:.3f}<br>비직경 δ: %{y:.3f}<extra>Cordier 참고 곡선</extra>",
                )
            )

            # Add user design point
            fig_cor.add_trace(
                go.Scatter(
                    x=[ns],
                    y=[ds],
                    mode="markers",
                    marker=dict(
                        color=status_color,
                        size=16,
                        symbol="star-diamond",
                        line=dict(color="white", width=1.5),
                    ),
                    name=f"설계점 · {status_label}",
                    hovertemplate=(
                        "비속도 σ: %{x:.3f}<br>비직경 δ: %{y:.3f}"
                        f"<br>σ 참고 범위: {lo:g}–{hi:g}<extra>{status_label}</extra>"
                    ),
                )
            )

            fig_cor.update_layout(
                xaxis_title="무차원 비속도 σ",
                yaxis_title="무차원 비직경 δ",
                xaxis_type="log",
                yaxis_type="log",
                plot_bgcolor="#FFFFFF",
                paper_bgcolor="rgba(0,0,0,0)",
                margin=dict(l=40, r=40, t=40, b=40),
                height=460,
                legend=dict(x=0.02, y=0.02, bordercolor="#E2E8F0", borderwidth=1),
            )
            # 로그 축에 참고 범위 전체와 설계점을 모두 포함한다 (횡류팬 하한 포함).
            x_limits = [*cordier_ns, lo, hi]
            if ns > 0:
                x_limits.append(ns)
            fig_cor.update_xaxes(
                showgrid=True,
                gridcolor="#EEF2F6",
                linecolor="#E2E8F0",
                range=[np.log10(min(x_limits)) - 0.12, np.log10(max(x_limits)) + 0.12],
            )
            fig_cor.update_yaxes(showgrid=True, gridcolor="#EEF2F6", linecolor="#E2E8F0")
            fig_cor.update_layout(title_text="Cordier 설계 위치 · 비속도와 비직경")
            _show_chart(fig_cor, width="stretch")
            st.caption(
                f"녹색 영역: {selected_preset}의 비속도 σ 기준 참고 범위 ({lo:g}–{hi:g}). "
                "비직경 δ의 허용 범위나 실제 효율을 보장하지 않으며, 기존 경험식 판정 기준을 시각화한 영역입니다."
            )

            if selected_preset == "횡류팬 (Cross-flow)":
                st.info(
                    "ℹ️ 횡류팬은 관류(cross-flow) 구조상 Cordier 상사 좌표에 엄밀히 대응하지 않습니다. σ 값과 판정은 참고 지표로만 활용하십시오."
                )

            if is_normal:
                st.success(
                    f"✅ σ 기준 범위 내: 무차원 비속도(σ={ns:.3f})가 **{selected_preset}**의 대표 범위(σ≈{lo}–{hi}) 내에 위치합니다."
                )
            else:
                st.warning(
                    f"⚠️ σ 기준 범위 밖: 현재 σ={ns:.3f}는 **{selected_preset}**의 대표 범위(σ≈{lo}–{hi})를 벗어나 있습니다. 목표 유량·정압 대비 회전수를 재검토하거나 다른 팬 형식을 고려하십시오."
                )


# =============================================================================
# TAB 6: 실험 데이터 기반 파생 설계
# =============================================================================
def render_tab_experiments(settings):
    """실험 기반 설계 화면. 전역 입력은 settings로 전달한다."""
    fan_category = settings["fan_category"]
    is_axial = settings["is_axial"]
    dim_label = settings["dim_label"]
    dim_key = settings["dim_key"]
    flow_unit = settings["flow_unit"]
    pressure_unit = settings["pressure_unit"]
    power_unit = settings["power_unit"]
    flow_unit_disp = settings["flow_unit_disp"]
    pressure_unit_disp = settings["pressure_unit_disp"]
    power_unit_disp = settings["power_unit_disp"]
    air_density_1 = settings["air_density_1"]
    air_density_2 = settings["air_density_2"]
    air_viscosity_1 = settings["air_viscosity_1"]
    air_viscosity_2 = settings["air_viscosity_2"]
    use_eff_correction = settings["use_eff_correction"]

    st.markdown("### 실험 데이터베이스 기반 AI 모델 추천 및 파생 설계")
    st.markdown(
        "사내에 축적된 다수의 송풍기 기종별 실험 데이터베이스(CSV/Excel)를 일괄 업로드하여, 신규 목표 사양($Q_{req}, P_{req}$)을 충족할 수 있는 최적의 모태 팬 모델을 AI가 검색·스케일링하여 순위별로 추천합니다."
    )

    # Built-in Sample Database CSV
    SAMPLE_DB_CSV = """Model_Name,D,H,N,Q,P,W,Lw,Power_Type,Mfg_Method
Centrifugal_Turbo_A,250,60,1800,0.0,220.0,80.0,58.0,Torque,3D
Centrifugal_Turbo_A,250,60,1800,5.0,215.0,95.0,58.5,Torque,3D
Centrifugal_Turbo_A,250,60,1800,10.0,200.0,110.0,59.0,Torque,3D
Centrifugal_Turbo_A,250,60,1800,15.0,175.0,125.0,60.5,Torque,3D
Centrifugal_Turbo_A,250,60,1800,20.0,140.0,135.0,62.0,Torque,3D
Centrifugal_Turbo_A,250,60,1800,25.0,90.0,140.0,64.5,Torque,3D
Centrifugal_Turbo_A,250,60,1800,30.0,20.0,130.0,67.0,Torque,3D
Centrifugal_Sirocco_B,150,50,2000,0.0,80.0,25.0,52.0,Motor_Input,2D_Injection
Centrifugal_Sirocco_B,150,50,2000,2.0,78.0,32.0,52.5,Motor_Input,2D_Injection
Centrifugal_Sirocco_B,150,50,2000,4.0,75.0,40.0,53.0,Motor_Input,2D_Injection
Centrifugal_Sirocco_B,150,50,2000,6.0,68.0,48.0,54.0,Motor_Input,2D_Injection
Centrifugal_Sirocco_B,150,50,2000,8.0,55.0,55.0,55.5,Motor_Input,2D_Injection
Centrifugal_Sirocco_B,150,50,2000,10.0,35.0,62.0,57.0,Motor_Input,2D_Injection
Centrifugal_Sirocco_B,150,50,2000,12.0,10.0,68.0,59.0,Motor_Input,2D_Injection
Axial_HighFlow_C,400,100,1450,0.0,95.0,120.0,64.0,Torque,2D_Slide
Axial_HighFlow_C,400,100,1450,15.0,90.0,135.0,64.5,Torque,2D_Slide
Axial_HighFlow_C,400,100,1450,30.0,80.0,150.0,65.0,Torque,2D_Slide
Axial_HighFlow_C,400,100,1450,45.0,65.0,160.0,66.0,Torque,2D_Slide
Axial_HighFlow_C,400,100,1450,60.0,48.0,165.0,67.5,Torque,2D_Slide
Axial_HighFlow_C,400,100,1450,75.0,28.0,160.0,69.5,Torque,2D_Slide
Axial_HighFlow_C,400,100,1450,90.0,5.0,145.0,72.0,Torque,2D_Slide
CrossFlow_Slim_D,120,300,1500,0.0,45.0,15.0,45.0,Torque,2D_Injection
CrossFlow_Slim_D,120,300,1500,1.0,43.0,18.0,46.0,Torque,2D_Injection
CrossFlow_Slim_D,120,300,1500,2.0,40.0,21.0,47.0,Torque,2D_Injection
CrossFlow_Slim_D,120,300,1500,3.0,35.0,24.0,48.0,Torque,2D_Injection
CrossFlow_Slim_D,120,300,1500,4.0,28.0,26.0,50.0,Torque,2D_Injection
CrossFlow_Slim_D,120,300,1500,5.0,18.0,27.0,52.0,Torque,2D_Injection
CrossFlow_Slim_D,120,300,1500,6.0,5.0,25.0,54.0,Torque,2D_Injection
"""

    col6_1, col6_2 = st.columns([1.1, 2])

    with col6_1:
        with st.container(border=True, key="panel_4438"):
            st.markdown("#### 📁 1. 실험 데이터베이스 로드")
            uploaded_file = st.file_uploader(
                "실험 DB 파일 업로드 (.csv, .xlsx)",
                type=["csv", "xlsx"],
                help="기종명, 직경(D), 폭(H), 회전수(N), 풍량(Q), 정압(P), 동력(W), 소음(Lw) 등을 담은 파일",
                key="tab6_upload",
            )

            # Show download template link
            st.download_button(
                label="📥 표준 DB 포맷 템플릿 다운로드",
                data=SAMPLE_DB_CSV,
                file_name="fan_test_database_template.csv",
                mime="text/csv",
                key="tab6_download_template",
                on_click="ignore",
            )

            # Read database
            if uploaded_file is not None:
                try:
                    if uploaded_file.name.endswith(".xlsx"):
                        df_raw = pd.read_excel(uploaded_file)
                    else:
                        df_raw = pd.read_csv(uploaded_file)
                    st.success("데이터베이스 파일 로드 완료!")
                except Exception as ex:
                    st.error(f"파일 로드 실패: {ex}")
                    return
            else:
                st.info(
                    "예시 데이터로 기능을 확인할 수 있습니다. 실제 설계에는 측정 데이터를 업로드하십시오."
                )
                df_raw = pd.read_csv(io.StringIO(SAMPLE_DB_CSV))

            # Column Sanitization & Mapping
            df_cleaned = df_raw.copy()
            df_cleaned.columns = [c.strip().lower() for c in df_cleaned.columns]

            # Show preview
            with st.expander("로드된 로우 데이터셋 미리보기", expanded=False):
                st.dataframe(df_cleaned, width="stretch")

        with st.container(border=True, key="panel_4479"):
            st.markdown("#### 2. 신규 설계 목표 사양")
            st.caption(
                "ℹ️ 이 탭은 상단 설정의 글로벌 단위와 무관하게 CMM·Pa·W·dB(A) 고정 단위로 동작합니다."
            )
            t6_q_req = st.number_input("목표 풍량 Q_req (CMM)", value=15.0, min_value=0.1, key="tab6_q_req")
            t6_p_req = st.number_input("목표 정압 P_req (Pa)", value=120.0, min_value=0.1, key="tab6_p_req")

            st.markdown("##### 설계 및 제작 제약 조건")
            t6_d_max = st.number_input(
                "최대 허용 직경 D_max (mm)", value=800.0, min_value=50.0, key="tab6_d_max", format="%.0f"
            )
            t6_n_max = st.number_input(
                "최대 허용 속도 N_max (RPM)", value=4000.0, min_value=100.0, key="tab6_n_max", format="%.0f"
            )

            st.markdown("##### 동력 및 실험 보정")
            t6_motor_eff = st.slider("모터 소비전력 데이터용 모터 효율 (%)", 50, 99, 82, key="tab6_motor_eff")

            col_c1, col_c2, col_c3 = st.columns(3)
            with col_c1:
                c_q = st.slider("유량 보정 C_q", 0.80, 1.20, 1.00, 0.01, key="tab6_cq")
            with col_c2:
                c_p = st.slider("정압 보정 C_p", 0.80, 1.20, 1.00, 0.01, key="tab6_cp")
            with col_c3:
                c_w = st.slider("동력 보정 C_w", 0.80, 1.50, 1.00, 0.01, key="tab6_cw")

            t6_preference = st.radio(
                "설계 후보 정렬 기준",
                [
                    "최고 효율 순 (Max Efficiency)",
                    "최소 동력 순 (Min Power)",
                    "최소 소음 순 (Min Noise)",
                    "최소 직경 순 (Min Diameter)",
                ],
                key="tab6_preference",
            )

    with col6_2:
        try:
            # Map required columns safely
            col_mapping = {
                "model_name": ["model_name", "model", "기종명", "기종", "name"],
                "d": ["d", "diameter", "직경", "d1", "d ("],
                "h": ["h", "width", "height", "폭", "높이", "길이", "l", "h1", "h (", "l ("],
                "n": ["n", "rpm", "회전수", "n1", "n ("],
                "q": ["q", "flow", "flow_rate", "풍량", "유량", "cmm", "cmh", "cfm", "q ("],
                "p": ["p", "pressure", "static", "정압", "압력", "mmaq", "kpa", "p ("],
                "w": ["w", "power", "동력", "소비전력", "watt", "w ("],
                "lw": ["lw", "noise", "db", "db(a)", "소음", "소음레벨", "sound"],
            }

            # power_type/mfg_method는 별도 처리 컬럼이므로 수치 매핑 후보에서 제외
            reserved_cols = {"power_type", "mfg_method"}
            candidate_cols = [c for c in df_cleaned.columns if c not in reserved_cols]

            mapped_df = pd.DataFrame()
            used_cols = set()
            for key, options in col_mapping.items():
                matched_col = None
                # 1순위: 헤더 정확 일치
                for opt in options:
                    if opt in candidate_cols and opt not in used_cols:
                        matched_col = opt
                        break
                # 2순위: 부분 일치 ("q (cmm)" 같은 단위 병기 헤더 대응, 2글자 이상 키워드만)
                if matched_col is None:
                    for opt in options:
                        if len(opt) < 2:
                            continue
                        matched_col = next(
                            (c for c in candidate_cols if c not in used_cols and opt in c), None
                        )
                        if matched_col is not None:
                            break
                if matched_col is not None:
                    used_cols.add(matched_col)
                    mapped_df[key] = df_cleaned[matched_col]
                else:
                    if key == "lw":
                        mapped_df[key] = np.nan
                    elif key == "h" and is_axial:
                        # 축류팬은 독립 높이가 없으므로 데이터셋에 H 열을 요구하지 않는다.
                        mapped_df[key] = np.nan
                    else:
                        raise ValueError(
                            f"필수 컬럼 '{key}'에 매핑되는 컬럼을 찾을 수 없습니다. 템플릿을 참고하여 헤더를 교정하십시오."
                        )

            if is_axial:
                # 이후 공통 그룹/표시 로직용 자리값일 뿐 계산 엔진에서는 사용하지 않는다.
                mapped_df["h"] = mapped_df["d"]

            # Optional columns
            mapped_df["power_type"] = (
                df_cleaned["power_type"] if "power_type" in df_cleaned.columns else "Torque"
            )
            mapped_df["mfg_method"] = df_cleaned["mfg_method"] if "mfg_method" in df_cleaned.columns else "3D"

            # 수치 검증: 문자/빈칸은 NaN으로 변환 후 필수 수치가 결측인 행 제거
            for num_key in ["d", "h", "n", "q", "p", "w", "lw"]:
                mapped_df[num_key] = pd.to_numeric(mapped_df[num_key], errors="coerce")
            n_rows_before = len(mapped_df)
            mapped_df = mapped_df.replace([np.inf, -np.inf], np.nan).dropna(
                subset=["d", "h", "n", "q", "p", "w"]
            )
            mapped_df = mapped_df[(mapped_df["q"] >= 0) & (mapped_df["p"] >= 0) & (mapped_df["w"] > 0)]
            if len(mapped_df) < n_rows_before:
                st.warning(
                    f"⚠️ 수치로 해석할 수 없거나 비어 있는 {n_rows_before - len(mapped_df)}개 행을 제외하고 분석합니다."
                )
            if len(mapped_df) == 0:
                raise ValueError("유효한 수치 데이터 행이 없습니다")

            # 동일 기종이 복수의 시험 조건(D/H/N)으로 존재하면 조건별로 분리하여 평가
            cond_counts = (
                mapped_df.drop_duplicates(subset=["model_name", "d", "h", "n"]).groupby("model_name").size()
            )
            grouped = mapped_df.groupby(["model_name", "d", "h", "n"])
            matched_results = []
            failed_models = []

            st.caption(
                "각 기종의 측정 PQ 구간에서 161개 격자점과 원본 측정점을 탐색합니다. "
                "기하상사(H/L∝D)와 치수·RPM 제약을 적용한 후보 중 정렬 기준에 맞는 조합을 선택합니다."
            )
            for (model_name, d1, h1, n1), group in grouped:
                disp_name = (
                    f"{model_name} (D{d1:.0f}·N{n1:.0f})"
                    if cond_counts.get(model_name, 1) > 1
                    else model_name
                )
                if d1 <= 0 or h1 <= 0 or n1 <= 0:
                    failed_models.append(f"{disp_name}: D/H/N 값이 0 이하")
                    continue
                q1, p1, w1, db1 = clean_curve_arrays(
                    group["q"], group["p"], group["w"], group["lw"], required=(0, 1, 2)
                )
                if len(q1) < 2:
                    failed_models.append(f"{disp_name}: 서로 다른 유효 풍량점이 2개 미만")
                    continue
                p_type = str(group["power_type"].iloc[0]).strip().lower()
                w1_shaft = w1 * (t6_motor_eff / 100.0) if p_type == "motor_input" else w1
                eff1 = np.divide(q1 * p1 / 60, w1_shaft, out=np.zeros_like(q1), where=w1_shaft > 0)
                if np.any(eff1 > 1):
                    failed_models.append(
                        f"{disp_name}: Q·P/W 효율이 100%를 초과합니다. 단위와 동력 종류를 확인하십시오."
                    )
                    continue
                mfg_norm = str(group["mfg_method"].iloc[0]).strip().lower()
                if "injection" in mfg_norm or "한방" in mfg_norm:
                    eff_cap, mfg_lbl = 0.55, "2D 한방사출"
                elif "slide" in mfg_norm or "슬라이드" in mfg_norm:
                    eff_cap, mfg_lbl = 0.65, "2D 슬라이드"
                else:
                    eff_cap, mfg_lbl = 0.85, "3D 트위스티드"
                candidates = find_derived_candidates(
                    q1,
                    p1,
                    w1_shaft,
                    db1,
                    d1,
                    h1,
                    n1,
                    t6_q_req,
                    t6_p_req,
                    air_density_1,
                    air_density_2,
                    t6_d_max,
                    t6_n_max,
                    use_eff_correction,
                    eff_cap,
                    c_q,
                    c_p,
                    c_w,
                    fan_category,
                    mu1=air_viscosity_1,
                    mu2=air_viscosity_2,
                )
                if t6_preference == "최소 소음 순 (Min Noise)":
                    candidates = [candidate for candidate in candidates if np.isfinite(candidate["noise"])]
                if not candidates:
                    failed_models.append(
                        f"{disp_name}: 제약을 만족하는 후보 또는 선택 기준의 유효 데이터가 없음"
                    )
                    continue
                best = min(
                    candidates, key=lambda candidate: derived_candidate_score(candidate, t6_preference)
                )
                d2, h2, n2 = best["d2"], best["x2"], best["n2"]
                v_tip2 = np.pi * d2 * n2 / (60 * 1000)
                sigma = (n2 * 2 * np.pi / 60) * np.sqrt(t6_q_req / 60) / (t6_p_req / air_density_2) ** 0.75
                matched_results.append(
                    {
                        "모델 기종명": disp_name,
                        "제작 공법": mfg_lbl,
                        "예측 직경 D₂": f"{d2:.1f} mm",
                        "예측 폭 H₂": f"{h2:.1f} mm",
                        "예측 RPM N₂": f"{n2:.0f} RPM",
                        "예측 축동력 (W)": best["w"],
                        "예측 정압효율 (%)": best["eff"] * 100,
                        "예측 소음 dB(A)": best["noise"],
                        "소음 근거": (
                            "측정 범위 내 보간·상사" if np.isfinite(best["noise"]) else "데이터 없음"
                        ),
                        "팁속도 (m/s)": v_tip2,
                        "비속도 (σ)": sigma,
                        "후보 수": len(candidates),
                        "기준 풍량 (CMM)": best["q_ref"],
                        "q1": q1,
                        "p1": p1,
                        "w1": w1_shaft,
                        "db1": db1,
                        "eff1": eff1,
                        "r_D_val": best["r_d"],
                        "r_N_val": best["r_n"],
                        "mfg_val": mfg_norm,
                        "eff_cap_val": eff_cap,
                        "d1": d1,
                        "h1": h1,
                        "n1": n1,
                        "d2_val": d2,
                        "h2_val": h2,
                        "n2_val": n2,
                        "q_ref": best["q_ref"],
                        "p_ref": best["p_ref"],
                        "w_ref": best["w_ref"],
                        "noise_ref": best["noise_ref"],
                    }
                )

            if failed_models:
                st.warning(
                    "⚠️ 다음 기종은 계산 오류로 결과에서 제외되었습니다 — " + " / ".join(failed_models)
                )

            # Sort results based on user preference
            if t6_preference == "최고 효율 순 (Max Efficiency)":
                matched_results = sorted(matched_results, key=lambda x: x["예측 정압효율 (%)"], reverse=True)
            elif t6_preference == "최소 동력 순 (Min Power)":
                matched_results = sorted(matched_results, key=lambda x: x["예측 축동력 (W)"])
            elif t6_preference == "최소 소음 순 (Min Noise)":
                matched_results = sorted(matched_results, key=lambda x: x["예측 소음 dB(A)"])
            else:
                matched_results = sorted(
                    matched_results, key=lambda x: float(x["예측 직경 D₂"].split(" ")[0])
                )

            if len(matched_results) > 0:
                st.markdown("### 설계 후보 모델 매칭 및 랭킹 결과")
                st.markdown(
                    f"설계 및 직경/속도 제약 조건을 통과한 총 **{len(matched_results)}개** 기종의 예측 리스트입니다."
                )

                # Prepare view DataFrame
                view_cols = ["모델 기종명", "제작 공법", "예측 직경 D₂"]
                if not is_axial:
                    view_cols.append("예측 폭 H₂")
                view_cols += [
                    "예측 RPM N₂",
                    "예측 축동력 (W)",
                    "예측 정압효율 (%)",
                    "예측 소음 dB(A)",
                    "비속도 (σ)",
                ]
                view_df = pd.DataFrame(matched_results)[view_cols]
                view_df["소음 근거"] = [row["소음 근거"] for row in matched_results]
                st.dataframe(view_df.style.format(precision=2, na_rep="—"), width="stretch", hide_index=True)

                # Highlight top model
                top_model = matched_results[0]
                top_geometry = (
                    f"직경 **{top_model['예측 직경 D₂']}** (상사비: {top_model['r_D_val']:.2f}배), "
                    if is_axial
                    else f"직경 **{top_model['예측 직경 D₂']}** (상사비: {top_model['r_D_val']:.2f}배), 폭 **{top_model['예측 폭 H₂']}**, "
                )
                noise_description = (
                    f"{top_model['예측 소음 dB(A)']:.1f} dB(A)"
                    if np.isfinite(top_model["예측 소음 dB(A)"])
                    else "데이터 없음"
                )
                st.success(
                    f"**설계 후보 최적 기종**: **{top_model['모델 기종명']}** ({top_model['제작 공법']})\n\n"
                    f"- 동 기종을 {top_geometry}회전수 **{top_model['예측 RPM N₂']}**로 스케일링할 경우 "
                    f"목표 성능점에 도달하며 소비 전력은 **{top_model['예측 축동력 (W)']:.1f} W**, 예측 소음은 **{noise_description}** 로 분석됩니다."
                )

                st.markdown("---")
                st.markdown("#### 상세 예측 곡선 시각화 (선택 기종)")
                selected_model_name = st.selectbox(
                    "성능 곡선을 시각화할 기종 선택",
                    [x["모델 기종명"] for x in matched_results],
                    key="tab6_selected_model",
                )

                # Find details of selected model
                sel_m = [x for x in matched_results if x["모델 기종명"] == selected_model_name][0]

                # Reconstruct scaled curves for plotting
                q1_raw = sel_m["q1"]
                p1_raw = sel_m["p1"]
                w1_shaft = sel_m["w1"]
                db1_raw = sel_m["db1"]
                eff1_raw = sel_m["eff1"]

                r_D = sel_m["r_D_val"]
                r_N = sel_m["r_N_val"]
                mfg = sel_m["mfg_val"]
                eff_cap = sel_m["eff_cap_val"]

                # 선택한 후보점을 곡선에 포함해 표·마커·곡선의 계산을 정확히 일치시킨다.
                q_plot = np.unique(np.r_[q1_raw, sel_m["q_ref"]])
                p_plot = np.interp(q_plot, q1_raw, p1_raw)
                w_plot = np.interp(q_plot, q1_raw, w1_shaft)
                noise_plot = np.array(
                    [np.nan if (v := _optional_interp(qv, q1_raw, db1_raw)) is None else v for qv in q_plot]
                )
                q2_scaled, p2_scaled, w2_scaled, db2_scaled, eff2_scaled = scale_derived_curve(
                    q_plot,
                    p_plot,
                    w_plot,
                    noise_plot,
                    sel_m["d1"],
                    sel_m["h1"],
                    sel_m["n1"],
                    sel_m["d2_val"],
                    sel_m["h2_val"],
                    sel_m["n2_val"],
                    air_density_1,
                    air_density_2,
                    use_eff_correction,
                    eff_cap,
                    c_q,
                    c_p,
                    c_w,
                    fan_category,
                    mu1=air_viscosity_1,
                    mu2=air_viscosity_2,
                )
                eff2_scaled = eff2_scaled * 100
                st.caption(
                    f"선택 후보: 기준 Q={sel_m['q_ref']:.2f} CMM · 유효 후보 {sel_m['후보 수']}개 · "
                    f"소음: {sel_m['소음 근거']}"
                )
                if not np.isfinite(sel_m["예측 소음 dB(A)"]):
                    st.info("목표점 소음은 유효 측정 범위를 벗어나거나 데이터가 없어 계산하지 않았습니다.")
                if abs(r_D - 1) > 0.2 or abs(r_N - 1) > 0.2:
                    st.caption(
                        "형상 또는 RPM 변화가 ±20%를 넘습니다. 후보는 상사 모델 예측이며 추가 시험으로 확인하십시오."
                    )

                # Plotting
                fig_t6 = make_subplots(specs=[[{"secondary_y": True}]])
                fig_t6.add_trace(
                    go.Scatter(
                        x=q1_raw,
                        y=p1_raw,
                        mode="lines+markers",
                        name="원형 실험 PQ",
                        line=dict(color="#64748B", width=1.5, dash="dot"),
                    ),
                    secondary_y=False,
                )
                fig_t6.add_trace(
                    go.Scatter(
                        x=q2_scaled,
                        y=p2_scaled,
                        mode="lines+markers",
                        name="예측보정 PQ",
                        line=dict(color="#0D9488", width=2.5),
                    ),
                    secondary_y=False,
                )

                fig_t6.add_trace(
                    go.Scatter(
                        x=q1_raw,
                        y=eff1_raw * 100.0,
                        mode="lines+markers",
                        name="원형 정압효율",
                        line=dict(color="#A7C4BC", width=1.2, dash="dot"),
                    ),
                    secondary_y=True,
                )
                fig_t6.add_trace(
                    go.Scatter(
                        x=q2_scaled,
                        y=eff2_scaled,
                        mode="lines+markers",
                        name="예측보정 정압효율",
                        line=dict(color="#0F766E", width=2.0),
                    ),
                    secondary_y=True,
                )

                # Duct resistance line
                q_sys_plot_t6 = np.linspace(0, max(q2_scaled.max(), q1_raw.max()) * 1.15, 100)
                p_sys_plot_t6 = (t6_p_req / t6_q_req**2) * (q_sys_plot_t6**2)
                fig_t6.add_trace(
                    go.Scatter(
                        x=q_sys_plot_t6,
                        y=p_sys_plot_t6,
                        mode="lines",
                        name="요구 저항곡선",
                        line=dict(color="#2563EB", width=1.2, dash="dash"),
                    ),
                    secondary_y=False,
                )

                # Markers
                fig_t6.add_trace(
                    go.Scatter(
                        x=[t6_q_req],
                        y=[t6_p_req],
                        mode="markers",
                        marker=dict(color="#0D9488", size=14, symbol="star"),
                        name="목표 운전점",
                    ),
                    secondary_y=False,
                )

                fig_t6.update_layout(
                    title=dict(
                        text=f"[{selected_model_name}] 스케일링 예측 성능선도 (이중 축 효율 포함)",
                        font=dict(
                            family="Inter, Pretendard, Malgun Gothic, sans-serif", size=16, color="#0F172A"
                        ),
                    ),
                    xaxis_title="풍량 (CMM)",
                    plot_bgcolor="#FFFFFF",
                    paper_bgcolor="rgba(0,0,0,0)",
                    legend=dict(
                        x=0.02, y=0.98, bordercolor="#E2E8F0", borderwidth=1, bgcolor="rgba(255,255,255,0.7)"
                    ),
                    height=420,
                    margin=dict(l=40, r=40, t=55, b=40),
                )
                fig_t6.update_xaxes(showgrid=True, gridcolor="#EEF2F6", linecolor="#E2E8F0")
                fig_t6.update_yaxes(
                    title_text="정압 (Pa)",
                    showgrid=True,
                    gridcolor="#EEF2F6",
                    linecolor="#E2E8F0",
                    secondary_y=False,
                )
                fig_t6.update_yaxes(
                    title_text="정압효율 (%)", range=[0, 100], showgrid=False, secondary_y=True
                )
                _show_chart(fig_t6, width="stretch")

                # Power and Noise Curves for selected
                col_sd1, col_sd2 = st.columns(2)
                with col_sd1:
                    fig_w = go.Figure()
                    fig_w.add_trace(
                        go.Scatter(
                            x=q1_raw,
                            y=w1_shaft,
                            mode="lines+markers",
                            name="원형 실험 축동력",
                            line=dict(color="#64748B", dash="dot"),
                        )
                    )
                    fig_w.add_trace(
                        go.Scatter(
                            x=q2_scaled,
                            y=w2_scaled,
                            mode="lines+markers",
                            name="예측보정 축동력",
                            line=dict(color="#0D9488"),
                        )
                    )
                    fig_w.add_trace(
                        go.Scatter(
                            x=[t6_q_req],
                            y=[sel_m["예측 축동력 (W)"]],
                            mode="markers",
                            marker=dict(color="#0D9488", size=11, symbol="diamond"),
                            name="목표점 동력",
                        )
                    )
                    fig_w.update_layout(
                        title="풍량 대 축동력 곡선 비교",
                        xaxis_title="풍량 (CMM)",
                        yaxis_title="축동력 (W)",
                        plot_bgcolor="#FFFFFF",
                        paper_bgcolor="rgba(0,0,0,0)",
                        margin=dict(l=40, r=20, t=40, b=40),
                        height=300,
                    )
                    fig_w.update_xaxes(showgrid=True, gridcolor="#EEF2F6", linecolor="#E2E8F0")
                    fig_w.update_yaxes(showgrid=True, gridcolor="#EEF2F6", linecolor="#E2E8F0")
                    _show_chart(fig_w, width="stretch")
                with col_sd2:
                    fig_db = go.Figure()
                    fig_db.add_trace(
                        go.Scatter(
                            x=q1_raw,
                            y=db1_raw,
                            mode="lines+markers",
                            name="원형 실험 소음",
                            line=dict(color="#64748B", dash="dot"),
                        )
                    )
                    fig_db.add_trace(
                        go.Scatter(
                            x=q2_scaled,
                            y=db2_scaled,
                            mode="lines+markers",
                            name="예측보정 소음",
                            line=dict(color="#D97706"),
                        )
                    )
                    fig_db.add_trace(
                        go.Scatter(
                            x=[t6_q_req],
                            y=[sel_m["예측 소음 dB(A)"]],
                            mode="markers",
                            marker=dict(color="#D97706", size=11, symbol="diamond"),
                            name="목표점 소음",
                        )
                    )
                    fig_db.update_layout(
                        title="풍량 대 소음 레벨 곡선 비교",
                        xaxis_title="풍량 (CMM)",
                        yaxis_title="소음 레벨 dB(A)",
                        plot_bgcolor="#FFFFFF",
                        paper_bgcolor="rgba(0,0,0,0)",
                        margin=dict(l=40, r=20, t=40, b=40),
                        height=300,
                    )
                    fig_db.update_xaxes(showgrid=True, gridcolor="#EEF2F6", linecolor="#E2E8F0")
                    fig_db.update_yaxes(showgrid=True, gridcolor="#EEF2F6", linecolor="#E2E8F0")
                    _show_chart(fig_db, width="stretch")

                # AI Matching statement
                st.markdown("#### AI 제작 적합성 및 계통 매칭 제언")
                v_tip = sel_m["팁속도 (m/s)"]
                sigma = sel_m["비속도 (σ)"]
                mfg_lbl = sel_m["제작 공법"]

                adv_cols = st.columns([1, 2])
                with adv_cols[0]:
                    st.metric("목표 설계 비속도 (σ)", f"{sigma:.3f}")
                    st.metric("목표 팁 주속도 (V_tip)", f"{v_tip:.1f} m/s")
                with adv_cols[1]:
                    if "injection" in mfg and sigma > 1.2:
                        st.error(
                            f"⚠️ **금형 기술 경고 (High Specific Speed vs 2D Mold)**\n\n"
                            f"현재 목표 설계점의 비속도(σ = {sigma:.2f})는 고유동 축류/혼류 대역에 해당하여 **비틀림(Twist) 깃 설계가 강력히 요구**됩니다. "
                            f"2D 한방 사출 공법은 금형 탈형 각도로 인해 비틀림 각을 구현할 수 없어 날개 선단 및 후단부에서 극심한 박리가 일어나며, "
                            f"이 경우 실제 효율은 설정된 한계치({eff_cap*100:.0f}%)보다도 더 낮아질 위험이 있습니다. 슬라이드 금형(2D) 또는 3D 형상 적용을 추천합니다."
                        )
                    elif "injection" in mfg and eff1_raw.max() > eff_cap:
                        st.warning(
                            f"⚠️ **효율 한계 패널티 검출**\n\n"
                            f"입력하신 실험 모태 프로토타입의 최대 효율은 {eff1_raw.max()*100:.1f}%로 우수하지만, "
                            f"금형 방식으로 선택하신 '{mfg_lbl}'의 형상적 마찰/누설 한계({eff_cap*100:.0f}%)로 인해, "
                            f"스케일업 시 효율 개선 효과를 제한(Cap={eff_cap*100:.0f}%)하고 손실을 가중 반영하여 안전 마진을 반영했습니다."
                        )
                    else:
                        st.success(
                            f"✅ **제조 공법 적합성 통과**\n\n"
                            f"현재 기동 비속도(σ = {sigma:.2f}) 및 금형 구조는 물리적 운전 대역에 부합합니다. "
                            f"레이놀즈 수 변화에 따른 효율 개선(Cap: {eff_cap*100:.0f}% 이내)이 정상적으로 스케일링 엔진에 반영되었습니다."
                        )
            else:
                st.warning(
                    "⚠️ 요구 설계 및 최대 치수/속도 제약 조건을 만족하는 실험 모델을 데이터베이스 내에서 찾을 수 없습니다. 제약 조건이나 목표 요구 성능점을 완화하십시오."
                )

        except Exception as err_t6:
            st.error(
                f"실측 데이터베이스 분석 중 에러 발생: {err_t6}. 로드된 파일의 데이터 포맷과 컬럼 매핑이 정확한지 확인하십시오."
            )


# =============================================================================
# TAB 3: 기준팬 1:1 비교
# =============================================================================
def render_tab_comparison(settings):
    """1:1 비교 화면. 전역 입력은 settings로 전달한다."""
    fan_category = settings["fan_category"]
    is_axial = settings["is_axial"]
    dim_label = settings["dim_label"]
    dim_key = settings["dim_key"]
    flow_unit = settings["flow_unit"]
    pressure_unit = settings["pressure_unit"]
    power_unit = settings["power_unit"]
    flow_unit_disp = settings["flow_unit_disp"]
    pressure_unit_disp = settings["pressure_unit_disp"]
    power_unit_disp = settings["power_unit_disp"]
    air_density_1 = settings["air_density_1"]
    air_density_2 = settings["air_density_2"]
    air_viscosity_1 = settings["air_viscosity_1"]
    air_viscosity_2 = settings["air_viscosity_2"]
    use_eff_correction = settings["use_eff_correction"]

    st.markdown("### 기준 팬과 목표 스케일 팬의 1:1 다이렉트 상사 예측")
    st.markdown(
        "확실한 설계점(기준점)을 아는 경우, 치수와 RPM 변화에 다른 한 개의 목표 운전점 값의 변화를 일대일로 비교 분석합니다."
    )

    col3_1, col3_2 = st.columns(2)

    with col3_1:
        with st.container(border=True, key="panel_4833"):
            st.markdown("#### 1. 기준 팬 (Baseline Fan)")

            subc_b1, subc_b2 = st.columns(2)
            with subc_b1:
                base_D = st.number_input(
                    "기준 직경 D1 (mm)", value=300.0, min_value=10.0, step=10.0, key="b3_d1", format="%.0f"
                )
                if is_axial:
                    st.text_input("기준 독립 높이 H1", value="미적용", disabled=True, key="b3_x1_axial")
                    base_X = base_D
                else:
                    base_X = st.number_input(
                        f"기준 {dim_label}1 (mm)",
                        value=100.0,
                        min_value=5.0,
                        step=5.0,
                        key="b3_x1",
                        format="%.0f",
                    )
                base_N = st.number_input(
                    "기준 회전수 N1 (RPM)",
                    value=1500.0,
                    min_value=10.0,
                    step=50.0,
                    key="b3_n1",
                    format="%.0f",
                )
            with subc_b2:
                base_Q = unit_number_input(
                    f"기준 풍량 Q1 ({flow_unit_disp})",
                    value=30.0,
                    min_value=0.1,
                    step=1.0,
                    key="b3_q1",
                    quantity="flow",
                    unit=flow_unit,
                )
                base_P = unit_number_input(
                    f"기준 정압 P1 ({pressure_unit_disp})",
                    value=20.0,
                    min_value=0.1,
                    step=1.0,
                    key="b3_p1",
                    quantity="pressure",
                    unit=pressure_unit,
                )
                base_W = unit_number_input(
                    f"기준 동력 W1 ({power_unit_disp})",
                    value=150.0,
                    min_value=0.1,
                    step=10.0,
                    key="b3_w1",
                    quantity="power",
                    unit=power_unit,
                )

            base_noise = st.number_input(
                "기준 소음 Lw1 (dB(A))", value=60.0, min_value=10.0, max_value=150.0, step=1.0, key="b3_db1"
            )

    with col3_2:
        with st.container(border=True, key="panel_4853"):
            st.markdown("#### 2. 목표 팬 설계안 (Target Fan Design)")

            target_D = st.number_input(
                "목표 직경 D2 (mm)", value=350.0, min_value=10.0, step=10.0, key="b3_d2", format="%.0f"
            )
            if is_axial:
                st.text_input("목표 독립 높이 H2", value="미적용", disabled=True, key="b3_x2_axial")
                target_X = target_D
            else:
                target_X = st.number_input(
                    f"목표 {dim_label}2 (mm)",
                    value=120.0,
                    min_value=5.0,
                    step=5.0,
                    key="b3_x2",
                    format="%.0f",
                )

            target_input_mode = st.selectbox(
                "목표 사양 입력값 선택",
                ["목표 회전수 N2 (RPM) 지정", "목표 정압 P2 지정", "목표 풍량 Q2 지정"],
                key="b3_input_mode",
            )

            if target_input_mode == "목표 회전수 N2 (RPM) 지정":
                target_N = st.number_input(
                    "목표 회전수 N2 (RPM)",
                    value=1700.0,
                    min_value=10.0,
                    step=50.0,
                    key="b3_n2",
                    format="%.0f",
                )
            elif target_input_mode == "목표 정압 P2 지정":
                target_P_in = unit_number_input(
                    f"목표 정압 P2 ({pressure_unit_disp})",
                    value=25.0,
                    min_value=0.1,
                    step=1.0,
                    key="b3_p2_in",
                    quantity="pressure",
                    unit=pressure_unit,
                )
            else:  # 목표 풍량 Q2 지정
                target_Q_in = unit_number_input(
                    f"목표 풍량 Q2 ({flow_unit_disp})",
                    value=35.0,
                    min_value=0.1,
                    step=1.0,
                    key="b3_q2_in",
                    quantity="flow",
                    unit=flow_unit,
                )

            st.markdown("##### 계산 적용 옵션")
            base_eff_chk = float(
                static_efficiency(
                    convert_flow(base_Q, flow_unit, "CMM (m³/min)"),
                    convert_pressure(base_P, pressure_unit, "Pa"),
                    convert_power(base_W, power_unit, "W"),
                )
            )
            st.metric("기준 정압효율 · Q·P/W", f"{base_eff_chk * 100:.2f}%")
            st.caption("입력한 풍량·정압·축동력으로 효율 보정 기준을 계산합니다.")

    if not 0 < base_eff_chk <= 1:
        st.error("기준 Q·P/W 정압효율은 0% 초과, 100% 이하여야 합니다. 단위와 축동력을 확인하십시오.")
        return

    # Scale calculation
    q_b_base = convert_flow(base_Q, flow_unit, "CMM (m³/min)")
    p_b_base = convert_pressure(base_P, pressure_unit, "Pa")
    w_b_base = convert_power(base_W, power_unit, "W")

    # Solve N2_calc based on selected target input mode
    if target_input_mode == "목표 회전수 N2 (RPM) 지정":
        target_N_calc = target_N
    elif target_input_mode == "목표 정압 P2 지정":
        p_t_target_pa = convert_pressure(target_P_in, pressure_unit, "Pa")
        if p_b_base > 0 and p_t_target_pa >= 0:
            target_N_calc = (
                base_N
                * np.sqrt(p_t_target_pa / p_b_base)
                * (base_D / target_D)
                * np.sqrt(air_density_1 / air_density_2)
            )
        else:
            target_N_calc = base_N
            st.warning(
                "경고: 기준 정압 또는 목표 정압 조건이 유효하지 않아 회전수 계산에 기준 회전수(N1)를 적용했습니다."
            )
    else:  # 목표 풍량 Q2 지정
        q_t_target_cmm = convert_flow(target_Q_in, flow_unit, "CMM (m³/min)")
        if q_b_base > 0 and q_t_target_cmm >= 0:
            if is_axial:
                target_N_calc = base_N * (q_t_target_cmm / q_b_base) * ((base_D / target_D) ** 3)
            else:
                target_N_calc = (
                    base_N * (q_t_target_cmm / q_b_base) * ((base_D / target_D) ** 2) * (base_X / target_X)
                )
        else:
            target_N_calc = base_N
            st.warning(
                "경고: 기준 풍량 또는 목표 풍량 조건이 유효하지 않아 회전수 계산에 기준 회전수(N1)를 적용했습니다."
            )

    q_t_sc, p_t_sc, w_t_sc, db_t_sc, t_eff = run_similarity_scaling(
        q_b_base,
        p_b_base,
        w_b_base,
        base_noise,
        base_D,
        base_X,
        base_N,
        target_D,
        target_X,
        target_N_calc,
        air_density_1,
        air_density_2,
        use_eff_correction,
        base_eff_chk,
        fan_category,
        mu1=air_viscosity_1,
        mu2=air_viscosity_2,
    )

    target_Q_res = convert_flow(q_t_sc, "CMM (m³/min)", flow_unit)
    target_P_res = convert_pressure(p_t_sc, "Pa", pressure_unit)
    target_W_res = convert_power(w_t_sc, "W", power_unit)

    # Deltas
    n_del = ((target_N_calc / base_N) - 1.0) * 100
    q_del = ((target_Q_res / base_Q) - 1.0) * 100
    p_del = ((target_P_res / base_P) - 1.0) * 100
    w_del = ((target_W_res / base_W) - 1.0) * 100
    db_del = db_t_sc - base_noise

    n_arrow, n_cls = delta_arrow_class(n_del, higher_is_bad=True)
    q_arrow3, q_cls3 = delta_arrow_class(q_del, higher_is_bad=False)
    p_arrow3, p_cls3 = delta_arrow_class(p_del, higher_is_bad=False)
    w_arrow3, w_cls3 = delta_arrow_class(w_del, higher_is_bad=True)
    db_arrow3, db_cls3 = delta_arrow_class(db_del, higher_is_bad=True)

    with st.container(border=True, key="panel_4929"):
        st.markdown("#### 상사법칙 예측 일대일 결과 요약")

        res_col1, res_col2, res_col3, res_col4, res_col5 = st.columns(5)
        with res_col1:
            st.markdown(
                f"""
            <div class="metric-box">
                <div class="title">목표 회전수 (N2)</div>
                <div class="value">{target_N_calc:.0f} RPM</div>
                <div class="delta {n_cls}">{n_arrow} {n_del:+.1f}% (기준: {base_N:.0f})</div>
            </div>
            """,
                unsafe_allow_html=True,
            )
        with res_col2:
            st.markdown(
                f"""
            <div class="metric-box">
                <div class="title">목표 풍량 (Q2)</div>
                <div class="value">{target_Q_res:.2f} <span style="font-size: 1.0rem; font-weight: normal; color: #64748B;">{flow_unit_disp}</span></div>
                <div class="delta {q_cls3}">{q_arrow3} {q_del:+.1f}% <span style="font-size: 0.8rem; font-weight: normal; color: #64748B;">(기준: {base_Q:.2f})</span></div>
            </div>
            """,
                unsafe_allow_html=True,
            )
        with res_col3:
            st.markdown(
                f"""
            <div class="metric-box">
                <div class="title">목표 정압 (P2)</div>
                <div class="value">{target_P_res:.2f} <span style="font-size: 1.0rem; font-weight: normal; color: #64748B;">{pressure_unit_disp}</span></div>
                <div class="delta {p_cls3}">{p_arrow3} {p_del:+.1f}% <span style="font-size: 0.8rem; font-weight: normal; color: #64748B;">(기준: {base_P:.2f})</span></div>
            </div>
            """,
                unsafe_allow_html=True,
            )
        with res_col4:
            st.markdown(
                f"""
            <div class="metric-box">
                <div class="title">목표 동력 (W2)</div>
                <div class="value">{target_W_res:.1f} <span style="font-size: 1.0rem; font-weight: normal; color: #64748B;">{power_unit_disp}</span></div>
                <div class="delta {w_cls3}">{w_arrow3} {w_del:+.1f}% <span style="font-size: 0.8rem; font-weight: normal; color: #64748B;">(기준: {base_W:.1f})</span></div>
            </div>
            """,
                unsafe_allow_html=True,
            )
        with res_col5:
            st.markdown(
                f"""
            <div class="metric-box">
                <div class="title">목표 소음 (Lw2)</div>
                <div class="value">{db_t_sc:.1f} dB(A)</div>
                <div class="delta {db_cls3}">{db_arrow3} {db_del:+.1f} dB(A) (기준: {base_noise:.1f})</div>
            </div>
            """,
                unsafe_allow_html=True,
            )

        st.markdown("---")
        st.markdown("##### 물리 현상 해설 코멘터리")

        # Generate automatic structural physical commentary (단일 소스: build_scaling_commentary)
        commentary = build_scaling_commentary(
            base_D,
            target_D,
            base_X,
            target_X,
            base_N,
            target_N_calc,
            dim_label,
            use_eff_correction,
            base_eff_chk,
            t_eff,
            md=True,
            fan_category=fan_category,
        )
        for cm in commentary:
            st.write(cm)

        # Tab 3 AI Prescriptions & Statement (실계산 기반)
        st.markdown("---")
        st.markdown("#### AI 설계 스케일링 처방전 & 소견서")
        if is_axial:
            st.markdown(
                "축류팬에는 독립 폭/높이 자유도가 없으므로 직경과 RPM만 바꾸면서 풍량과 정압을 동시에 "
                "보존할 수 없습니다. 아래에서는 **정압·팁속도를 유지**하는 직경/RPM 조합과 그에 따른 "
                "풍량·동력·소음 변화를 실계산합니다."
            )
            ai_goal = "축류팬 정압 유지 검토"
        else:
            st.markdown(
                "현재 예측된 목표 팬을 기준으로, **운전점(풍량·정압)을 정확히 유지하는 듀티 보존 재조합**(P 유지 ⇒ N₃=N₂/r_D, Q 유지 ⇒ H₃=H₂/r_D)을 상사법칙으로 직접 계산하여 제안합니다. 아래 수치는 서술이 아니라 본 툴의 예측 엔진이 산출한 값입니다."
            )
            ai_goal = st.radio(
                "선호하는 AI 최적화 방향 선택",
                [
                    "동력 절감 최우선 (Maximize Efficiency)",
                    "소음 최소화 최우선 (Minimize Noise)",
                    "팬 크기 축소 최우선 (Downsizing Size)",
                ],
                horizontal=True,
                key="tab3_ai_goal",
            )

        v_tip_t = (np.pi * target_D * target_N_calc) / (60.0 * 1000.0)

        def _rebalance_live(r_d_prop):
            """현재 목표 팬(D2,X2,N2)에서 듀티 보존 재조합(D3=r_D·D2, N3=N2/r_D, H3=H2/r_D)의
            성능을 run_similarity_scaling으로 실계산하여 반환."""
            d3 = target_D * r_d_prop
            n3 = target_N_calc / r_d_prop
            x3 = target_X if is_axial else target_X / r_d_prop
            q3, p3, w3, db3, eff3 = run_similarity_scaling(
                q_t_sc,
                p_t_sc,
                w_t_sc,
                db_t_sc,
                target_D,
                target_X,
                target_N_calc,
                d3,
                x3,
                n3,
                air_density_2,
                air_density_2,
                use_eff_correction,
                t_eff,
                fan_category,
                mu1=air_viscosity_2,
                mu2=air_viscosity_2,
            )
            v_tip3 = (np.pi * d3 * n3) / (60.0 * 1000.0)
            return {
                "d3": d3,
                "n3": n3,
                "x3": x3,
                "dq": (q3 / q_t_sc - 1.0) * 100 if q_t_sc > 0 else 0.0,
                "dp": (p3 / p_t_sc - 1.0) * 100 if p_t_sc > 0 else 0.0,
                "dw": (w3 / w_t_sc - 1.0) * 100 if w_t_sc > 0 else 0.0,
                "ddb": db3 - db_t_sc,
                "deff": (eff3 - t_eff) * 100,
                "dvtip": (v_tip3 / v_tip_t - 1.0) * 100 if v_tip_t > 0 else 0.0,
                "eff3": eff3,
            }

        def _show_rebalance_table(rb):
            dimension_row = (
                ""
                if is_axial
                else f"| {dim_label}₃ | {rb['x3']:.1f} mm | {(rb['x3']/target_X-1)*100:+.1f}% |"
            )
            st.markdown(f"""
| 항목 | 제안값 | 목표 대비 변화 (실계산) |
|---|---|---|
| 직경 D₃ | {rb['d3']:.1f} mm | {(rb['d3']/target_D-1)*100:+.1f}% |
{dimension_row}
| 회전수 N₃ | {rb['n3']:.0f} RPM | {(rb['n3']/target_N_calc-1)*100:+.1f}% |
| 풍량 변화 | — | **{rb['dq']:+.2f}%** |
| 정압 변화 | — | **{rb['dp']:+.2f}%** |
| 축동력 변화 | — | **{rb['dw']:+.2f}%** |
| 정압효율 변화 | η = {rb['eff3']*100:.1f}% | **{rb['deff']:+.2f}%p** |
| 소음 변화 | — | **{rb['ddb']:+.2f} dB(A)** |
| 팁속도 변화 | — | {rb['dvtip']:+.2f}% |
""")

        if is_axial:
            r_d_ax = st.slider(
                "정압 유지 직경비 r_D",
                min_value=0.80,
                max_value=1.20,
                value=1.05,
                step=0.01,
                key="tab3_rd_axial",
            )
            rb = _rebalance_live(r_d_ax)
            st.info("💡 **축류팬 정압 유지안: N₃=N₂/r_D (팁속도 일정)**")
            _show_rebalance_table(rb)
            st.caption(
                f"축류팬 법칙상 정압은 유지되지만 풍량은 D²에 비례해 {rb['dq']:+.2f}% 변합니다. "
                f"따라서 동일 풍량까지 필요하면 블레이드 피치, 허브비 또는 복수 팬 구성 등 별도 설계 자유도가 필요합니다. "
                f"예측 소음 변화는 {rb['ddb']:+.2f} dB(A)입니다."
            )
        elif ai_goal == "동력 절감 최우선 (Maximize Efficiency)":
            r_d_up = st.slider(
                "직경 확대 비율 r_D", min_value=1.02, max_value=1.20, value=1.08, step=0.01, key="tab3_rd_up"
            )
            rb = _rebalance_live(r_d_up)
            st.info("💡 **설계 제안: 저속 대구경 듀티 보존 스케일링**")
            _show_rebalance_table(rb)
            st.caption(
                f"물리 해석: 듀티 보존 재조합에서 정압 유지 조건(P ∝ (N·D)²)에 의해 팁속도는 원리적으로 불변이며, "
                f"동력 절감분은 전적으로 레이놀즈수 상승에 따른 Ackeret 효율 개선(Δη = {rb['deff']:+.2f}%p)에서 나옵니다. "
                + (
                    ""
                    if use_eff_correction
                    else "⚠️ 현재 상단 설정에서 효율 보정이 꺼져 있어 개선분이 0으로 계산됩니다. 보정을 켜야 이 처방의 효과가 반영됩니다."
                )
            )
        elif ai_goal == "소음 최소화 최우선 (Minimize Noise)":
            st.info("💡 **모델 진단: 듀티 보존 재조합으로는 소음 저감 불가 (모델 항등성)**")
            rb = _rebalance_live(1.05)
            st.markdown(f"""
- **모델 검증 결과**: 본 툴의 소음 상사법칙 Lw₂ = Lw₁ + 50log(N) + 60log(D) + 10log(H)는 KN 모델(Lw = KN + 10logQ + 20logP)과 수학적으로 등가입니다. 따라서 **운전점(Q, P)을 유지하는 어떤 D/N/H 재조합도 예측 소음을 바꾸지 못합니다** — 예컨대 D+5% 듀티 보존 재조합의 실계산 결과는 ΔLw = {rb['ddb']:+.2f} dB(A)로 사실상 0입니다. 마찬가지로 정압 유지 조건에서 팁속도(∝ N·D)도 불변입니다.
- **실효적인 소음 저감 경로 (실계산)**:
  1. **요구 운전점 하향**: 계통 저항을 저감(덕트 확관, 곡관 저감 등)해 동일 유량을 더 낮은 정압으로 달성하거나, 요구 유량 자체를 낮추십시오. 동일 팬 회전수 5% 감속 시 ΔLw = {50*np.log10(0.95):+.2f} dB(A) (Q −5.0%, P −9.75% 동반).
  2. **특성 소음 계수(KN) 자체 개선**: 블레이드 형상 최적화(입출구각, 캠버), 컷오프 간극 확대, 세레이션/바이오미미크리 트레일링엣지 등 설계 품질 개선은 KN을 낮추는 경로로, 상사법칙 범위 밖의 실험/CFD 검증 영역입니다.
""")
        else:  # 크기 축소 최우선
            r_d_dn = st.slider(
                "직경 축소 비율 r_D", min_value=0.80, max_value=0.98, value=0.90, step=0.01, key="tab3_rd_dn"
            )
            rb = _rebalance_live(r_d_dn)
            st.info("💡 **설계 제안: 고속 소구경 듀티 보존 스케일링**")
            _show_rebalance_table(rb)
            st.caption(
                f"물리 해석: 직경 축소분만큼 회전수·폭을 상향해 운전점을 유지합니다. 예측 소음 변화는 ΔLw = {rb['ddb']:+.2f} dB(A)로 본 모델상 미미하나, "
                f"레이놀즈수 하락에 의한 효율 저하(Δη = {rb['deff']:+.2f}%p)로 축동력이 {rb['dw']:+.2f}% 변화합니다. "
                f"단, 폭({dim_label}) 확대 비율이 커질수록 b/D 비가 원형 설계에서 멀어져 상사 가정의 유효성이 저하되므로 ±20~30% 이내 스케일을 권장합니다."
            )

        st.markdown("##### 제안서용 공학적 소견서 (복사하여 활용 가능)")

        # 소견서용 σ 실계산 및 밴드 판정 (조건부 문구)
        _q_t_m3s = q_t_sc / 60.0
        _y_head_t = p_t_sc / air_density_2 if air_density_2 > 0 else 0.0
        _omega_t = target_N_calc * 2.0 * np.pi / 60.0
        if _y_head_t > 0 and _q_t_m3s > 0:
            sigma_t3 = _omega_t * np.sqrt(_q_t_m3s) / (_y_head_t**0.75)
        else:
            sigma_t3 = 0.0
        if fan_category == CENTRIFUGAL_CATEGORY:
            _lo3, _hi3 = 0.15, 1.80
        elif fan_category == AXIAL_CATEGORY:
            _lo3, _hi3 = 1.50, 8.00
        else:
            _lo3, _hi3 = 0.05, 0.35

        if sigma_t3 == 0.0:
            sigma_stmt = "설계점 비속도는 입력 조건 미비로 판정하지 못했습니다."
        elif _lo3 <= sigma_t3 <= _hi3:
            sigma_stmt = f"설계점 무차원 비속도 σ = {sigma_t3:.3f}은 선택된 {fan_category}의 대표 효율 밴드(σ≈{_lo3}–{_hi3}) 내에 부합합니다."
        else:
            sigma_stmt = f"설계점 무차원 비속도 σ = {sigma_t3:.3f}은 선택된 {fan_category}의 대표 효율 밴드(σ≈{_lo3}–{_hi3})를 벗어나므로, 기종 적합성 재검토 또는 회전수·직경 재배분을 권장합니다."

        if v_tip_t < 35.0:
            vtip_stmt = f"약 {v_tip_t:.1f} m/s로 저소음 안정 운전 영역에 위치합니다"
        elif v_tip_t < 55.0:
            vtip_stmt = f"약 {v_tip_t:.1f} m/s로 표준 산업 운전 영역에 위치합니다"
        else:
            vtip_stmt = f"약 {v_tip_t:.1f} m/s로 고속 운전 영역에 해당하여 소음·진동 대책 병행이 필요합니다"

        # Format the commentary based on current specs
        geometry_statement = (
            f"직경 {target_D:.1f} mm, 운전 속도 {target_N_calc:.0f} RPM(독립 높이 미적용)"
            if is_axial
            else f"직경 {target_D:.1f} mm, {dim_label} {target_X:.1f} mm, 운전 속도 {target_N_calc:.0f} RPM"
        )
        statement_text = f"""[기술 소견서]
본 송풍기 설계 스펙은 {geometry_statement}으로 구성되었으며, 요구 운전점에서 날개 끝 주변 속도(Tip Speed)는 {vtip_stmt}. 
상사 법칙 스케일링 기준에 따라 기존 팬 대비 {q_del:+.1f}%의 유량 편차와 {p_del:+.1f}%의 정압 편차를 추종하며, 소요 동력은 {target_W_res:.1f} {power_unit_disp} 수준으로 예측됩니다. 
{sigma_stmt} 공진 위험 회피를 위해 정격 RPM 주변의 구조 진동 댐핑 마진을 확보하는 것을 권장합니다."""

        st.text_area(
            "Engineering Statement (Korean)", value=statement_text, height=200, key="tab3_ai_commentary"
        )

        st.markdown("---")
        st.markdown("##### 공학 보고서 내보내기")
        st.markdown(
            "본 1:1 상사 비교 예측 결과 및 물리 수식이 포함된 파워포인트(PPT) 슬라이드 보고서를 즉시 생성하여 다운로드합니다."
        )

        # PPT용 코멘터리는 동일 소스에서 평문(md=False)으로 생성
        ppt_commentary = build_scaling_commentary(
            base_D,
            target_D,
            base_X,
            target_X,
            base_N,
            target_N_calc,
            dim_label,
            use_eff_correction,
            base_eff_chk,
            t_eff,
            md=False,
            fan_category=fan_category,
        )

        try:
            if st.button("1:1 비교 보고서 생성", key="tab3_generate_report", type="primary"):
                ppt_data = create_ppt_report(
                    fan_category=fan_category,
                    dim_label=dim_label,
                    base_D=base_D,
                    base_X=base_X,
                    base_N=base_N,
                    base_Q=base_Q,
                    base_P=base_P,
                    base_W=base_W,
                    base_noise=base_noise,
                    target_D=target_D,
                    target_X=target_X,
                    target_N_calc=target_N_calc,
                    target_Q_res=target_Q_res,
                    target_P_res=target_P_res,
                    target_W_res=target_W_res,
                    db_t_sc=db_t_sc,
                    flow_unit=flow_unit_disp,
                    pressure_unit=pressure_unit_disp,
                    power_unit=power_unit_disp,
                    air_density_1=air_density_1,
                    air_density_2=air_density_2,
                    use_eff_correction=use_eff_correction,
                    n_del=n_del,
                    q_del=q_del,
                    p_del=p_del,
                    w_del=w_del,
                    db_del=db_del,
                    commentary_list=ppt_commentary,
                )

                st.download_button(
                    label="📥 PPTX 보고서 다운로드 (PowerPoint Report)",
                    data=ppt_data,
                    file_name=f"fan_scaling_report_{datetime.date.today().strftime('%Y%m%d')}.pptx",
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    on_click="ignore",
                )
        except Exception as ppt_err:
            st.error(f"PPT 보고서 생성 중 에러 발생: {ppt_err}")


# =============================================================================
# TAB 4: 최적화 및 파라미터 스윕
# =============================================================================
def render_tab_sweep(settings):
    """설계 탐색 화면. 전역 입력은 settings로 전달한다."""
    fan_category = settings["fan_category"]
    is_axial = settings["is_axial"]
    dim_label = settings["dim_label"]
    dim_key = settings["dim_key"]
    flow_unit = settings["flow_unit"]
    pressure_unit = settings["pressure_unit"]
    power_unit = settings["power_unit"]
    flow_unit_disp = settings["flow_unit_disp"]
    pressure_unit_disp = settings["pressure_unit_disp"]
    power_unit_disp = settings["power_unit_disp"]
    air_density_1 = settings["air_density_1"]
    air_density_2 = settings["air_density_2"]
    air_viscosity_1 = settings["air_viscosity_1"]
    air_viscosity_2 = settings["air_viscosity_2"]
    use_eff_correction = settings["use_eff_correction"]

    st.markdown("### 최적화 및 설계 매개변수 다차원 스윕 (Parameter Sweep)")
    if is_axial:
        st.markdown(
            "축류팬의 독립 높이 변수를 제외하고 직경과 RPM을 연속 스윕하여 사양 변화 추이를 분석합니다."
        )
    else:
        st.markdown(
            "직경, 높이/길이, RPM 세 가지 주요 변수 중 두 가지를 지정한 범위에서 연속 스윕(Sweep)하여 사양 변화 추이를 2D 컨투어 및 3D 표면으로 분석합니다."
        )

    col4_1, col4_2 = st.columns([1, 2.2])

    with col4_1:
        with st.container(border=True, key="panel_5212"):
            st.markdown("#### 1. 스윕 고정 사양 및 스케일 기준")

            ref_D = st.number_input(
                "기준 직경 D1 (mm)", value=300.0, min_value=10.0, key="b4_d1", format="%.0f"
            )
            if is_axial:
                st.text_input("기준 독립 높이 H1", value="미적용", disabled=True, key="b4_x1_axial")
                ref_X = ref_D
            else:
                ref_X = st.number_input(
                    f"기준 {dim_label}1 (mm)", value=100.0, min_value=5.0, key="b4_x1", format="%.0f"
                )
            ref_N = st.number_input(
                "기준 회전수 N1 (RPM)", value=1500.0, min_value=10.0, key="b4_n1", format="%.0f"
            )

            ref_Q = unit_number_input(
                f"기준 풍량 Q1 ({flow_unit_disp})", value=30.0, key="b4_q1", quantity="flow", unit=flow_unit
            )
            ref_P = unit_number_input(
                f"기준 정압 P1 ({pressure_unit_disp})",
                value=20.0,
                key="b4_p1",
                quantity="pressure",
                unit=pressure_unit,
            )
            ref_W = unit_number_input(
                f"기준 동력 W1 ({power_unit_disp})",
                value=150.0,
                key="b4_w1",
                quantity="power",
                unit=power_unit,
            )
            ref_db = st.number_input("기준 소음 Lw1 (dB(A))", value=60.0, key="b4_db1")

            st.markdown("#### 2. 스윕 매개변수(X, Y축) 선택")
            sweep_options = (
                ["X축: 직경(D) / Y축: 회전수(N)"]
                if is_axial
                else [
                    "X축: 직경(D) / Y축: 회전수(N)",
                    f"X축: 직경(D) / Y축: {dim_label}",
                    f"X축: {dim_label} / Y축: 회전수(N)",
                ]
            )
            sweep_axes = st.selectbox(
                "X축 및 Y축 파라미터 구성", sweep_options, key=f"tab4_sweep_axes_{dim_key}"
            )

            # Range inputs based on axes selection
            if sweep_axes == "X축: 직경(D) / Y축: 회전수(N)":
                x_var = "D"
                y_var = "N"
                z_fixed = "X"
                if is_axial:
                    z_fixed_val = ref_X
                    st.caption("독립 높이 H는 스윕에서 제외됩니다.")
                else:
                    z_fixed_val = st.number_input(
                        f"고정 {dim_label}2 (mm)", value=100.0, step=5.0, format="%.0f"
                    )

                x_min = st.number_input("직경 D 범위 최소 (mm)", value=200.0, step=10.0, format="%.0f")
                x_max = st.number_input("직경 D 범위 최대 (mm)", value=400.0, step=10.0, format="%.0f")
                y_min = st.number_input("회전수 N 범위 최소 (RPM)", value=1000.0, step=100.0, format="%.0f")
                y_max = st.number_input("회전수 N 범위 최대 (RPM)", value=2000.0, step=100.0, format="%.0f")

            elif sweep_axes == f"X축: 직경(D) / Y축: {dim_label}":
                x_var = "D"
                y_var = "X"
                z_fixed = "N"
                z_fixed_val = st.number_input("고정 회전수 N2 (RPM)", value=1500.0, step=50.0, format="%.0f")

                x_min = st.number_input("직경 D 범위 최소 (mm)", value=200.0, step=10.0, format="%.0f")
                x_max = st.number_input("직경 D 범위 최대 (mm)", value=400.0, step=10.0, format="%.0f")
                y_min = st.number_input(f"{dim_label} 범위 최소 (mm)", value=50.0, step=5.0, format="%.0f")
                y_max = st.number_input(f"{dim_label} 범위 최대 (mm)", value=200.0, step=5.0, format="%.0f")

            else:  # X: X_dim, Y: N
                x_var = "X"
                y_var = "N"
                z_fixed = "D"
                z_fixed_val = st.number_input("고정 직경 D2 (mm)", value=300.0, step=10.0, format="%.0f")

                x_min = st.number_input(f"{dim_label} 범위 최소 (mm)", value=50.0, step=5.0, format="%.0f")
                x_max = st.number_input(f"{dim_label} 범위 최대 (mm)", value=200.0, step=5.0, format="%.0f")
                y_min = st.number_input("회전수 N 범위 최소 (RPM)", value=1000.0, step=100.0, format="%.0f")
                y_max = st.number_input("회전수 N 범위 최대 (RPM)", value=2000.0, step=100.0, format="%.0f")

            st.markdown("#### 3. 모니터링 출력값 선택")
            output_to_show = st.selectbox(
                "시각화 대상 공력 성능",
                ["풍량 (Flow Rate)", "정압 (Static Pressure)", "축동력 (Shaft Power)", "소음 (Sound Level)"],
            )

    with col4_2:
        with st.container(border=True, key="panel_5285"):
            st.markdown(f"#### 📈 {output_to_show} 매개변수 스윕 공간 매핑")

            if x_min <= 0 or y_min <= 0 or z_fixed_val <= 0 or x_min >= x_max or y_min >= y_max:
                st.error("스윕 치수와 RPM은 양수이고, 각 최소값은 최대값보다 작아야 합니다.")
                return
            if ref_Q <= 0 or ref_P <= 0 or ref_W <= 0:
                st.error("기준 풍량·정압·동력은 0보다 커야 합니다.")
                return
            # Grid Generation
            x_grid = np.linspace(x_min, x_max, 25)
            y_grid = np.linspace(y_min, y_max, 25)

            X_mat, Y_mat = np.meshgrid(x_grid, y_grid)

            # Setup static factors

            # 격자 → (d2, x2, n2) 배열 매핑 (run_similarity_scaling이 배열 안전하므로 1회 벡터 호출)
            if x_var == "D" and y_var == "N":
                TD2, TN2 = X_mat, Y_mat
                TX2 = np.full_like(X_mat, z_fixed_val)
            elif x_var == "D" and y_var == "X":
                TD2, TX2 = X_mat, Y_mat
                TN2 = np.full_like(X_mat, z_fixed_val)
            else:  # x_var == "X" and y_var == "N"
                TX2, TN2 = X_mat, Y_mat
                TD2 = np.full_like(X_mat, z_fixed_val)

            # Convert inputs to base units (스칼라)
            qb = convert_flow(ref_Q, flow_unit, "CMM (m³/min)")
            pb = convert_pressure(ref_P, pressure_unit, "Pa")
            wb = convert_power(ref_W, power_unit, "W")
            eff_peak_ref = float(static_efficiency(qb, pb, wb))
            if not 0 < eff_peak_ref <= 1:
                st.error("기준 Q·P/W 정압효율은 0% 초과, 100% 이하여야 합니다. 단위와 축동력을 확인하십시오.")
                return
            if use_eff_correction:
                st.caption(f"기준 Q·P/W 정압효율 η₁={eff_peak_ref:.4f}를 사용해 각 조건의 효율을 보정합니다.")

            q_sc, p_sc, w_sc, db_sc, _ = run_similarity_scaling(
                qb,
                pb,
                wb,
                ref_db,
                ref_D,
                ref_X,
                ref_N,
                TD2,
                TX2,
                TN2,
                air_density_1,
                air_density_2,
                use_eff_correction,
                eff_peak_ref,
                fan_category,
                mu1=air_viscosity_1,
                mu2=air_viscosity_2,
            )

            # Convert back (배열 연산)
            Z_Q = convert_flow(q_sc, "CMM (m³/min)", flow_unit)
            Z_P = convert_pressure(p_sc, "Pa", pressure_unit)
            Z_W = convert_power(w_sc, "W", power_unit)
            Z_db = db_sc

            # Select active Z
            if output_to_show == "풍량 (Flow Rate)":
                Z_active = Z_Q
                colorbar_title = f"풍량 ({flow_unit})"
                colorscale = "Blues"
            elif output_to_show == "정압 (Static Pressure)":
                Z_active = Z_P
                colorbar_title = f"정압 ({pressure_unit})"
                colorscale = "Reds"
            elif output_to_show == "축동력 (Shaft Power)":
                Z_active = Z_W
                colorbar_title = f"동력 ({power_unit})"
                colorscale = "Viridis"
            else:
                Z_active = Z_db
                colorbar_title = "소음 dB(A)"
                colorscale = "Plasma"

            x_label_title = "직경 D2 (mm)" if x_var == "D" else f"{dim_label}2 (mm)"
            y_label_title = "회전수 N2 (RPM)" if y_var == "N" else f"{dim_label}2 (mm)"

            # 3D vs 2D plot selector
            chart_mode = st.radio(
                "시각화 차트 형태",
                ["2D 등고선 등고 맵 (2D Contour Plot)", "3D 표면 공간 맵 (3D Surface Plot)"],
                horizontal=True,
            )

            if chart_mode == "2D 등고선 등고 맵 (2D Contour Plot)":
                fig_sw = go.Figure(
                    data=go.Contour(
                        z=Z_active,
                        x=x_grid,
                        y=y_grid,
                        colorscale=colorscale,
                        colorbar=dict(title=colorbar_title),
                        contours=dict(showlabels=True, labelfont=dict(size=12, color="white")),
                    )
                )
                fig_sw.update_layout(
                    xaxis_title=x_label_title,
                    yaxis_title=y_label_title,
                    margin=dict(l=40, r=40, t=40, b=40),
                    height=450,
                )
            else:
                fig_sw = go.Figure(
                    data=[
                        go.Surface(
                            z=Z_active,
                            x=x_grid,
                            y=y_grid,
                            colorscale=colorscale,
                            colorbar=dict(title=colorbar_title),
                        )
                    ]
                )
                fig_sw.update_layout(
                    scene=dict(
                        xaxis_title=x_label_title, yaxis_title=y_label_title, zaxis_title=colorbar_title
                    ),
                    margin=dict(l=40, r=40, t=40, b=40),
                    height=450,
                )

            view_label = "2D 등고선" if chart_mode.startswith("2D") else "3D 표면"
            fig_sw.update_layout(title_text=f"{colorbar_title} 설계 맵 · {view_label}")
            _show_chart(fig_sw, width="stretch")

            st.markdown("##### 설계 최적 사양 타겟 검색기 (Target Search)")
            st.markdown(
                "목표 풍량과 압력을 만족하면서, **동력을 최소화**하거나 **소음을 최소화**하는 형상 설계점을 검색합니다."
            )

            c_t1, c_t2, c_t3 = st.columns(3)
            with c_t1:
                req_flow = unit_number_input(
                    f"필요 최소 풍량 ({flow_unit_disp})",
                    value=float(ref_Q * 1.1),
                    key="req_f",
                    quantity="flow",
                    unit=flow_unit,
                    value_is_display=True,
                )
            with c_t2:
                req_press = unit_number_input(
                    f"필요 최소 정압 ({pressure_unit_disp})",
                    value=float(ref_P * 1.1),
                    key="req_p",
                    quantity="pressure",
                    unit=pressure_unit,
                    value_is_display=True,
                )
            with c_t3:
                opt_criteria = st.selectbox(
                    "최적화 지향점", ["축동력 최소화 (Min Power)", "소음 최소화 (Min Noise)"]
                )

            # Find candidates within grid that satisfy requirements
            valid_indices = np.where((Z_Q >= req_flow) & (Z_P >= req_press))

            if len(valid_indices[0]) > 0:
                candidates = []
                for idx, jdy in zip(valid_indices[0], valid_indices[1]):
                    # shape mapping
                    if x_var == "D" and y_var == "N":
                        td = x_grid[jdy]
                        tn = y_grid[idx]
                        tx = z_fixed_val
                    elif x_var == "D" and y_var == "X":
                        td = x_grid[jdy]
                        tx = y_grid[idx]
                        tn = z_fixed_val
                    else:
                        tx = x_grid[jdy]
                        tn = y_grid[idx]
                        td = z_fixed_val

                    candidate = {
                        "직경 D2 (mm)": round(td, 1),
                        "회전수 N2 (RPM)": round(tn, 0),
                        f"예측 풍량 ({flow_unit})": round(Z_Q[idx, jdy], 2),
                        f"예측 정압 ({pressure_unit})": round(Z_P[idx, jdy], 2),
                        f"축동력 ({power_unit})": round(Z_W[idx, jdy], 1),
                        "소음 (dB(A))": round(Z_db[idx, jdy], 1),
                    }
                    if not is_axial:
                        candidate[f"{dim_label}2 (mm)"] = round(tx, 1)
                    candidates.append(candidate)

                df_candidates = pd.DataFrame(candidates)
                if opt_criteria == "축동력 최소화 (Min Power)":
                    df_sorted = df_candidates.sort_values(by=f"축동력 ({power_unit})").head(5)
                else:
                    df_sorted = df_candidates.sort_values(by="소음 (dB(A))").head(5)

                st.write("**추천 최적 설계 조합 Top 5:**")
                st.dataframe(df_sorted, width="stretch", hide_index=True)
            else:
                st.info(
                    "현재 스윕 설계 범위 내에서 입력하신 필요 유량과 정압 조건을 모두 충족하는 조합이 존재하지 않습니다. 스윕 범위를 넓혀 보십시오."
                )


# =============================================================================
# TAB 5: 덕트 압력 손실 계산기 (Duct Pressure Loss Calculator)
# =============================================================================
def render_tab_duct(settings):
    """덕트 손실 화면. 전역 입력은 settings로 전달한다."""
    fan_category = settings["fan_category"]
    is_axial = settings["is_axial"]
    dim_label = settings["dim_label"]
    dim_key = settings["dim_key"]
    flow_unit = settings["flow_unit"]
    pressure_unit = settings["pressure_unit"]
    power_unit = settings["power_unit"]
    flow_unit_disp = settings["flow_unit_disp"]
    pressure_unit_disp = settings["pressure_unit_disp"]
    power_unit_disp = settings["power_unit_disp"]
    air_density_1 = settings["air_density_1"]
    air_density_2 = settings["air_density_2"]
    air_viscosity_1 = settings["air_viscosity_1"]
    air_viscosity_2 = settings["air_viscosity_2"]
    use_eff_correction = settings["use_eff_correction"]

    st.markdown("### 유체역학식 기반 배관/덕트 정압 손실 계산기")
    st.markdown(
        "덕트의 형상(원형/사각)과 재질(마찰계수), 그리고 곡로 등의 장애 조건에 맞춰 배관의 압력 손실(정압 강하)을 상세히 계산할 수 있는 독립형 엔지니어링 유틸리티 툴입니다."
    )

    with st.container(border=True, key="panel_5469"):
        st.markdown("#### 1. 유량 및 기하학적 설계 매개변수")

        calc_q = unit_number_input(
            f"설계 풍량 Q ({flow_unit_disp})",
            value=35.0,
            min_value=0.1,
            step=1.0,
            key="tab5_q",
            quantity="flow",
            unit=flow_unit,
        )

        col_shape1, col_shape2 = st.columns([1, 2])
        with col_shape1:
            duct_shape = st.selectbox(
                "덕트 형상 선택", ["원형 덕트 (Circular)", "사각 덕트 (Rectangular)"], key="tab5_duct_shape"
            )
        with col_shape2:
            if duct_shape == "원형 덕트 (Circular)":
                duct_dia = st.number_input(
                    "덕트 내경 D (mm)",
                    value=250.0,
                    min_value=10.0,
                    step=10.0,
                    key="tab5_duct_dia",
                    format="%.0f",
                )
                duct_area = np.pi * ((duct_dia / 1000.0) ** 2) / 4.0
                D_h = duct_dia / 1000.0
            else:
                col_rect1, col_rect2 = st.columns(2)
                with col_rect1:
                    duct_w = st.number_input(
                        "덕트 가로 폭 W (mm)",
                        value=300.0,
                        min_value=10.0,
                        step=10.0,
                        key="tab5_duct_w",
                        format="%.0f",
                    )
                with col_rect2:
                    duct_h = st.number_input(
                        "덕트 세로 높이 H (mm)",
                        value=200.0,
                        min_value=10.0,
                        step=10.0,
                        key="tab5_duct_h",
                        format="%.0f",
                    )
                duct_area = (duct_w / 1000.0) * (duct_h / 1000.0)
                D_h = (2.0 * duct_w * duct_h) / (duct_w + duct_h) / 1000.0

        st.markdown("---")
        st.markdown("#### 2. 배관 연장 및 마찰 계수 설정")
        col_rough1, col_rough2 = st.columns(2)
        with col_rough1:
            duct_len = st.number_input(
                "덕트 직관 길이 L (m)", value=10.0, min_value=0.0, step=1.0, key="tab5_duct_len"
            )
        with col_rough2:
            roughness_sel = st.selectbox(
                "덕트 내부 조도 (ε)",
                [
                    "아연도철판 (Galvanized Iron - 0.15 mm)",
                    "상업용 강관 (Steel - 0.045 mm)",
                    "매끄러운 배관 (PVC/Plastic - 0.0015 mm)",
                    "수동 입력 (Custom)",
                ],
                key="tab5_roughness_sel",
            )

            if roughness_sel == "아연도철판 (Galvanized Iron - 0.15 mm)":
                eps = 0.15 / 1000.0
            elif roughness_sel == "상업용 강관 (Steel - 0.045 mm)":
                eps = 0.045 / 1000.0
            elif roughness_sel == "매끄러운 배관 (PVC/Plastic - 0.0015 mm)":
                eps = 0.0015 / 1000.0
            else:
                custom_rough = st.number_input(
                    "수동 조도 ε (mm)", value=0.15, min_value=0.0001, format="%.4f", key="tab5_custom_rough"
                )
                eps = custom_rough / 1000.0

        st.markdown("---")
        st.markdown("#### 3. 부차적 손실 및 피팅류 구성")

        st.markdown("##### 피팅류 및 계통 구성 요소별 저항계수(K) 설정")
        col_fit1, col_fit2, col_fit3 = st.columns(3)

        with col_fit1:
            st.markdown("**1) 곡관류 (Elbows)**")
            elbow_preset = st.selectbox(
                "곡관 타입",
                [
                    "표준 90도 곡관 (R/D=1.0) [K=0.30]",
                    "완만한 90도 곡관 (R/D=1.5) [K=0.18]",
                    "급격한 90도 직각 곡관 (Miter) [K=1.20]",
                    "표준 45도 곡관 [K=0.15]",
                    "안내깃(Vane) 설치 사각 곡관 [K=0.25]",
                    "없음 (None)",
                    "직접 입력 (Custom)",
                ],
                index=5,
                key="tab5_elbow_preset",
            )
            if elbow_preset == "없음 (None)":
                elbow_k = 0.0
                default_elbows = 0
            elif elbow_preset == "직접 입력 (Custom)":
                elbow_k = st.number_input(
                    "곡관 K값 직접 입력", value=0.30, min_value=0.0, step=0.05, key="tab5_elbow_k_custom"
                )
                default_elbows = 1
            else:
                elbow_presets_map = {
                    "표준 90도 곡관 (R/D=1.0) [K=0.30]": 0.30,
                    "완만한 90도 곡관 (R/D=1.5) [K=0.18]": 0.18,
                    "급격한 90도 직각 곡관 (Miter) [K=1.20]": 1.20,
                    "표준 45도 곡관 [K=0.15]": 0.15,
                    "안내깃(Vane) 설치 사각 곡관 [K=0.25]": 0.25,
                }
                elbow_k = elbow_presets_map[elbow_preset]
                default_elbows = 2
            elbows = st.number_input(
                "곡관 개수", value=default_elbows, min_value=0, step=1, key="tab5_elbows"
            )

        with col_fit2:
            st.markdown("**2) 입출구 및 단면 변화 (Reducers/Ports)**")
            fitting_preset = st.selectbox(
                "피팅 타입",
                [
                    "벨마우스 입구 (Bellmouth Inlet) [K=0.05]",
                    "직각 플랜지 입구 (Square Inlet) [K=0.50]",
                    "돌출 입구 (Borda Inlet) [K=0.80]",
                    "급격 확대관 (Sudden Expansion) [K=1.00]",
                    "급격 축소관 (Sudden Contraction) [K=0.50]",
                    "자유 토출구 (Free Discharge) [K=1.00]",
                    "없음 (None)",
                    "직접 입력 (Custom)",
                ],
                index=6,
                key="tab5_fitting_preset",
            )
            if fitting_preset == "없음 (None)":
                fitting_k = 0.0
                default_fittings = 0
            elif fitting_preset == "직접 입력 (Custom)":
                fitting_k = st.number_input(
                    "피팅 K값 직접 입력", value=0.50, min_value=0.0, step=0.05, key="tab5_fitting_k_custom"
                )
                default_fittings = 1
            else:
                fitting_presets_map = {
                    "벨마우스 입구 (Bellmouth Inlet) [K=0.05]": 0.05,
                    "직각 플랜지 입구 (Square Inlet) [K=0.50]": 0.50,
                    "돌출 입구 (Borda Inlet) [K=0.80]": 0.80,
                    "급격 확대관 (Sudden Expansion) [K=1.00]": 1.00,
                    "급격 축소관 (Sudden Contraction) [K=0.50]": 0.50,
                    "자유 토출구 (Free Discharge) [K=1.00]": 1.00,
                }
                fitting_k = fitting_presets_map[fitting_preset]
                default_fittings = 1
            fittings_qty = st.number_input(
                "피팅 개수", value=default_fittings, min_value=0, step=1, key="tab5_fittings_qty"
            )

        with col_fit3:
            st.markdown("**3) 댐퍼 및 밸브류 (Dampers/Valves)**")
            damper_preset = st.selectbox(
                "댐퍼/밸브 타입",
                [
                    "버터플라이 댐퍼 (완전 열림) [K=0.30]",
                    "버터플라이 댐퍼 (30도 닫힘) [K=5.00]",
                    "버터플라이 댐퍼 (45도 닫힘) [K=15.0]",
                    "멀티블레이드 댐퍼 (완전 열림) [K=0.50]",
                    "멀티블레이드 댐퍼 (45도 닫힘) [K=20.0]",
                    "체크 댐퍼 (Check Damper) [K=1.50]",
                    "없음 (None)",
                    "직접 입력 (Custom)",
                ],
                index=6,  # 없음 (None)
                key="tab5_damper_preset",
            )
            if damper_preset == "없음 (None)":
                damper_k = 0.0
                default_dampers = 0
            elif damper_preset == "직접 입력 (Custom)":
                damper_k = st.number_input(
                    "댐퍼 K값 직접 입력", value=0.50, min_value=0.0, step=0.05, key="tab5_damper_k_custom"
                )
                default_dampers = 1
            else:
                damper_presets_map = {
                    "버터플라이 댐퍼 (완전 열림) [K=0.30]": 0.30,
                    "버터플라이 댐퍼 (30도 닫힘) [K=5.00]": 5.00,
                    "버터플라이 댐퍼 (45도 닫힘) [K=15.0]": 15.0,
                    "멀티블레이드 댐퍼 (완전 열림) [K=0.50]": 0.50,
                    "멀티블레이드 댐퍼 (45도 닫힘) [K=20.0]": 20.0,
                    "체크 댐퍼 (Check Damper) [K=1.50]": 1.50,
                }
                damper_k = damper_presets_map[damper_preset]
                default_dampers = 1
            dampers_qty = st.number_input(
                "댐퍼/밸브 개수", value=default_dampers, min_value=0, step=1, key="tab5_dampers_qty"
            )

        K_total = (elbows * elbow_k) + (fittings_qty * fitting_k) + (dampers_qty * damper_k)

    # Calculations
    rho = air_density_2
    mu = air_viscosity_2

    q_cmm = convert_flow(calc_q, flow_unit, "CMM (m³/min)")
    q_m3s = q_cmm / 60.0

    if duct_area > 0:
        v_flow = q_m3s / duct_area
    else:
        v_flow = 0.0

    Re = (rho * v_flow * D_h) / mu if mu > 0 else 0

    if Re > 0:
        if Re < 2300:
            f_factor = 64.0 / Re
        else:
            rel_rough = eps / D_h
            term = (rel_rough / 3.7) ** 1.11 + 6.9 / Re
            if term > 0:
                inv_sqrt_f = -1.8 * np.log10(term)
                f_factor = (1.0 / inv_sqrt_f) ** 2
            else:
                f_factor = 0.02
    else:
        f_factor = 0.0

    dp_friction = f_factor * (duct_len / D_h) * (rho * (v_flow**2) / 2.0) if D_h > 0 else 0.0
    dp_minor = K_total * (rho * (v_flow**2) / 2.0)
    dp_total_pa = dp_friction + dp_minor
    sys_p_calc = convert_pressure(dp_total_pa, "Pa", pressure_unit)

    with st.container(border=True, key="panel_5664"):
        st.markdown("#### 유체역학식 기반 계산 결과 리포트")

        c_res1, c_res2, c_res3, c_res4 = st.columns(4)
        with c_res1:
            st.markdown(
                f"""
            <div class="metric-box">
                <div class="title">덕트 내 유속</div>
                <div class="value">{v_flow:.2f} <span style="font-size: 1.0rem; font-weight: normal; color: #64748B;">m/s</span></div>
                <div class="delta" style="color: #64748B;">면적: {duct_area:.3f} m²</div>
            </div>
            """,
                unsafe_allow_html=True,
            )
        with c_res2:
            st.markdown(
                f"""
            <div class="metric-box">
                <div class="title">레이놀즈 수 (Re)</div>
                <div class="value">{Re:.0f}</div>
                <div class="delta" style="color: #64748B;">{"층류 (Laminar)" if Re < 2300 else "난류 (Turbulent)"}</div>
            </div>
            """,
                unsafe_allow_html=True,
            )
        with c_res3:
            st.markdown(
                f"""
            <div class="metric-box">
                <div class="title">마찰 계수 (f)</div>
                <div class="value">{f_factor:.4f}</div>
                <div class="delta" style="color: #64748B;">Haaland 식 근사치</div>
            </div>
            """,
                unsafe_allow_html=True,
            )
        with c_res4:
            st.markdown(
                f"""
            <div class="metric-box">
                <div class="title">총 정압 강하</div>
                <div class="value">{sys_p_calc:.2f} <span style="font-size: 1.0rem; font-weight: normal; color: #64748B;">{pressure_unit_disp}</span></div>
                <div class="delta" style="color: #64748B;">직관마찰:{convert_pressure(dp_friction, "Pa", pressure_unit):.1f} / 피팅부차:{convert_pressure(dp_minor, "Pa", pressure_unit):.1f}</div>
            </div>
            """,
                unsafe_allow_html=True,
            )

        st.info(
            "💡 본 계산기는 독립적으로 가동되는 덕트 설계 도구입니다. 산출된 총 정압 강하 값을 참고하여 각 설계 탭의 정격 요구 정압에 기입하여 활용하십시오."
        )


def run_self_tests():
    """실행 파일 안의 계산 회귀 테스트. python fan_similarity_v10.py --self-test."""
    import unittest
    import tempfile
    from concurrent.futures import ThreadPoolExecutor
    from unittest.mock import patch

    class RegressionTests(unittest.TestCase):
        def test_chart_title_is_explicit_and_existing_title_is_preserved(self):
            with (
                patch.object(st, "session_state", {}),
                patch.object(st, "plotly_chart") as render,
                patch.object(st, "expander"),
                patch.object(st, "caption"),
                patch.object(st, "markdown") as markdown,
            ):
                _show_chart(go.Figure())
                serialized = render.call_args.args[0].to_plotly_json()
                self.assertEqual(serialized["layout"]["title"]["text"], "")
                markdown.assert_not_called()
                # 모바일: 기존 제목은 툴바와 겹치지 않도록 차트 위 HTML 제목으로 옮겨 보존한다.
                _show_chart(go.Figure(layout=dict(title_text="기존 차트 제목")))
                self.assertEqual(render.call_args.args[0].layout.title.text, "")
                self.assertIn("기존 차트 제목", markdown.call_args.args[0])

        def test_analytic_system_intersection(self):
            result = find_system_intersection([0, 100], [100, 0], 0.01)
            self.assertAlmostEqual(result, 61.80339887498949, places=10)
            self.assertAlmostEqual(100 - result, 0.01 * result**2, places=10)

        def test_operating_efficiency_matches_interpolated_power(self):
            q, p, w = [10, 30], [300, 100], [100, 100]
            op_q = find_system_intersection(q, p, 0.5)
            self.assertAlmostEqual(op_q, 20)
            op_eta = static_efficiency(op_q, 0.5 * op_q**2, np.interp(op_q, q, w))
            self.assertAlmostEqual(float(op_eta), 2 / 3)
            grid, eta = efficiency_curve_points(q, p, w, [op_q])
            self.assertAlmostEqual(float(eta[np.flatnonzero(grid == op_q)[0]]), 2 / 3)
            self.assertTrue(np.isnan(static_efficiency(20, 200, 0)))
            self.assertTrue(np.isnan(static_efficiency(20, 200, np.nan)))

        def test_final_efficiency_cap_after_all_coefficients(self):
            for cw, expected in [(0.8, 0.55), (1.0, 0.55), (1.5, 0.5)]:
                result = scale_derived_curve(
                    [30],
                    [300],
                    [200],
                    [60],
                    300,
                    100,
                    1500,
                    300,
                    100,
                    1500,
                    1.2,
                    1.2,
                    False,
                    0.55,
                    c_q=1.1,
                    c_p=0.9,
                    c_w=cw,
                )
                q, p, w, _, eta = result
                self.assertAlmostEqual(float(eta[0]), expected)
                np.testing.assert_allclose(eta, q * p / 60 / w)
                self.assertLessEqual(float(eta[0]), 0.55 + 1e-12)

        def test_low_efficiency_correction_conserves_power_identity(self):
            eta = np.array([0.001, 1 / 15, 0.5, 0.99, 1.0])
            q, p = np.full(5, 30.0), np.full(5, 20.0)
            w = 10 / eta
            identity = run_similarity_scaling(
                q, p, w, 60.0, 300, 100, 1500, 300, 100, 1500, 1.2, 1.2, True, 0.6
            )
            np.testing.assert_allclose(identity[2], w, rtol=1e-12)
            np.testing.assert_allclose(identity[4], eta, rtol=1e-12)
            result = run_similarity_scaling(
                q, p, w, 60.0, 300, 100, 1500, 600, 200, 1500, 1.2, 1.2, True, 0.6
            )
            expected = 1 - (1 - eta) * (0.5 + 0.5 * 4 ** (-0.2))
            np.testing.assert_allclose(result[4], expected)
            np.testing.assert_allclose(result[0] * result[1] / 60 / result[2], expected)
            # 非양수 예측 영역: 명시된 기본 상사 fallback, 양수 동력과 실제 효율 유지.
            small = run_similarity_scaling(
                30.0, 20.0, 150.0, 60.0, 300, 100, 1500, 150, 50, 1500, 1.2, 1.2, True
            )
            self.assertAlmostEqual(float(small[2]), 150 / 32)
            self.assertAlmostEqual(float(small[4]), 1 / 15)

        def test_reynolds_density_and_viscosity_ratio(self):
            args = (30.0, 300.0, 300.0, 60.0, 300, 100, 1500, 300, 100, 1500, 1.2)
            density_half = run_similarity_scaling(*args, 0.6, True, mu1=2e-5, mu2=2e-5)
            viscosity_double = run_similarity_scaling(*args, 1.2, True, mu1=2e-5, mu2=4e-5)
            expected = 1 - 0.5 * (0.5 + 0.5 * 2**0.2)
            self.assertAlmostEqual(float(density_half[4]), expected)
            self.assertAlmostEqual(float(viscosity_double[4]), expected)
            cancelled = run_similarity_scaling(*args, 0.6, True, mu1=2e-5, mu2=1e-5)
            self.assertAlmostEqual(float(cancelled[4]), 0.5)
            self.assertAlmostEqual(float(cancelled[2]), 150.0)
            for mu in [0, -1, np.nan]:
                with self.assertRaises(ValueError):
                    run_similarity_scaling(*args, 1.2, True, mu2=mu)

        def test_viscosity_temperature_and_wrapper_propagation(self):
            self.assertAlmostEqual(float(calc_air_viscosity(15)), 1.7894e-5, delta=1e-9)
            self.assertGreater(calc_air_viscosity(100), calc_air_viscosity(20))
            with self.assertRaises(ValueError):
                calc_air_viscosity(-273.15)
            args = (
                [20, 50, 80],
                [200, 100, 20],
                [200, 200, 150],
                [65, 60, 70],
                300.0,
                100.0,
                1000.0,
                100.0,
                200.0,
                1.2,
                0.9,
                800.0,
                4000.0,
                True,
                0.55,
            )
            candidates = find_derived_candidates(*args, c_w=0.8, mu1=1.8e-5, mu2=2.5e-5)
            self.assertTrue(candidates)
            c = candidates[0]
            result = scale_derived_curve(
                [c["q_ref"]],
                [c["p_ref"]],
                [c["w_ref"]],
                [c["noise_ref"]],
                300.0,
                100.0,
                1000.0,
                c["d2"],
                c["x2"],
                c["n2"],
                1.2,
                0.9,
                True,
                0.55,
                c_w=0.8,
                mu1=1.8e-5,
                mu2=2.5e-5,
            )
            np.testing.assert_allclose(
                [v[0] for v in result], [c[k] for k in ["q", "p", "w", "noise", "eff"]]
            )
            curve = _fan_set_curve_at_rpm(
                [30.0],
                [300.0],
                [300.0],
                [60.0],
                300,
                100,
                1500,
                300,
                100,
                1500,
                1.2,
                0.6,
                2,
                "병렬 (Parallel)",
                AXIAL_CATEGORY,
                True,
                mu1=2e-5,
                mu2=1e-5,
            )
            np.testing.assert_allclose([curve[i][0] for i in range(3)], [60, 150, 300])

        def test_intersection_unsorted_duplicates_and_invalid(self):
            result = find_system_intersection([100, 0, 0, np.nan], [0, 100, 100, 20], 0.01)
            self.assertAlmostEqual(result, 61.80339887498949)
            self.assertIsNone(find_system_intersection([10], [20], 1))

        def test_intersection_multiple_and_boundaries(self):
            self.assertEqual(find_system_intersection([0, 1, 2, 3], [1, -1, 1, -1], 0), 2.5)
            self.assertEqual(find_curve_intersection([0, 1, 2, 3], [1, -1, 1, -1], [0] * 4), 2.5)
            self.assertEqual(find_system_intersection([0, 1], [1, 0], 0), 1.0)
            self.assertEqual(find_system_intersection([0, 1], [0, 0], 0), 1.0)

        def test_scaling_identity_and_rpm_laws(self):
            for category in [CENTRIFUGAL_CATEGORY, AXIAL_CATEGORY, CROSSFLOW_CATEGORY]:
                args = (10.0, 100.0, 50.0, 60.0, 300.0, 100.0, 1000.0, 300.0, 100.0)
                identity = run_similarity_scaling(*args, 1000.0, 1.2, 1.2, True, 0.6, category)
                np.testing.assert_allclose(identity[:4], [10, 100, 50, 60])
                doubled = run_similarity_scaling(*args, 2000.0, 1.2, 1.2, False, 0.6, category)
                np.testing.assert_allclose(doubled[:3], [20, 400, 400])
                self.assertAlmostEqual(doubled[3] - 60, 50 * np.log10(2))

        def test_axial_geometry_ignores_height(self):
            result = run_similarity_scaling(
                10.0,
                100.0,
                50.0,
                60.0,
                300.0,
                0.0,
                1000.0,
                600.0,
                0.0,
                1000.0,
                1.2,
                1.2,
                fan_category=AXIAL_CATEGORY,
            )
            np.testing.assert_allclose(result[:3], [80, 400, 1600])

        def test_efficiency_array_zero_does_not_disable_other_points(self):
            result = run_similarity_scaling(
                np.array([0.0, 10.0]),
                np.array([100.0, 100.0]),
                np.array([50.0, 50.0]),
                np.array([np.nan, 60.0]),
                300.0,
                100.0,
                1000.0,
                300.0,
                100.0,
                2000.0,
                1.2,
                1.2,
                True,
                np.array([0.0, 0.6]),
            )
            self.assertEqual(result[2][0], 400.0)
            self.assertLess(result[2][1], 400.0)
            self.assertTrue(np.isnan(result[3][0]))

        def test_invalid_geometry_and_environment(self):
            with self.assertRaises(ValueError):
                run_similarity_scaling(1.0, 1.0, 1.0, 60.0, 300.0, 100.0, 0.0, 300.0, 100.0, 1000.0, 1.2, 1.2)
            for temp, alt in [(-273.15, 0), (20, 45000), (np.nan, 0)]:
                with self.subTest(temp=temp, alt=alt), self.assertRaises(ValueError):
                    calc_air_density(temp, alt)
            self.assertAlmostEqual(calc_air_density(20, 0), 1.2041183164, places=8)

        def test_optional_interpolation_uses_q_and_preserves_unknown(self):
            self.assertEqual(_optional_interp(15, [10, 20], [100, 80]), 90)
            self.assertIsNone(_optional_interp(5, [10, 20], [100, 80]))
            self.assertIsNone(_optional_interp(11, [10], [100]))
            self.assertEqual(_optional_interp(10, [10], [100]), 100)
            self.assertIsNone(_optional_interp(10, [10, 20], [np.nan, np.nan]))
            np.testing.assert_allclose(interpolate_optional_curve([0, 1, 10], [60, np.nan, 80]), [60, 62, 80])

        def test_rpm_extrapolation_permission_and_provenance(self):
            args = (
                112.5,
                25.0,
                [0, 50, 100],
                [200, 150, 50],
                [100] * 3,
                [60] * 3,
                300.0,
                100.0,
                1000.0,
                300.0,
                100.0,
                1.2,
                1.2,
                1,
                "병렬 (Parallel)",
                None,
            )
            self.assertIsNone(_solve_required_rpm(*args))
            rpm, details = _solve_required_rpm(*args, allow_limited_extrapolation=True, return_details=True)
            self.assertAlmostEqual(rpm, 1000)
            self.assertTrue(details["used_extrapolation"])
            self.assertEqual(details["reference_q_max"], 100)
            self.assertAlmostEqual(details["extrapolation_pct"], 12.5)

        def test_rpm_inverse_with_parallel_and_series(self):
            for arrangement, q, p in [("병렬 (Parallel)", 320.0, 100.0), ("직렬 (Series)", 160.0, 200.0)]:
                result = _solve_required_rpm(
                    q,
                    p,
                    [20, 40, 60],
                    [200, 100, 0],
                    [100] * 3,
                    [60] * 3,
                    300.0,
                    100.0,
                    1000.0,
                    600.0,
                    200.0,
                    1.2,
                    1.2,
                    2,
                    arrangement,
                    AXIAL_CATEGORY,
                    preferred_rpm=500,
                )
                self.assertAlmostEqual(result, 500.0)

        def test_resistance_fit_duplicate_flows(self):
            k, exponent, error = fit_system_resistance([10, 10], [100, 100])
            self.assertAlmostEqual(k, 1.0)
            self.assertEqual(exponent, 2.0)
            self.assertAlmostEqual(error, 0.0)

        def test_curve_normalization_and_optional_completion(self):
            frame = pd.DataFrame(
                {
                    "풍량 Q (CMM)": [20, 10, 10, np.inf],
                    "정압 P (Pa)": [100, 200, 200, 100],
                    "정압효율 ηs (%)": [50, 50, 50, 50],
                }
            )
            clean, _, _ = _prepare_optional_fan_curve(frame)
            self.assertEqual(len(clean), 2)
            np.testing.assert_allclose(clean["축동력 W (W)"], [200 / 3, 200 / 3])
            self.assertTrue(clean["소음 Lw (dB(A))"].isna().all())

        def test_derived_design_search_and_curve_consistency(self):
            args = (
                [20, 50, 80],
                [200, 100, 20],
                [200, 200, 150],
                [65, 60, 70],
                300.0,
                100.0,
                1000.0,
                100.0,
                200.0,
                1.2,
                1.2,
                800.0,
                4000.0,
                False,
                0.85,
            )
            candidates = find_derived_candidates(*args)
            self.assertGreater(len(candidates), 2)
            self.assertGreater(np.ptp([c["d2"] for c in candidates]), 10.0)
            candidate = next(c for c in candidates if c["q_ref"] == 50)
            self.assertAlmostEqual(candidate["d2"], 300 * 2**0.25)
            self.assertAlmostEqual(candidate["n2"], 1000 * 2**0.25)
            np.testing.assert_allclose([candidate["q"], candidate["p"], candidate["w"]], [100, 200, 800])
            self.assertTrue(all(c["d2"] <= 800 and c["n2"] <= 4000 for c in candidates))
            predicted = scale_derived_curve(
                [50],
                [100],
                [200],
                [60],
                300,
                100,
                1000,
                candidate["d2"],
                candidate["x2"],
                candidate["n2"],
                1.2,
                1.2,
                False,
                0.85,
            )
            np.testing.assert_allclose([a[0] for a in predicted[:3]], [100, 200, 800])

        def test_missing_noise_not_ranked_as_measurement(self):
            candidates = find_derived_candidates(
                [20, 50, 80],
                [200, 100, 20],
                [200, 200, 150],
                [np.nan] * 3,
                300,
                100,
                1000,
                100,
                200,
                1.2,
                1.2,
                800,
                4000,
                True,
                0.85,
            )
            self.assertTrue(candidates)
            self.assertTrue(all(np.isnan(c["noise"]) for c in candidates))
            self.assertTrue(
                all(np.isinf(derived_candidate_score(c, "최소 소음 순 (Min Noise)")) for c in candidates)
            )

        def test_unit_roundtrip_and_table_preservation(self):
            canonical = pd.DataFrame(
                dict(q=[10.0, 35.0], p=[100.0, 250.0], w=[150.0, 300.0], noise=[60.0, np.nan])
            )
            for fu in UNIT_FLOW:
                for pu in UNIT_PRESSURE:
                    for wu in UNIT_POWER:
                        restored = pq_canonical_frame(pq_display_frame(canonical, fu, pu, wu), fu, pu, wu)
                        np.testing.assert_allclose(restored, canonical, equal_nan=True)

        def test_scalar_input_changes_unit_without_reinterpreting_number(self):
            state = {}

            def fake_input(label, **kwargs):
                key = kwargs["key"]
                if key not in state:
                    state[key] = kwargs["value"]
                return state[key]

            with (
                patch.object(st, "session_state", state),
                patch.object(st, "number_input", side_effect=fake_input),
            ):
                self.assertEqual(
                    unit_number_input("Q", quantity="flow", unit="CMM (m³/min)", value=35, key="q"), 35
                )
                state["q"] = 42.0
                self.assertEqual(
                    unit_number_input("Q", quantity="flow", unit="CMH (m³/hr)", value=35, key="q"), 2520
                )
                self.assertEqual(
                    unit_number_input("Q", quantity="flow", unit="CMM (m³/min)", value=35, key="q"), 42
                )

        def test_project_storage_isolated_atomic_and_legacy_read_only(self):
            with tempfile.TemporaryDirectory() as directory:
                db = Path(directory) / "inputs.sqlite3"
                legacy = Path(directory) / "v9.json"
                legacy.write_text('{"q": 42}')
                self.assertEqual(read_project_inputs(db, "default", legacy), {"q": 42})
                self.assertEqual(read_project_inputs(db, "other", legacy), {})
                write_project_inputs(db, "A", {"q": 10, "p": 20})
                write_project_inputs(db, "B", {"q": 99})
                with ThreadPoolExecutor(max_workers=4) as pool:
                    list(pool.map(lambda i: write_project_inputs(db, "A", {f"field{i}": i}), range(12)))
                write_project_inputs(db, "A", {"q": 11})
                state = read_project_inputs(db, "A")
                self.assertEqual(state["q"], 11)
                self.assertEqual(state["p"], 20)
                self.assertTrue(all(state[f"field{i}"] == i for i in range(12)))
                self.assertEqual(read_project_inputs(db, "B"), {"q": 99})
                self.assertEqual(legacy.read_text(), '{"q": 42}')

    result = unittest.TextTestRunner(verbosity=2).run(
        unittest.defaultTestLoader.loadTestsFromTestCase(RegressionTests)
    )
    if not result.wasSuccessful():
        raise SystemExit(1)


if __name__ == "__main__":
    import sys

    if "--self-test" in sys.argv:
        run_self_tests()
    else:
        main()
